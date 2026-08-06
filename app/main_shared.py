from __future__ import annotations

from app.main_pages import *

import asyncio
import base64
import copy
import csv
import fcntl
import hashlib
import hmac
import html
import io
import json
import math
import os
import platform
import queue
import random
import re
import secrets
import shlex
import shutil
import socket
import sqlite3
import subprocess
import sys
import tempfile
import threading
import time
import uuid
import zipfile
from collections import Counter, defaultdict
from contextlib import contextmanager
from dataclasses import replace
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple
import xml.etree.ElementTree as ET
from zoneinfo import ZoneInfo
from urllib.parse import parse_qs, quote, urlparse

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from fastapi import Body, FastAPI, File, Form, HTTPException, Request, Response, UploadFile, WebSocket, WebSocketDisconnect
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, RedirectResponse, StreamingResponse
from fastapi.routing import APIRoute
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

from app.async_pipeline import CircuitBreaker, TokenBucketRateLimiter, fingerprint_payload
from app.ad_production_closure import MetaActivityReadonlyService, persist_activity_changes
from app.meta_ad_account_access import (
    MetaAdAccountAccessPolicy,
    access_summary,
    classify_meta_exception,
)
from app.meta_api_budget import MetaRateLimitBlocked
from app.ad_daily_report import (
    FixtureRealConversionProvider,
    RECOMMENDATION_RULE_VERSION,
    TugaoRealConversionProvider,
    build_daily_report_from_dashboard_snapshot,
    export_daily_report_xlsx,
    load_persisted_daily_report,
    persist_daily_report,
    report_from_dict,
    recommendation_history_payload,
    recommendation_review_payload,
    report_to_dict,
)
from app.ad_creative_intelligence import (
    CREATIVE_FEATURE_FLAGS,
    MetaCreativeSyncService,
    build_creative_intelligence_payload,
    creative_asset_from_meta_payload,
    ensure_creative_intelligence_tables,
    normalize_meta_ad_account_id,
    normalize_feature_flags,
    persist_creative_assets,
)
from app.creative_image_generation import (
    CreativeImageGenerationBrief,
    ExternalImageProviderConfig,
    PROVIDER_CHATGPT_PRO_MANUAL,
    PROVIDER_HERMES_IMAGE2_AGENT,
    PROVIDER_LOCAL_PRODUCTION_PNG,
    InvalidCreativeImageError,
    approve_creative_experiment_generation,
    chatgpt_pro_workbench_status,
    claim_hermes_image2_generation_task,
    claim_next_chatgpt_pro_job,
    cancel_hermes_image2_generation_task,
    confirm_creative_experiment_binding,
    create_chatgpt_pro_job,
    create_hermes_image2_generation_job,
    create_review_record,
    creative_image_auto_approval_eligible,
    create_feed_image_generation,
    creative_experiment_binding_status,
    enrich_creative_image_names,
    creative_image_download_filename,
    detect_creative_experiment_binding,
    delete_chatgpt_pro_job,
    fail_hermes_image2_generation_task,
    get_chatgpt_pro_job,
    get_creative_generation_task,
    get_creative_pro_generation_status,
    generated_image_experiment_tracking,
    heartbeat_hermes_image2_generation_task,
    latest_generated_images,
    list_hermes_image2_generation_tasks,
    list_chatgpt_pro_jobs,
    mark_generated_image_adopted,
    mark_chatgpt_pro_job_completed,
    reject_creative_experiment_binding,
    retry_hermes_image2_generation_task,
    save_chatgpt_pro_uploaded_image,
    save_hermes_image2_uploaded_image,
    external_image_provider_readiness,
    next_hermes_image2_generation_task,
    start_hermes_image2_generation_task,
    update_chatgpt_pro_job_analysis,
    update_chatgpt_pro_job_generation_plan,
    upload_final_creative_asset_hash,
)
from app.im_diagnostics import (
    claim_im_llm_diagnosis_task,
    complete_im_llm_diagnosis_task,
    create_im_llm_diagnosis_task,
    create_im_llm_diagnosis_tasks_for_latest_run,
    ensure_im_diagnostics_tables,
    fail_im_llm_diagnosis_task,
    generate_im_diagnosis_fixtures,
    get_im_llm_diagnosis_task,
    im_conversation_detail,
    im_conversations_payload,
    im_diagnostics_summary,
    next_im_llm_diagnosis_task,
    persist_im_diagnostics_payload,
    review_im_conversation_diagnosis,
    run_im_diagnosis,
    update_im_script_suggestion_status,
)
from app.im_diagnostics_api import (
    DEFAULT_IM_DIAGNOSTICS_BASE_URL,
    TimeTradeImDiagnosticsClient,
    aggregate_reception_mode_daily,
    fetch_im_diagnostics_payload,
)
from app.im_result_message_facts import im_result_message_detail_rows
from app.tugao_bi import TugaoBindSuccessClient, sync_tugao_bind_success_events
from app.approval_accounts import (
    WHATSAPP_APPROVAL_RUNTIME_CONFIG_KEYS,
    apply_baileys_runtime_assignment_defaults as _apply_baileys_runtime_assignment_defaults,
    apply_whatsapp_approval_runtime_defaults as _apply_whatsapp_approval_runtime_defaults,
    baileys_default_provider_mode_for_responsible_type as _baileys_default_provider_mode_for_responsible_type,
    default_baileys_provider_base_url as _default_baileys_provider_base_url,
)
from app.baileys_accounts import (
    default_baileys_account_id_for_account_key,
    first_baileys_account_id,
    resolve_baileys_account_id_for_card,
)
from app.crm_adapter import LiveCrmAdapter
from app.lark_cli_adapter import LarkCliReplyAdapter
from app.native_ocr import normalize_native_ocr_fields
from app.ocr_adapter import RapidOcrAdapter
from app.ops_auth import (
    OPS_AUTH_ALLOWED_ROLES,
    OPS_AUTH_BUSINESS_ROLES,
    OPS_AUTH_INTERNAL_HEADER,
    OPS_AUTH_ROLE_ADMIN,
    OPS_AUTH_ROLE_CUSTOMER_SERVICE,
    OPS_AUTH_ROLE_INTERNAL,
    OPS_AUTH_ROLE_OPERATOR,
    OPS_AUTH_ROLE_SUPER_ADMIN,
    OPS_AUTH_SESSION_COOKIE,
    normalize_ops_role,
    ops_role_is_business,
)
from app.operation_tasks import (
    build_whatsapp_approval_task_envelope,
    effective_whatsapp_approval_task_wait_timeout,
    is_whatsapp_approval_operation_task_type,
    operation_task_account_key_from_object_key,
    operation_task_lease_expiry_status,
    operation_task_is_terminal_status,
    parse_operation_task_row,
    operation_task_should_retry,
    operation_task_terminal_failure_status,
    whatsapp_approval_operation_from_task_type,
    whatsapp_approval_task_specs,
)
from app.production_ops import (
    build_success_notifications,
    expand_notify_profile_targets,
    fetch_json,
    format_lark_alert,
    load_json_state,
    requester_fingerprint,
    save_json_state,
    should_suppress_lark_alert,
)
from app.registration_group_truth import (
    APPROVAL_TRUTH_PENDING_TTL_SECONDS,
    APPROVAL_TRUTH_UNKNOWN_TTL_SECONDS,
    APPROVAL_TRUTH_ZERO_TTL_SECONDS,
    build_approval_queue_truth_from_truth_state,
    build_truth_state,
    normalize_int_or_none,
    now_utc,
    serialize_membership_verifier,
)
from app.registration_truth_diagnostics import (
    build_diagnostic_approval_queue_truth_view,
    build_pending_truth_match_keys,
    normalize_pending_truth_history_entry,
    pending_truth_snapshot_group_state,
    select_pending_truth_confirmed_empty_candidate,
    select_pending_truth_confirmed_pending_candidate,
)
from app.realtime_approval_state import RealtimeApprovalStateStore
from app.report_routes import create_report_router
from app.schema_migrations import apply_schema_migration_registry
from app.timo_incremental_materialization import ensure_timo_incremental_schema
from app.streamer_analytics import ensure_streamer_analytics_views
from app.streamer_analytics_routes import create_streamer_analytics_router
from app.streamer_history_export import (
    build_streamer_history_xlsx,
    fetch_linky_streamer_profile,
    fetch_linky_streamer_history,
    load_covered_dates,
    load_local_revenue_rows,
    lookup_streamer_first_join,
    merge_revenue_calendar,
    normalize_history_app,
    normalize_streamer_id,
    normalize_timo_revenue_export_row,
    uncovered_dates,
)
from app.streamer_roi import ensure_streamer_roi_tables
from app.secret_redaction import compact_runtime_log_payload as _compact_runtime_log_payload
from app.secret_redaction import summarize_startup_health_payload as _summarize_startup_health_payload
from app.secret_redaction import redact_sensitive_payload as _redact_sensitive_payload
from app.sqlite_observability import connect_observed_sqlite, sqlite_observability_snapshot
from app.sqlite_bootstrap import ensure_sqlite_ready, sqlite_busy_timeout_ms
from app.sqlite_write_queue import (
    SQLiteWriteQueueError,
    db_writer_enabled,
    db_writer_required,
    submit_sqlite_write_job,
)
from app.timo_auth_station import (
    AuthStationDeviceBindingRequest,
    TimoAuthStationService,
    create_timo_auth_station_public_router,
    create_timo_auth_station_router,
    _station_device_readiness,
)
from app.timo_guild_executor import TIMO_DEFAULT_API_BASE_URL, TimoGuildExecutor

SUGO_DEFAULT_API_BASE_URL = 'https://union.sugo.com/union_leader/api'
SOGO_DEFAULT_API_BASE_URL = SUGO_DEFAULT_API_BASE_URL
SUGO_APP_NAME = 'sugo'
SUGO_LEGACY_APP_NAME = 'sogo'
SUGO_APP_NAMES = (SUGO_APP_NAME, SUGO_LEGACY_APP_NAME)
from app.whatsapp_approval_runtime import (
    BAILEYS_PROVIDER_MODES,
    DefaultWhatsAppApprovalRuntimeAdapter,
    RUNTIME_MODE_KEYS,
    resolve_whatsapp_approval_provider_mode,
)
from app.whatsapp_login_state import enrich_whatsapp_login_state, map_whatsapp_login_state
def _build_ops_client_version(*, respect_override: bool = True) -> str:
    override = str(os.getenv('OPS_CLIENT_VERSION') or os.getenv('APP_VERSION') or '').strip()
    if respect_override and override:
        return override[:80]
    hasher = hashlib.sha256()
    for path in [Path(__file__).resolve(), Path(__file__).resolve().parent / 'static' / 'ops' / 'common.js']:
        try:
            stat = path.stat()
            hasher.update(str(path.name).encode('utf-8'))
            hasher.update(str(stat.st_mtime_ns).encode('utf-8'))
            hasher.update(str(stat.st_size).encode('utf-8'))
        except OSError:
            hasher.update(str(path).encode('utf-8'))
    return hasher.hexdigest()[:16]


OPS_CLIENT_VERSION = _build_ops_client_version()
OPS_CLIENT_STARTED_AT = datetime.now(timezone.utc).isoformat()
GROUP_ATMOSPHERE_BUSINESS_TIMEZONE = ZoneInfo('Asia/Shanghai')


def _ops_runtime_version_state() -> Dict[str, Any]:
    current_version = _build_ops_client_version()
    stale = bool(current_version and current_version != OPS_CLIENT_VERSION)
    return {
        'version': current_version,
        'process_version': OPS_CLIENT_VERSION,
        'started_at': OPS_CLIENT_STARTED_AT,
        'stale': stale,
        'message_cn': '后端进程仍在运行旧代码，请重启 8011 后再操作。' if stale else '后端进程版本与磁盘代码一致。',
    }


def _group_atmosphere_business_date(now: Optional[datetime] = None) -> str:
    current = now or datetime.now(timezone.utc)
    if current.tzinfo is None:
        current = current.replace(tzinfo=timezone.utc)
    return current.astimezone(GROUP_ATMOSPHERE_BUSINESS_TIMEZONE).date().isoformat()


def _group_atmosphere_business_day_bounds_utc(now: Optional[datetime] = None) -> tuple[str, str]:
    current = now or datetime.now(timezone.utc)
    if current.tzinfo is None:
        current = current.replace(tzinfo=timezone.utc)
    local_date = current.astimezone(GROUP_ATMOSPHERE_BUSINESS_TIMEZONE).date()
    local_start = datetime(
        local_date.year,
        local_date.month,
        local_date.day,
        tzinfo=GROUP_ATMOSPHERE_BUSINESS_TIMEZONE,
    )
    utc_start = local_start.astimezone(timezone.utc)
    utc_end = utc_start + timedelta(days=1)
    return utc_start.isoformat(), utc_end.isoformat()


PHONE_PREFIX_COUNTRY_MAP = {
    '1': 'United States',
    '7': 'Russia',
    '20': 'Egypt',
    '27': 'South Africa',
    '31': 'Netherlands',
    '32': 'Belgium',
    '33': 'France',
    '34': 'Spain',
    '39': 'Italy',
    '44': 'United Kingdom',
    '49': 'Germany',
    '52': 'Mexico',
    '53': 'Cuba',
    '55': 'Brazil',
    '56': 'Chile',
    '57': 'Colombia',
    '58': 'Venezuela',
    '60': 'Malaysia',
    '61': 'Australia',
    '62': 'Indonesia',
    '63': 'Philippines',
    '64': 'New Zealand',
    '65': 'Singapore',
    '66': 'Thailand',
    '81': 'Japan',
    '82': 'South Korea',
    '84': 'Vietnam',
    '86': 'China',
    '90': 'Turkey',
    '91': 'India',
    '92': 'Pakistan',
    '93': 'Afghanistan',
    '94': 'Sri Lanka',
    '95': 'Myanmar',
    '98': 'Iran',
    '212': 'Morocco',
    '213': 'Algeria',
    '216': 'Tunisia',
    '218': 'Libya',
    '220': 'Gambia',
    '221': 'Senegal',
    '233': 'Ghana',
    '234': 'Nigeria',
    '251': 'Ethiopia',
    '254': 'Kenya',
    '255': 'Tanzania',
    '256': 'Uganda',
    '351': 'Portugal',
    '352': 'Luxembourg',
    '353': 'Ireland',
    '354': 'Iceland',
    '355': 'Albania',
    '356': 'Malta',
    '357': 'Cyprus',
    '358': 'Finland',
    '380': 'Ukraine',
    '420': 'Czech Republic',
    '852': 'Hong Kong',
    '853': 'Macau',
    '855': 'Cambodia',
    '856': 'Laos',
    '880': 'Bangladesh',
    '886': 'Taiwan',
    '961': 'Lebanon',
    '962': 'Jordan',
    '963': 'Syria',
    '964': 'Iraq',
    '965': 'Kuwait',
    '966': 'Saudi Arabia',
    '971': 'United Arab Emirates',
    '972': 'Israel',
    '973': 'Bahrain',
    '974': 'Qatar',
    '975': 'Bhutan',
    '976': 'Mongolia',
    '977': 'Nepal',
    '998': 'Uzbekistan',
}

COUNTRY_LABEL_ALIASES = {
    'id': 'Indonesia',
    '62': 'Indonesia',
    '+62': 'Indonesia',
    'indonesia': 'Indonesia',
    'br': 'Brazil',
    '55': 'Brazil',
    '+55': 'Brazil',
    'brazil': 'Brazil',
    'mx': 'Mexico',
    '52': 'Mexico',
    '+52': 'Mexico',
    'mexico': 'Mexico',
    've': 'Venezuela',
    '58': 'Venezuela',
    '+58': 'Venezuela',
    'venezuela': 'Venezuela',
    '委内瑞拉': 'Venezuela',
    'cl': 'Chile',
    '56': 'Chile',
    '+56': 'Chile',
    'chile': 'Chile',
    '智利': 'Chile',
    'co': 'Colombia',
    '57': 'Colombia',
    '+57': 'Colombia',
    'colombia': 'Colombia',
    '哥伦比亚': 'Colombia',
}

PHONE_LOCALIZED_NUMBER_RULES = {
    'Venezuela': {
        'country_code': '58',
        'national_lengths': {10},
        'local_lengths': {11},
        'trunk_prefixes': ('0',),
        'leading_digits': ('4',),
    },
    'Chile': {
        'country_code': '56',
        'national_lengths': {9},
        'local_lengths': {9},
        'trunk_prefixes': (),
        'leading_digits': ('9',),
    },
    'Colombia': {
        'country_code': '57',
        'national_lengths': {10},
        'local_lengths': {10},
        'trunk_prefixes': (),
        'leading_digits': ('3',),
    },
}


def _parse_config_list(value: Any) -> List[str]:
    if isinstance(value, (list, tuple, set)):
        raw_items = value
    else:
        raw_items = re.split(r'[\s,]+', str(value or '').strip())
    seen = set()
    items: List[str] = []
    for item in raw_items:
        normalized = str(item or '').strip()
        if not normalized or normalized in seen:
            continue
        seen.add(normalized)
        items.append(normalized)
    return items


def _dashboard_float(value: Any) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value) if math.isfinite(float(value)) else 0.0
    text = str(value).strip()
    if not text or text.upper() in {'N/A', 'NA', 'NULL', '-'}:
        return 0.0
    cleaned = re.sub(r'[^0-9.\-]', '', text.replace(',', ''))
    if cleaned in {'', '-', '.', '-.'}:
        return 0.0
    try:
        parsed = float(cleaned)
    except ValueError:
        return 0.0
    return parsed if math.isfinite(parsed) else 0.0


def _normalized_csv_key(value: str) -> str:
    text = str(value or '').strip().lower()
    if re.search(r'[\u4e00-\u9fff]', text):
        return re.sub(r'[\s_\-()/（）]+', '', text)
    return re.sub(r'[^a-z0-9]+', '_', text).strip('_')


def _csv_lookup(row: Dict[str, Any], candidates: List[str]) -> str:
    normalized = {_normalized_csv_key(key): value for key, value in (row or {}).items()}
    for candidate in candidates:
        key = _normalized_csv_key(candidate)
        if key in normalized:
            return str(normalized.get(key) or '').strip()
    return ''


def _extract_meta_ad_id_from_text(value: Any) -> str:
    return str(_extract_meta_ads_manager_context(value).get('ad_id') or '')


def _meta_ad_account_candidates_for_context(
    configured_account_ids: Iterable[Any],
    *,
    link_context: Optional[Dict[str, str]] = None,
    body: Optional[Dict[str, Any]] = None,
) -> List[str]:
    candidates: List[str] = []
    for value in [
        (body or {}).get('account_id'),
        (link_context or {}).get('account_id'),
        *(configured_account_ids or []),
    ]:
        normalized = normalize_meta_ad_account_id(value)
        if normalized and normalized not in candidates:
            candidates.append(normalized)
    return candidates


def _extract_meta_ads_manager_context(value: Any) -> Dict[str, str]:
    raw = str(value or '').strip()
    if not raw:
        return {}
    if re.fullmatch(r'\d{6,}', raw):
        return {'ad_id': raw}

    def _candidate_from_token(token: Any) -> str:
        text = str(token or '').strip()
        for part in re.split(r'[\s,]+', text):
            normalized = part.strip()
            if re.fullmatch(r'\d{6,}', normalized):
                return normalized
        return ''

    context: Dict[str, str] = {}
    query = urlparse(raw).query
    params = parse_qs(query)
    mappings = {
        'ad_id': ('selected_ad_ids', 'selected_ad_id', 'ad_id', 'ad_ids'),
        'account_id': ('act', 'account_id', 'ad_account_id'),
        'business_id': ('business_id', 'bm_id'),
        'campaign_id': ('selected_campaign_ids', 'campaign_id', 'campaign_ids'),
        'adset_id': ('selected_adset_ids', 'adset_id', 'adset_ids'),
    }
    for output_key, query_keys in mappings.items():
        for key in query_keys:
            for item in params.get(key, []):
                candidate = _candidate_from_token(item)
                if candidate:
                    context[output_key] = candidate
                    break
            if context.get(output_key):
                break

    for output_key, query_keys in mappings.items():
        if context.get(output_key):
            continue
        pattern = r'(?:' + '|'.join(re.escape(key) for key in query_keys) + r')=([^&#\s]+)'
        for match in re.finditer(pattern, raw, re.I):
            candidate = _candidate_from_token(match.group(1))
            if candidate:
                context[output_key] = candidate
                break
    return context


def _csv_event_lookup(row: Dict[str, Any], event_name: str, *, metric: str = 'unique_users') -> str:
    suffixes = {
        'unique_users': ['unique users', 'unique user', 'users'],
        'event_counter': ['event counter', 'events', 'counter', 'count'],
        'sales': ['sales in usd', 'sales', 'revenue'],
    }.get(metric, [])
    candidates = [event_name]
    for suffix in suffixes:
        candidates.extend([
            f'{event_name} ({suffix})',
            f'{event_name}_{suffix}',
            f'{event_name} {suffix}',
        ])
    return _csv_lookup(row, candidates)


def _dashboard_response_text(response: Any) -> str:
    content = getattr(response, 'content', b'')
    if isinstance(content, bytes) and content:
        for encoding in ('utf-8-sig', 'utf-8', 'gb18030'):
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
    return str(getattr(response, 'text', '') or '')


def _parse_dashboard_date(value: Any) -> Optional[datetime.date]:
    text = str(value or '').strip()
    if not text:
        return None
    text = text[:10]
    for fmt in ('%Y-%m-%d', '%Y/%m/%d', '%d/%m/%Y', '%m/%d/%Y'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(text).date()
    except ValueError:
        return None


def _empty_ad_metrics() -> Dict[str, float]:
    return {
        'cost': 0.0,
        'installs': 0.0,
        'af_installs': 0.0,
        'registrations': 0.0,
        'meta_installs': 0.0,
        'meta_registrations': 0.0,
        'af_registrations': 0.0,
        'onsite_registrations': 0.0,
        'high_value_users': 0.0,
        'im_entries': 0.0,
        'auto_apply_message_users': 0.0,
        'im_first_replies': 0.0,
        'im_step2_triggers': 0.0,
        'im_manual_reply_3': 0.0,
        'im_user_message_ge_5_users': 0.0,
        'im_link_clicks': 0.0,
        'im_link_click_users': 0.0,
        'link_click_users': 0.0,
        'linky_register_users': 0.0,
        'bind_success_users': 0.0,
        'crm_succeed_users': 0.0,
        'high_intent_im_users': 0.0,
        'guild_joins': 0.0,
        'promotion_guild_joins': 0.0,
        'organic_guild_joins': 0.0,
        'meta_guild_joins': 0.0,
        'af_guild_joins': 0.0,
        'purchases': 0.0,
        'revenue': 0.0,
        'clicks': 0.0,
        'link_clicks': 0.0,
        'impressions': 0.0,
        'reach': 0.0,
    }


def _finalize_ad_metrics(metrics: Dict[str, float]) -> Dict[str, Any]:
    result = {key: round(float(value or 0.0), 4) for key, value in _empty_ad_metrics().items()}
    result.update({key: round(float(metrics.get(key) or 0.0), 4) for key in result.keys()})
    cost = float(result.get('cost') or 0.0)
    installs = float(result.get('installs') or 0.0)
    af_installs = float(result.get('af_installs') or 0.0)
    meta_installs = float(result.get('meta_installs') or 0.0)
    meta_registrations = float(result.get('meta_registrations') or 0.0)
    af_registrations = float(result.get('af_registrations') or 0.0)
    onsite_registrations = float(result.get('onsite_registrations') or 0.0)
    raw_registrations = float(result.get('registrations') or 0.0)
    registrations = onsite_registrations or meta_registrations or raw_registrations or af_registrations
    high_value_users = float(result.get('high_value_users') or 0.0)
    im_entries = float(result.get('im_entries') or 0.0)
    im_first_replies = float(result.get('im_first_replies') or 0.0)
    im_manual_reply_3 = float(result.get('im_manual_reply_3') or 0.0)
    im_user_message_ge_5_users = float(result.get('im_user_message_ge_5_users') or 0.0)
    im_child_cap = im_entries if im_entries > 0 else 0.0
    result['im_first_replies'] = round(min(im_first_replies, im_child_cap), 4)
    result['im_manual_reply_3'] = round(min(im_manual_reply_3, im_child_cap), 4)
    result['im_user_message_ge_5_users'] = round(min(im_user_message_ge_5_users, im_child_cap), 4)
    im_first_replies = float(result.get('im_first_replies') or 0.0)
    im_manual_reply_3 = float(result.get('im_manual_reply_3') or 0.0)
    high_intent_im_users = float(result.get('high_intent_im_users') or 0.0)
    if not high_intent_im_users:
        high_intent_im_users = max(
            float(result.get('im_link_click_users') or 0.0),
            float(result.get('link_click_users') or 0.0),
            float(result.get('linky_register_users') or 0.0),
            float(result.get('bind_success_users') or 0.0),
        )
    high_intent_im_users = min(high_intent_im_users, im_child_cap)
    result['high_intent_im_users'] = round(high_intent_im_users, 4)
    guild_joins = float(result.get('guild_joins') or 0.0)
    promotion_guild_joins = float(result.get('promotion_guild_joins') or 0.0)
    organic_guild_joins = float(result.get('organic_guild_joins') or 0.0)
    meta_guild_joins = float(result.get('meta_guild_joins') or 0.0)
    af_guild_joins = float(result.get('af_guild_joins') or 0.0)
    clicks = float(result.get('clicks') or 0.0)
    link_clicks = float(result.get('link_clicks') or 0.0) or clicks
    impressions = float(result.get('impressions') or 0.0)
    reach = float(result.get('reach') or 0.0)
    install_basis = (meta_installs if cost and meta_installs else 0.0) or installs or meta_installs
    result['registrations'] = round(registrations, 4)
    result['onsite_registrations'] = round(onsite_registrations, 4)
    result['installs'] = round(install_basis, 4)
    result['cpi'] = round(cost / install_basis, 4) if install_basis else 0.0
    result['meta_cpi'] = round(cost / meta_installs, 4) if meta_installs else 0.0
    result['af_cpi'] = round(cost / af_installs, 4) if af_installs else 0.0
    result['meta_registration_cost'] = round(cost / meta_registrations, 4) if meta_registrations else 0.0
    result['af_registration_cost'] = round(cost / af_registrations, 4) if af_registrations else 0.0
    result['meta_join_cost'] = round(cost / meta_guild_joins, 4) if meta_guild_joins else 0.0
    result['af_join_cost'] = round(cost / af_guild_joins, 4) if af_guild_joins else 0.0
    result['install_gap'] = round(meta_installs - af_installs, 4)
    result['install_gap_rate'] = round((meta_installs - af_installs) / af_installs, 4) if af_installs else 0.0
    result['registration_gap'] = round(meta_registrations - af_registrations, 4)
    result['registration_gap_rate'] = round((meta_registrations - af_registrations) / af_registrations, 4) if af_registrations else 0.0
    result['join_gap'] = round(meta_guild_joins - guild_joins, 4)
    result['join_gap_rate'] = round((meta_guild_joins - guild_joins) / guild_joins, 4) if guild_joins else 0.0
    result['registration_cost'] = round(cost / onsite_registrations, 4) if onsite_registrations else 0.0
    result['high_value_cost'] = round(cost / high_value_users, 4) if high_value_users else 0.0
    result['im_cost'] = round(cost / im_entries, 4) if im_entries else 0.0
    result['join_cost'] = round(cost / guild_joins, 4) if guild_joins else 0.0
    result['roas'] = round(float(result.get('revenue') or 0.0) / cost, 4) if cost else 0.0
    result['cpm'] = round(cost / impressions * 1000, 4) if impressions else 0.0
    result['cpc'] = round(cost / link_clicks, 4) if link_clicks else 0.0
    result['all_click_cpc'] = round(cost / clicks, 4) if clicks else 0.0
    result['ctr'] = round(clicks / impressions, 4) if impressions else 0.0
    result['frequency'] = round(impressions / reach, 4) if reach else 0.0
    result['install_rate'] = round(install_basis / clicks, 4) if clicks else 0.0
    if not guild_joins and (promotion_guild_joins or organic_guild_joins):
        guild_joins = promotion_guild_joins + organic_guild_joins
        result['guild_joins'] = round(guild_joins, 4)
    result['registration_rate'] = round(onsite_registrations / install_basis, 4) if install_basis else 0.0
    result['high_value_rate'] = round(high_value_users / onsite_registrations, 4) if onsite_registrations else 0.0
    result['im_entry_rate'] = round(im_entries / high_value_users, 4) if high_value_users else 0.0
    result['im_first_reply_rate'] = round(im_first_replies / im_entries, 4) if im_entries else 0.0
    result['im_reply_3_rate'] = round(im_manual_reply_3 / im_entries, 4) if im_entries else 0.0
    result['high_intent_im_rate'] = round(high_intent_im_users / im_entries, 4) if im_entries else 0.0
    result['join_rate'] = round(guild_joins / im_entries, 4) if im_entries else 0.0
    return result


def _add_ad_metrics(target: Dict[str, float], row: Dict[str, Any], keys: Optional[List[str]] = None) -> None:
    for key in (keys or list(_empty_ad_metrics().keys())):
        target[key] = float(target.get(key) or 0.0) + float((row or {}).get(key) or 0.0)


def _normalize_ad_fact_account_value(row: Dict[str, Any]) -> str:
    data_source = str((row or {}).get('data_source') or '').strip().lower()
    if data_source in {'tugaofunnel', 'tugao_funnel', 'tugao_onsite_funnel'}:
        # Tugao funnel rows may contain external_app values such as Linky. They
        # describe the app-side funnel context, not the paid-media ad account.
        return ''
    return str((row or {}).get('app_id') or '').strip()


def _ad_fact_grain_key(row: Dict[str, Any]) -> Tuple[str, ...]:
    return (
        str((row or {}).get('date') or '').strip(),
        str((row or {}).get('data_source') or '').strip() or 'Unknown',
        str((row or {}).get('platform') or '').strip() or 'Unknown',
        _normalize_ad_fact_account_value(row),
        str((row or {}).get('appsflyer_app_id') or '').strip(),
        str((row or {}).get('country') or '').strip() or 'Unknown',
        str((row or {}).get('media_source') or '').strip(),
        str((row or {}).get('campaign') or '').strip() or '未命名',
        str((row or {}).get('ad_group') or '').strip(),
        str((row or {}).get('ad') or '').strip(),
        str((row or {}).get('source_type') or '').strip(),
    )


def _ad_fact_row_id(row: Dict[str, Any]) -> str:
    raw = json.dumps(_ad_fact_grain_key(row), ensure_ascii=False, sort_keys=False)
    return hashlib.sha256(raw.encode('utf-8')).hexdigest()


def _ad_country_is_unknown(value: Any) -> bool:
    return str(value or '').strip().lower() in {'', 'unknown', '未命名', '未知'}


AD_HISTORICAL_SETTLEMENT_UNALLOCATED_COUNTRY = 'HistoricalSettlementUnallocated'


def _ad_prepare_country_dimension(row: Dict[str, Any]) -> Dict[str, Any]:
    """Keep historical settlements explicit instead of disguising them as a country."""
    item = dict(row or {})
    data_source = str(item.get('data_source') or '').strip().lower()
    account_id = str(item.get('account_id') or item.get('ad_account_id') or '').strip().lower()
    historical_recovery = bool(item.get('historical_recovery')) or account_id == 'archived_settled_accounts'
    if data_source == 'meta' and historical_recovery and _ad_country_is_unknown(item.get('country')):
        item['country'] = AD_HISTORICAL_SETTLEMENT_UNALLOCATED_COUNTRY
        item['country_attribution_status'] = 'historical_settlement_unallocated'
        item['historical_recovery'] = True
    return item


def _ad_country_enrichment_key(row: Dict[str, Any], *, include_ad: bool = True) -> Tuple[str, ...]:
    return (
        str((row or {}).get('date') or '').strip(),
        str((row or {}).get('platform') or '').strip().lower(),
        _ad_account_label_or_missing((row or {}).get('app_id')).lower(),
        str((row or {}).get('campaign') or '').strip().lower(),
        str((row or {}).get('ad_group') or '').strip().lower(),
        str((row or {}).get('ad') or '').strip().lower() if include_ad else '',
    )


def _ad_meta_delivery_country_keys(row: Dict[str, Any]) -> List[Tuple[str, ...]]:
    date_key = str((row or {}).get('date') or '').strip()
    platform_key = str((row or {}).get('platform') or '').strip().lower()
    account_key = _normalize_ad_account_id_candidate(
        (row or {}).get('account_id') or (row or {}).get('ad_account_id')
    )
    campaign_id = str((row or {}).get('campaign_id') or '').strip()
    adset_id = str((row or {}).get('adset_id') or '').strip()
    ad_id = str((row or {}).get('ad_id') or '').strip()
    campaign_key = str((row or {}).get('campaign') or '').strip().casefold()
    ad_group_key = str((row or {}).get('ad_group') or '').strip().casefold()
    ad_key = str((row or {}).get('ad') or '').strip().casefold()
    if not date_key or not platform_key:
        return []
    keys: List[Tuple[str, ...]] = []
    if account_key and ad_id:
        keys.append(('ad_id', date_key, platform_key, account_key, ad_id))
    if account_key and adset_id:
        keys.append(('adset_id', date_key, platform_key, account_key, adset_id))
    if account_key and campaign_id:
        keys.append(('campaign_id', date_key, platform_key, account_key, campaign_id))
    if campaign_key:
        if account_key and ad_group_key and ad_key:
            keys.append(('ad', date_key, platform_key, account_key, campaign_key, ad_group_key, ad_key))
        if account_key and ad_group_key:
            keys.append(('ad_group', date_key, platform_key, account_key, campaign_key, ad_group_key))
        if account_key:
            keys.append(('campaign', date_key, platform_key, account_key, campaign_key))
        if ad_group_key and ad_key:
            keys.append(('ad_cross_account', date_key, platform_key, campaign_key, ad_group_key, ad_key))
        if ad_group_key:
            keys.append(('ad_group_cross_account', date_key, platform_key, campaign_key, ad_group_key))
        keys.append(('campaign_cross_account', date_key, platform_key, campaign_key))
    return keys


def _ad_enrich_countries_from_meta_delivery(rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Attribute countryless rows only from same-day Meta delivery facts."""
    countries_by_key: Dict[Tuple[str, ...], set[str]] = defaultdict(set)
    for row in rows or []:
        if not isinstance(row, dict):
            continue
        if str(row.get('data_source') or '').strip().lower() != 'meta':
            continue
        if str(row.get('country_attribution_status') or '').strip() != 'meta_delivery_country':
            continue
        country = normalize_country_label(row.get('country'))
        if _ad_country_is_unknown(country):
            continue
        for key in _ad_meta_delivery_country_keys(row):
            countries_by_key[key].add(country)

    enriched: List[Dict[str, Any]] = []
    for row in rows or []:
        if not isinstance(row, dict):
            continue
        status = str(row.get('country_attribution_status') or '').strip()
        country = str(row.get('country') or '').strip()
        needs_delivery_country = _ad_country_is_unknown(country) or status in {
            'unresolved_waiting_meta_delivery_country',
            'legacy_name_inferred',
        }
        if not needs_delivery_country:
            enriched.append(row)
            continue
        matched_country = ''
        matched_grain = ''
        ambiguous = False
        for key in _ad_meta_delivery_country_keys(row):
            candidates = countries_by_key.get(key, set())
            if len(candidates) == 1:
                matched_country = next(iter(candidates))
                matched_grain = key[0]
                break
            if len(candidates) > 1:
                ambiguous = True
                break
        if matched_country:
            updated = dict(row)
            updated['country'] = matched_country
            updated['country_attribution_status'] = 'meta_delivery_country_peer'
            updated['country_attribution_source'] = 'meta_insights_country_breakdown'
            updated['country_attribution_grain'] = matched_grain
            enriched.append(updated)
        elif ambiguous:
            updated = dict(row)
            updated['country'] = 'Unknown'
            updated['country_attribution_status'] = 'meta_delivery_country_ambiguous'
            updated['country_attribution_source'] = 'meta_insights_country_breakdown'
            enriched.append(updated)
        else:
            enriched.append(row)
    return enriched


def _ad_enrich_unknown_countries(rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    known_by_exact: Dict[Tuple[str, ...], str] = {}
    known_by_group: Dict[Tuple[str, ...], str] = {}
    for row in rows or []:
        if not isinstance(row, dict):
            continue
        country = str((row or {}).get('country') or '').strip()
        if _ad_country_is_unknown(country):
            continue
        exact_key = _ad_country_enrichment_key(row, include_ad=True)
        group_key = _ad_country_enrichment_key(row, include_ad=False)
        if exact_key and exact_key not in known_by_exact:
            known_by_exact[exact_key] = country
        if group_key and group_key not in known_by_group:
            known_by_group[group_key] = country
    enriched: List[Dict[str, Any]] = []
    for row in rows or []:
        if not isinstance(row, dict):
            continue
        if str(row.get('country_attribution_status') or '').strip() in {
            'unresolved_waiting_meta_delivery_country',
            'meta_delivery_country_ambiguous',
        }:
            enriched.append(row)
            continue
        if not _ad_country_is_unknown((row or {}).get('country')):
            enriched.append(row)
            continue
        exact_key = _ad_country_enrichment_key(row, include_ad=True)
        group_key = _ad_country_enrichment_key(row, include_ad=False)
        inferred = known_by_exact.get(exact_key) or known_by_group.get(group_key)
        if inferred:
            updated = dict(row)
            updated['country'] = inferred
            enriched.append(updated)
        else:
            enriched.append(row)
    return enriched


def _ad_daily_report_payload_has_unknown_country(payload: Dict[str, Any]) -> bool:
    def _has_literal_unknown_country(item: Dict[str, Any]) -> bool:
        return 'country' in item and _ad_country_is_unknown(item.get('country'))

    def _iter_items(value: Any) -> Iterable[Dict[str, Any]]:
        if isinstance(value, list):
            for item in value:
                if isinstance(item, dict):
                    yield item

    for item in _iter_items((payload or {}).get('ad_objects')):
        if _has_literal_unknown_country(item):
            return True
    for item in _iter_items((payload or {}).get('recommendations')):
        if _has_literal_unknown_country(item):
            return True
    for item in _iter_items((payload or {}).get('creative_test_plan')):
        if _has_literal_unknown_country(item):
            return True
    return False


def _ad_daily_report_payload_has_invalid_funnel_caps(payload: Dict[str, Any]) -> bool:
    def _iter_items(value: Any) -> Iterable[Dict[str, Any]]:
        if isinstance(value, list):
            for item in value:
                if isinstance(item, dict):
                    yield item

    for item in list(_iter_items((payload or {}).get('ad_objects'))) + list(_iter_items((payload or {}).get('recommendations'))):
        metrics = item.get('funnel_metrics') if isinstance(item.get('funnel_metrics'), dict) else None
        if metrics is None:
            evidence = item.get('evidence') if isinstance(item.get('evidence'), dict) else {}
            metrics = evidence.get('funnel_metrics') if isinstance(evidence.get('funnel_metrics'), dict) else {}
        im_entries = float((metrics or {}).get('im_entries') or 0.0)
        if im_entries <= 0:
            continue
        for key in ('user_engaged_im_users', 'high_intent_im_users'):
            if float((metrics or {}).get(key) or 0.0) > im_entries:
                return True
    return False


def _ad_daily_report_apply_funnel_caps(payload: Dict[str, Any]) -> Dict[str, Any]:
    def _normalize_label(value: Any) -> Any:
        if isinstance(value, str):
            return value.replace('链接/注册意向', '链接点击/注册/bind')
        if isinstance(value, list):
            return [_normalize_label(item) for item in value]
        if isinstance(value, dict):
            return {key: _normalize_label(item) for key, item in value.items()}
        return value

    def _iter_items(value: Any) -> Iterable[Dict[str, Any]]:
        if isinstance(value, list):
            for item in value:
                if isinstance(item, dict):
                    yield item

    def _cap_metrics(metrics: Dict[str, Any], *, spend: float = 0.0) -> None:
        im_entries = float((metrics or {}).get('im_entries') or 0.0)
        if im_entries <= 0:
            return
        for value_key, rate_key, cost_key in (
            ('user_engaged_im_users', 'user_engaged_im_rate', 'user_engaged_im_cost'),
            ('high_intent_im_users', 'high_intent_im_rate', 'high_intent_im_cost'),
        ):
            current = float((metrics or {}).get(value_key) or 0.0)
            if current <= im_entries:
                continue
            capped = im_entries
            metrics[value_key] = round(capped, 4)
            metrics[rate_key] = round(capped / im_entries, 6)
            if spend > 0 and capped > 0:
                metrics[cost_key] = round(spend / capped, 4)

    cleaned = _normalize_label(copy.deepcopy(payload or {}))
    for item in _iter_items(cleaned.get('ad_objects')):
        metrics = item.get('funnel_metrics') if isinstance(item.get('funnel_metrics'), dict) else None
        if metrics is not None:
            _cap_metrics(metrics, spend=float(item.get('spend') or 0.0))
    for item in _iter_items(cleaned.get('recommendations')):
        evidence = item.get('evidence') if isinstance(item.get('evidence'), dict) else {}
        metrics = evidence.get('funnel_metrics') if isinstance(evidence.get('funnel_metrics'), dict) else None
        if metrics is not None:
            _cap_metrics(metrics, spend=float(evidence.get('spend') or item.get('spend') or 0.0))
    return cleaned


def _ad_materialize_fact_rows(rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    buckets: Dict[Tuple[str, ...], Dict[str, Any]] = {}
    prepared_rows = [_ad_prepare_country_dimension(row) for row in (rows or []) if isinstance(row, dict)]
    prepared_rows = _ad_enrich_countries_from_meta_delivery(prepared_rows)
    for row in _ad_enrich_unknown_countries(prepared_rows):
        if not isinstance(row, dict):
            continue
        key = _ad_fact_grain_key(row)
        bucket = buckets.setdefault(key, {
            'date': key[0],
            'data_source': key[1],
            'platform': key[2],
            'app_id': key[3],
            'appsflyer_app_id': key[4],
            'target_app': _ad_dashboard_row_target_app(row),
            'account_id': row.get('account_id') or row.get('ad_account_id') or '',
            'account_name': row.get('account_name') or row.get('app_id') or '',
            'ad_account_id': row.get('ad_account_id') or row.get('account_id') or '',
            'external_app': row.get('external_app') or '',
            'country': key[5],
            'media_source': key[6],
            'campaign': key[7],
            'ad_group': key[8],
            'ad': key[9],
            'source_type': key[10],
            **_empty_ad_metrics(),
            'row_count': 0,
        })
        row_target_app = _ad_dashboard_row_target_app(row)
        if bucket.get('target_app') not in {'linky', 'timo'} and row_target_app in {'linky', 'timo'}:
            bucket['target_app'] = row_target_app
        for source_key, target_key in [
            ('account_id', 'account_id'),
            ('account_name', 'account_name'),
            ('ad_account_id', 'ad_account_id'),
            ('external_app', 'external_app'),
        ]:
            value = row.get(source_key)
            if not value and source_key == 'account_id':
                value = row.get('ad_account_id')
            if not value and source_key == 'ad_account_id':
                value = row.get('account_id')
            if value and not bucket.get(target_key):
                bucket[target_key] = value
        for metadata_key in (
            'historical_recovery',
            'historical_source_created_at',
            'country_attribution_status',
            'country_attribution_source',
            'country_attribution_grain',
            'campaign_id',
            'adset_id',
            'ad_id',
            'ad_name',
        ):
            value = row.get(metadata_key)
            if value in (None, ''):
                continue
            if metadata_key not in bucket:
                bucket[metadata_key] = value
            elif metadata_key in {'campaign_id', 'adset_id', 'ad_id'} and str(bucket[metadata_key]) != str(value):
                bucket[metadata_key] = ''
        bucket['row_count'] = int(bucket.get('row_count') or 0) + 1
        _add_ad_metrics(bucket, row)
    materialized = list(buckets.values())
    materialized.sort(key=lambda item: (
        str(item.get('date') or ''),
        str(item.get('data_source') or ''),
        str(item.get('platform') or ''),
        str(item.get('app_id') or ''),
        str(item.get('country') or ''),
        str(item.get('campaign') or ''),
        str(item.get('ad_group') or ''),
        str(item.get('ad') or ''),
    ))
    return materialized


def _aggregate_ad_metrics(rows: List[Dict[str, Any]]) -> Dict[str, Any]:
    metrics = _empty_ad_metrics()
    visible_rows = [row for row in rows if isinstance(row, dict)]
    platforms_with_media = {
        str(row.get('platform') or '').strip().lower()
        for row in visible_rows
        if str(row.get('data_source') or '').strip().lower() in {'meta', 'google', 'tiktok'}
    }
    platforms_with_af = {
        str(row.get('platform') or '').strip().lower()
        for row in visible_rows
        if str(row.get('data_source') or '').strip().lower() == 'appsflyer'
    }
    media_keys = ['cost', 'meta_installs', 'meta_registrations', 'meta_guild_joins', 'clicks', 'link_clicks', 'impressions', 'reach']
    af_keys = ['af_installs', 'af_registrations', 'registrations', 'af_guild_joins']
    af_media_fallback_keys = ['cost', 'installs', 'clicks', 'link_clicks', 'impressions', 'reach']
    tugao_funnel_keys = [
        'onsite_registrations', 'high_value_users', 'im_entries', 'im_first_replies',
        'auto_apply_message_users', 'im_step2_triggers', 'im_manual_reply_3',
        'im_user_message_ge_5_users', 'im_link_clicks', 'im_link_click_users',
        'link_click_users', 'linky_register_users', 'bind_success_users',
        'crm_succeed_users', 'high_intent_im_users',
    ]
    true_join_keys = ['guild_joins', 'promotion_guild_joins', 'organic_guild_joins']
    for row in visible_rows:
        source = str(row.get('data_source') or '').strip().lower()
        platform = str(row.get('platform') or '').strip().lower()
        if source == 'appsflyer':
            _add_ad_metrics(metrics, row, af_keys)
            if platform not in platforms_with_media:
                _add_ad_metrics(metrics, row, af_media_fallback_keys)
        elif source == 'bindsuccess':
            _add_ad_metrics(metrics, row, true_join_keys)
        elif source in {'marketingdiagnostics', 'marketing_diagnostics'}:
            keys = list(tugao_funnel_keys + true_join_keys)
            if platform not in platforms_with_media:
                keys.extend(media_keys)
            if platform not in platforms_with_af:
                keys.extend(af_keys)
            _add_ad_metrics(metrics, row, keys)
        elif source in {'tugaofunnel', 'tugao_funnel', 'tugao_onsite_funnel'}:
            _add_ad_metrics(metrics, row, tugao_funnel_keys + true_join_keys)
        elif source in {'meta', 'google', 'tiktok'}:
            _add_ad_metrics(metrics, row, media_keys)
        else:
            _add_ad_metrics(metrics, row)
    return _finalize_ad_metrics(metrics)


def _ad_missing_account_label() -> str:
    return '未归属广告账户'


def _ad_account_label_or_missing(value: Any) -> str:
    label = str(value or '').strip()
    if not label:
        return _ad_missing_account_label()
    if re.match(r'^[a-z][a-z0-9_]*(\.[a-z0-9_]+)+$', label, flags=re.IGNORECASE):
        return _ad_missing_account_label()
    return label


AD_DASHBOARD_TARGET_APP_LINKY_META_ACCOUNTS = {
    '1898261564216326',
    '1511281443796277',
    '1022472447112808',
    '2014618999169375',
    '865675816544216',
}
AD_DASHBOARD_TARGET_APP_TIMO_META_ACCOUNTS = {
    '1293506106236750',
    '1625526805175773',
    '1457588552349197',
    '1250000910496826',
}
AD_DASHBOARD_TARGET_APP_LINKY_ALIASES = {
    'linky',
}
AD_DASHBOARD_TARGET_APP_TIMO_ALIASES = {
    'com.timetrade.duitan',
    'duitan',
    'timo',
}
AD_DASHBOARD_TARGET_APP_LABELS = {
    'all': '全部',
    'linky': 'Linky',
    'timo': 'Timo',
    'inactive': '未启用',
}


def _normalize_ad_dashboard_target_app(value: Any) -> str:
    normalized = str(value or '').strip().lower()
    if normalized in {'', 'all', '全部', '__all__'}:
        return 'all'
    if normalized in {'linky', 'link'}:
        return 'linky'
    if normalized in {'timo'}:
        return 'timo'
    return ''


def _normalize_ad_account_id_candidate(value: Any) -> str:
    text = str(value or '').strip()
    if not text:
        return ''
    act_match = re.search(r'act[_\s-]*(\d{8,})', text, flags=re.IGNORECASE)
    if act_match:
        return act_match.group(1)
    digit_match = re.search(r'\b(\d{12,20})\b', text)
    return digit_match.group(1) if digit_match else ''


def _normalize_ad_app_alias_candidate(value: Any) -> str:
    text = str(value or '').strip().lower()
    if not text:
        return ''
    return re.sub(r'\s+', '', text.replace('act_', ''))


def _ad_dashboard_target_app_from_account(*, platform: Any, account_id: Any = '', account_label: Any = '', appsflyer_app_id: Any = '') -> str:
    normalized_appsflyer_app_id = _normalize_ad_app_alias_candidate(appsflyer_app_id)
    if normalized_appsflyer_app_id in AD_DASHBOARD_TARGET_APP_LINKY_ALIASES:
        return 'linky'
    if normalized_appsflyer_app_id in AD_DASHBOARD_TARGET_APP_TIMO_ALIASES:
        return 'timo'
    if str(platform or '').strip().lower() != 'meta':
        return 'inactive'
    normalized_account_id = _normalize_ad_account_id_candidate(account_id) or _normalize_ad_account_id_candidate(account_label)
    if normalized_account_id in AD_DASHBOARD_TARGET_APP_LINKY_META_ACCOUNTS:
        return 'linky'
    if normalized_account_id in AD_DASHBOARD_TARGET_APP_TIMO_META_ACCOUNTS:
        return 'timo'
    account_text = ' '.join([str(account_id or ''), str(account_label or '')]).strip().lower()
    if re.search(r'(^|[\s_-])tm($|[\s_-])', account_text):
        return 'timo'
    if re.search(r'(^|[\s_-])lk($|[\s_-])', account_text):
        return 'linky'
    return 'inactive'


def _ad_dashboard_row_target_app(row: Dict[str, Any]) -> str:
    explicit = str((row or {}).get('target_app') or '').strip().lower()
    data_source = str((row or {}).get('data_source') or '').strip().lower()
    external_target_app = _normalize_ad_dashboard_target_app((row or {}).get('external_app'))
    is_tugao_natural = (
        data_source in {'tugaofunnel', 'tugao_funnel', 'tugao_onsite_funnel'}
        and str((row or {}).get('platform') or '').strip().lower() == 'internal'
    )
    if is_tugao_natural and external_target_app in {'linky', 'timo'}:
        return external_target_app
    inferred = _ad_dashboard_target_app_from_account(
        platform=(row or {}).get('platform'),
        account_id=(row or {}).get('account_id') or (row or {}).get('ad_account_id'),
        account_label=(row or {}).get('app_id'),
        appsflyer_app_id=(row or {}).get('appsflyer_app_id'),
    )
    if inferred in {'linky', 'timo'}:
        return inferred
    return explicit if explicit in {'linky', 'timo', 'inactive'} else 'inactive'


def _normalize_ad_dashboard_match_value(value: Any) -> str:
    return re.sub(r'\s+', '', str(value or '').strip().lower())


def _ad_dashboard_match_key(row: Dict[str, Any], fields: Tuple[str, ...]) -> Tuple[str, ...]:
    return tuple(_normalize_ad_dashboard_match_value((row or {}).get(field)) for field in fields)


def _ad_dashboard_match_key_is_usable(key: Tuple[str, ...]) -> bool:
    invalid_values = {'未命名', 'none', 'null', 'unknown', '未知'}
    return bool(key) and all(part and part not in invalid_values for part in key)


def _ad_dashboard_build_target_app_matchers(rows: List[Dict[str, Any]]) -> List[Tuple[Tuple[str, ...], Dict[Tuple[str, ...], str]]]:
    levels: List[Tuple[str, ...]] = [
        ('country', 'campaign', 'ad_group', 'ad'),
        ('country', 'campaign', 'ad_group'),
        ('country', 'campaign'),
    ]
    buckets: Dict[Tuple[str, ...], Dict[Tuple[str, ...], set[str]]] = {level: {} for level in levels}
    for row in rows or []:
        if not isinstance(row, dict):
            continue
        target_app = _ad_dashboard_row_target_app(row)
        if target_app not in {'linky', 'timo'}:
            continue
        platform = str(row.get('platform') or '').strip().lower()
        if platform != 'meta':
            continue
        for level in levels:
            key = _ad_dashboard_match_key(row, level)
            if not _ad_dashboard_match_key_is_usable(key):
                continue
            buckets[level].setdefault(key, set()).add(target_app)
    matchers: List[Tuple[Tuple[str, ...], Dict[Tuple[str, ...], str]]] = []
    for level in levels:
        matchers.append((level, {key: next(iter(targets)) for key, targets in buckets[level].items() if len(targets) == 1}))
    return matchers


def _ad_dashboard_target_app_from_peer_rows(row: Dict[str, Any], matchers: List[Tuple[Tuple[str, ...], Dict[Tuple[str, ...], str]]]) -> str:
    if str((row or {}).get('platform') or '').strip().lower() != 'meta':
        return ''
    for fields, mapping in matchers:
        key = _ad_dashboard_match_key(row, fields)
        if not _ad_dashboard_match_key_is_usable(key):
            continue
        target_app = mapping.get(key)
        if target_app in {'linky', 'timo'}:
            return target_app
    return ''


def _ad_dashboard_target_app_options() -> List[Dict[str, str]]:
    return [
        {'value': 'all', 'label': AD_DASHBOARD_TARGET_APP_LABELS['all']},
        {'value': 'linky', 'label': AD_DASHBOARD_TARGET_APP_LABELS['linky']},
        {'value': 'timo', 'label': AD_DASHBOARD_TARGET_APP_LABELS['timo']},
    ]


def _empty_im_diagnostics_summary_payload(*, target_app: str = 'all', region: str = '') -> Dict[str, Any]:
    normalized_region = str(region or '').strip() or 'all'
    return {
        'ok': True,
        'target_app': _normalize_ad_dashboard_target_app(target_app) or 'all',
        'diagnosis_run_id': '',
        'region': normalized_region,
        'region_label': normalized_region if normalized_region != 'all' else '全部',
        'taxonomy_version': 'im_diagnosis_taxonomy_v1',
        'summary': {
            'sample_conversations': 0,
            'successful_conversations': 0,
            'lost_conversations': 0,
            'join_rate': 0.0,
        },
        'segments': [],
        'chain_steps': [],
        'result_message_facts': {
            'coverage_status': 'missing',
            'timezone': 'UTC+0',
            'metrics': {},
            'step_coverage': [],
        },
        'top_issues': [],
        'funnel_insights': [],
        'llm_tasks': {},
        'aggregates': [],
        'script_suggestions': [],
    }


def _ad_dashboard_apply_target_app_filter(rows: List[Dict[str, Any]], target_app: str) -> List[Dict[str, Any]]:
    normalized = _normalize_ad_dashboard_target_app(target_app)
    matchers = _ad_dashboard_build_target_app_matchers(rows or [])
    enriched = []
    for row in rows or []:
        if not isinstance(row, dict):
            continue
        row_target_app = _ad_dashboard_row_target_app(row)
        if row_target_app not in {'linky', 'timo'}:
            row_target_app = _ad_dashboard_target_app_from_peer_rows(row, matchers) or row_target_app
        enriched_row = row if row.get('target_app') == row_target_app else {**row, 'target_app': row_target_app}
        if normalized in {'', 'all'} or row_target_app == normalized:
            enriched.append(enriched_row)
    return enriched


def _ad_is_placeholder_account(value: Any) -> bool:
    label = str(value or '').strip()
    return not label or label == _ad_missing_account_label() or bool(re.match(r'^[a-z][a-z0-9_]*(\.[a-z0-9_]+)+$', label, flags=re.IGNORECASE))


def _ad_dashboard_build_account_matchers(rows: List[Dict[str, Any]]) -> List[Tuple[Tuple[str, ...], Dict[Tuple[str, ...], str]]]:
    levels: List[Tuple[str, ...]] = [
        ('date', 'platform', 'country', 'campaign', 'ad_group', 'ad'),
        ('date', 'platform', 'country', 'campaign', 'ad_group'),
        ('date', 'platform', 'country', 'campaign'),
        ('date', 'platform', 'campaign'),
    ]
    buckets: Dict[Tuple[str, ...], Dict[Tuple[str, ...], set[str]]] = {level: {} for level in levels}
    for row in rows or []:
        if not isinstance(row, dict):
            continue
        source = str(row.get('data_source') or '').strip().lower()
        if source not in {'meta', 'google', 'tiktok'}:
            continue
        account = str(row.get('app_id') or '').strip()
        if _ad_is_placeholder_account(account):
            continue
        for level in levels:
            key = _ad_dashboard_match_key(row, level)
            if not _ad_dashboard_match_key_is_usable(key):
                continue
            buckets[level].setdefault(key, set()).add(account)
    matchers: List[Tuple[Tuple[str, ...], Dict[Tuple[str, ...], str]]] = []
    for level in levels:
        matchers.append((level, {key: next(iter(accounts)) for key, accounts in buckets[level].items() if len(accounts) == 1}))
    return matchers


def _ad_dashboard_account_from_peer_rows(row: Dict[str, Any], matchers: List[Tuple[Tuple[str, ...], Dict[Tuple[str, ...], str]]]) -> str:
    if not _ad_is_placeholder_account((row or {}).get('app_id')):
        return ''
    source = str((row or {}).get('data_source') or '').strip().lower()
    if source in {'meta', 'google', 'tiktok'}:
        return ''
    for fields, mapping in matchers:
        key = _ad_dashboard_match_key(row, fields)
        if not _ad_dashboard_match_key_is_usable(key):
            continue
        account = mapping.get(key)
        if account and not _ad_is_placeholder_account(account):
            return account
    return ''


def _ad_dashboard_enrich_account_from_peer_rows(rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    matchers = _ad_dashboard_build_account_matchers(rows or [])
    enriched: List[Dict[str, Any]] = []
    for row in rows or []:
        if not isinstance(row, dict):
            continue
        peer_account = _ad_dashboard_account_from_peer_rows(row, matchers)
        if peer_account:
            enriched.append({**row, 'app_id': peer_account})
        else:
            enriched.append(row)
    return enriched


def _ad_period_summary(rows: List[Dict[str, Any]], start_date: datetime.date, end_date: datetime.date) -> Dict[str, Any]:
    visible_rows = []
    for row in rows:
        row_date = _parse_dashboard_date(row.get('date'))
        if row_date and start_date <= row_date <= end_date:
            visible_rows.append(row)
    return _aggregate_ad_metrics(visible_rows)


def _ad_period_overview_attribution_summary(platform: str, metrics: Dict[str, Any]) -> Dict[str, Any]:
    overview = dict(metrics or {})
    platform_label = str(platform or '').strip().lower()
    if platform_label == 'internal':
        overview['installs'] = round(float(overview.get('af_installs') or 0.0), 4)
        overview['cpi'] = round(float(overview.get('cost') or 0.0) / float(overview.get('installs') or 0.0), 4) if overview.get('installs') else 0.0
        return overview
    cost = float(overview.get('cost') or 0.0)
    af_installs = float(overview.get('af_installs') or 0.0)
    guild_joins = float(overview.get('guild_joins') or 0.0)
    overview['installs'] = round(af_installs, 4)
    overview['cpi'] = round(cost / af_installs, 4) if af_installs else 0.0
    overview['guild_joins'] = round(guild_joins, 4)
    overview['join_cost'] = round(cost / guild_joins, 4) if guild_joins else 0.0
    return overview


def _ad_period_country_breakdown(platform: str, rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    buckets: Dict[str, List[Dict[str, Any]]] = {}
    for row in rows or []:
        prepared = _ad_prepare_country_dimension(row)
        country = (
            AD_HISTORICAL_SETTLEMENT_UNALLOCATED_COUNTRY
            if prepared.get('country_attribution_status') == 'historical_settlement_unallocated'
            else normalize_country_label(prepared.get('country')) or 'Unknown'
        )
        buckets.setdefault(country or 'Unknown', []).append(row)
    breakdown: List[Dict[str, Any]] = []
    for country, country_rows in buckets.items():
        metrics = _ad_period_overview_attribution_summary(platform, _aggregate_ad_metrics(country_rows))
        breakdown.append({'country': country, **metrics})
    breakdown.sort(
        key=lambda item: (
            str(item.get('country') or '') == 'Unknown',
            -float(item.get('cost') or 0.0),
            -float(item.get('guild_joins') or 0.0),
            str(item.get('country') or ''),
        )
    )
    return breakdown


def _ad_rows_in_date_range(rows: List[Dict[str, Any]], start_date: datetime.date, end_date: datetime.date) -> List[Dict[str, Any]]:
    visible_rows: List[Dict[str, Any]] = []
    for row in rows:
        row_date = _parse_dashboard_date((row or {}).get('date'))
        if row_date and start_date <= row_date <= end_date:
            visible_rows.append(row)
    return visible_rows


def _ad_top_breakdown(rows: List[Dict[str, Any]], key: str, limit: int = 8) -> List[Dict[str, Any]]:
    buckets: Dict[str, List[Dict[str, Any]]] = {}
    for row in rows:
        label = str(row.get(key) or '').strip() or '未命名'
        buckets.setdefault(label, []).append(row)
    ranked = []
    for label, bucket_rows in buckets.items():
        ranked.append({'name': label, **_aggregate_ad_metrics(bucket_rows)})
    ranked.sort(key=lambda item: (float(item.get('cost') or 0.0), float(item.get('installs') or 0.0)), reverse=True)
    return ranked[:limit]


def _ad_comparison_series(rows: List[Dict[str, Any]], key: str, start_date: datetime.date, window_days: int, limit: int = 8) -> List[Dict[str, Any]]:
    buckets: Dict[str, List[Dict[str, Any]]] = {}
    for row in rows:
        label = str(row.get(key) or '').strip()
        if not label:
            continue
        buckets.setdefault(label, []).append(row)
    series: List[Dict[str, Any]] = []
    for label, bucket_rows in buckets.items():
        daily: List[Dict[str, Any]] = []
        for offset in range(window_days):
            day = start_date + timedelta(days=offset)
            daily.append({'date': day.isoformat(), **_ad_period_summary(bucket_rows, day, day)})
        metrics = _ad_period_summary(bucket_rows, start_date, start_date + timedelta(days=window_days - 1))
        series.append({
            'name': label,
            'metrics': metrics,
            'daily': daily,
        })
    series.sort(key=lambda item: (float((item.get('metrics') or {}).get('cost') or 0.0), float((item.get('metrics') or {}).get('installs') or 0.0)), reverse=True)
    return series[:limit]


AD_DASHBOARD_GLOBAL_FILTER_KEYS = ('data_source', 'app_id', 'media_source')
AD_DASHBOARD_FILTER_KEYS = AD_DASHBOARD_GLOBAL_FILTER_KEYS
AD_DASHBOARD_PLATFORM_NAMES = ('Meta', 'Google', 'TikTok', 'Internal')
AD_DASHBOARD_OVERVIEW_PLATFORM_NAMES = ('Meta', 'Google', 'TikTok', 'Internal')
AD_DASHBOARD_PLATFORM_FILTER_KEYS = ('target_app', 'app_id', 'country', 'campaign', 'ad_group', 'ad')
AD_DASHBOARD_PLATFORM_PARAM_PREFIXES = {
    'Meta': 'meta',
    'Google': 'google',
    'TikTok': 'tiktok',
    'Internal': 'natural',
}
AD_DASHBOARD_EVENT_MAPPINGS = [
    {
        'event': 'af_complete_registration',
        'label': '完成注册',
        'stage': '注册',
        'trigger': '用户注册成功后触发',
        'send_to_media': True,
        'media_events': {
            'Meta': 'fb_mobile_complete_registration',
            'Google': 'custom sign_up',
            'TikTok': 'custom sign_up',
        },
    },
    {
        'event': 'l1_task_high_value',
        'label': 'L1 高价值用户完成',
        'stage': '高价值',
        'trigger': '用户完成性别年龄选择且为女性 18-40 岁时触发',
        'send_to_media': True,
        'media_events': {
            'Meta': 'fb_mobile_tutorial_completion',
            'Google': 'tutorial_complete',
            'TikTok': 'CompleteTutorial',
        },
    },
    {
        'event': 'l1_task_low_value',
        'label': 'L1 普通价值用户完成',
        'stage': '分流',
        'trigger': '用户完成性别年龄选择但不属于高价值人群时触发',
        'send_to_media': False,
        'media_events': {
            'Meta': '--',
            'Google': '--',
            'TikTok': '--',
        },
    },
    {
        'event': 'l1_task_minor',
        'label': 'L1 未成年分支',
        'stage': '分流',
        'trigger': '用户完成性别年龄选择且年龄小于 18 岁时触发',
        'send_to_media': False,
        'media_events': {
            'Meta': '--',
            'Google': '--',
            'TikTok': '--',
        },
    },
    {
        'event': 'im_user_first_manual_reply',
        'label': '用户首次真人回复',
        'stage': 'IM',
        'trigger': '用户在 IM 中第一次主动发送真人消息',
        'send_to_media': True,
        'media_events': {
            'Meta': 'CUSTOM',
            'Google': 'CUSTOM',
            'TikTok': 'CUSTOM',
        },
    },
    {
        'event': 'im_apply_step2_triggered',
        'label': '报名第二步触发',
        'stage': 'IM',
        'trigger': '用户完成 R1 指定互动动作后，系统成功发送 R2 文案',
        'send_to_media': True,
        'media_events': {
            'Meta': 'CUSTOM',
            'Google': 'CUSTOM',
            'TikTok': 'CUSTOM',
        },
    },
    {
        'event': 'im_user_manual_reply_count_3',
        'label': '用户真人消息达到 3 条',
        'stage': 'IM',
        'trigger': '用户在同一 IM 会话中累计发送真人消息达到 3 条',
        'send_to_media': True,
        'media_events': {
            'Meta': 'CUSTOM',
            'Google': 'CUSTOM',
            'TikTok': 'CUSTOM',
        },
    },
    {
        'event': 'im_link_click',
        'label': 'IM 链接点击',
        'stage': '点击',
        'trigger': '用户点击 IM 消息中的链接',
        'send_to_media': True,
        'media_events': {
            'Meta': 'contact',
            'Google': 'session_star',
            'TikTok': 'CUSTOM',
        },
    },
    {
        'event': 'join_guild',
        'label': '加入公会',
        'stage': '入会',
        'trigger': '用户完成加入公会流程后触发',
        'send_to_media': True,
        'media_events': {
            'Meta': 'SubmitApplication',
            'Google': 'custom->unlock_achievement',
            'TikTok': 'JoinGroup',
        },
    },
]


def _ad_is_organic_media_source(value: Any) -> bool:
    return str(value or '').strip().lower() in {'internal', 'organic', 'organic_search', 'none', 'natural', '自然量'}


def _ad_platform_label(media_source: Any, *hints: Any) -> str:
    raw = str(media_source or '').strip()
    if raw.lower() == 'internal':
        return 'Internal'
    joined = ' '.join(str(value or '') for value in (media_source, *hints) if str(value or '').strip())
    lowered = joined.lower()
    if not lowered:
        return '未命名'
    if str(media_source or '').strip().lower() in {'organic', 'organic_search', 'none', 'natural', '自然量'}:
        return 'Internal'
    if any(marker in lowered for marker in ('google', 'googleads', 'google ads', 'google_ads', 'googleadwords', 'googleadwords_int', 'adwords', 'uac')):
        return 'Google'
    if any(marker in lowered for marker in ('tiktok', 'tik tok', 'bytedance', 'byte dance', 'bytedanceglobal', 'bytedanceglobal_int', 'tt4b', 'tiktokglobal', 'tiktokglobal_int', 'tiktok_int', 'tiktok ads', 'tiktok_ads', 'tiktok for business', 'tiktokforbusiness', 'musically', 'pangle')):
        return 'TikTok'
    if re.search(r'(^|[^a-z0-9])tk([^a-z0-9]|$)', lowered):
        return 'TikTok'
    if any(marker in lowered for marker in ('facebook', 'facebook_int', 'facebook ads', 'meta', 'meta ads', 'fbad', 'restricted', 'instagram')):
        return 'Meta'
    if re.search(r'(^|[^a-z0-9])(fb|ig)([^a-z0-9]|$)', lowered):
        return 'Meta'
    return raw or '未命名'


def _ad_country_label(*values: Any) -> str:
    joined = ' '.join(str(value or '') for value in values if str(value or '').strip()).strip()
    lowered = joined.lower()
    if not lowered:
        return 'Unknown'
    if any(marker in lowered for marker in ('indonesia', 'indonesian', 'idn', ' id ', '-id', '_id', '印尼')):
        return 'Indonesia'
    if any(marker in lowered for marker in ('brazil', 'brasil', ' br ', '-br', '_br', '巴西')):
        return 'Brazil'
    if any(marker in lowered for marker in ('mexico', ' mx ', '-mx', '_mx', '墨西哥')):
        return 'Mexico'
    if any(marker in lowered for marker in ('venezuela', ' ve ', '-ve', '_ve', '委内瑞拉')):
        return 'Venezuela'
    if any(marker in lowered for marker in ('chile', ' cl ', '-cl', '_cl', '智利')):
        return 'Chile'
    if any(marker in lowered for marker in ('colombia', ' co ', '-co', '_co', '哥伦比亚')):
        return 'Colombia'
    raw = str(values[0] or '').strip() if values else ''
    return normalize_country_label(raw) or 'Unknown'


def _normalize_ad_filter_values(values: Any) -> List[str]:
    if values is None:
        return []
    if isinstance(values, str):
        raw_items = values.split(',')
    else:
        raw_items = []
        for value in list(values or []):
            raw_items.extend(str(value or '').split(','))
    normalized: List[str] = []
    seen = set()
    for item in raw_items:
        text = str(item or '').strip()
        lowered = text.lower()
        if not text or lowered in {'all', '全部', '__all__'} or lowered in seen:
            continue
        seen.add(lowered)
        normalized.append(text)
    return normalized


def _normalize_ad_filters(filters: Optional[Dict[str, Any]]) -> Dict[str, List[str]]:
    normalized: Dict[str, List[str]] = {}
    for key in AD_DASHBOARD_FILTER_KEYS:
        values = _normalize_ad_filter_values((filters or {}).get(key))
        if values:
            normalized[key] = values
    return normalized


def _row_matches_ad_filters(row: Dict[str, Any], filters: Dict[str, List[str]]) -> bool:
    for key, values in (filters or {}).items():
        if not values:
            continue
        wanted = {str(value or '').strip().lower() for value in values if str(value or '').strip()}
        current = str((row or {}).get(key) or '').strip().lower()
        if current not in wanted:
            return False
    return True


def _filter_ad_rows(rows: List[Dict[str, Any]], filters: Dict[str, List[str]]) -> List[Dict[str, Any]]:
    if not filters:
        return list(rows)
    return [row for row in rows if _row_matches_ad_filters(row, filters)]


def _ad_filter_options(rows: List[Dict[str, Any]], limit: int = 200, keys: Tuple[str, ...] = AD_DASHBOARD_FILTER_KEYS) -> Dict[str, List[Dict[str, Any]]]:
    options: Dict[str, List[Dict[str, Any]]] = {}
    for key in keys:
        buckets: Dict[str, Dict[str, Any]] = {}
        for row in rows:
            label = str((row or {}).get(key) or '').strip()
            if not label:
                continue
            if key == 'app_id' and _ad_is_placeholder_account(label):
                continue
            bucket = buckets.setdefault(label, {'name': label, 'row_count': 0, 'metrics': _empty_ad_metrics()})
            bucket['row_count'] = int(bucket.get('row_count') or 0) + 1
            _add_ad_metrics(bucket['metrics'], row)
        ranked = []
        for bucket in buckets.values():
            ranked.append({
                'name': bucket['name'],
                'row_count': bucket['row_count'],
                'metrics': _finalize_ad_metrics(bucket['metrics']),
            })
        ranked.sort(
            key=lambda item: (
                float((item.get('metrics') or {}).get('cost') or 0.0),
                float((item.get('metrics') or {}).get('installs') or 0.0),
                int(item.get('row_count') or 0),
            ),
            reverse=True,
        )
        options[key] = ranked[:limit]
    return options


def _ad_platform_cards(rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    cards: List[Dict[str, Any]] = []
    for platform in AD_DASHBOARD_PLATFORM_NAMES:
        platform_rows = [row for row in rows if str((row or {}).get('platform') or '').strip().lower() == platform.lower()]
        metrics = _ad_period_summary(platform_rows, datetime.min.date(), datetime.max.date())
        campaigns = {str((row or {}).get('campaign') or '').strip() for row in platform_rows if str((row or {}).get('campaign') or '').strip()}
        ad_groups = {str((row or {}).get('ad_group') or '').strip() for row in platform_rows if str((row or {}).get('ad_group') or '').strip()}
        ads = {str((row or {}).get('ad') or '').strip() for row in platform_rows if str((row or {}).get('ad') or '').strip()}
        cards.append({
            'name': platform,
            'row_count': len(platform_rows),
            'campaign_count': len(campaigns),
            'ad_group_count': len(ad_groups),
            'ad_count': len(ads),
            **metrics,
        })
    return cards


def _normalize_ad_platform_filters(platform_filters: Optional[Dict[str, Any]]) -> Dict[str, Dict[str, List[str]]]:
    normalized: Dict[str, Dict[str, List[str]]] = {}
    for platform in AD_DASHBOARD_PLATFORM_NAMES:
        raw = (platform_filters or {}).get(platform) or {}
        platform_values = {
            key: _normalize_ad_filter_values(raw.get(key))
            for key in AD_DASHBOARD_PLATFORM_FILTER_KEYS
        }
        platform_values = {key: values for key, values in platform_values.items() if values}
        if platform_values:
            normalized[platform] = platform_values
    return normalized


def _row_matches_ad_platform_filters(row: Dict[str, Any], platform_filters: Dict[str, Dict[str, List[str]]]) -> bool:
    platform = str((row or {}).get('platform') or '').strip()
    filters = platform_filters.get(platform) or {}
    if not filters:
        return False
    for key, values in filters.items():
        wanted = {str(value or '').strip().lower() for value in values if str(value or '').strip()}
        if key == 'target_app':
            current = _ad_dashboard_row_target_app(row)
        else:
            current = str((row or {}).get(key) or '').strip().lower()
        if current not in wanted:
            return False
    return True


def _filter_ad_rows_by_platform_filters(rows: List[Dict[str, Any]], platform_filters: Dict[str, Dict[str, List[str]]]) -> List[Dict[str, Any]]:
    if not platform_filters:
        return list(rows)
    return [row for row in rows if _row_matches_ad_platform_filters(row, platform_filters)]


def _normalize_ad_platform_date_windows(
    platform_date_windows: Optional[Dict[str, Any]],
    *,
    fallback_start_date: datetime.date,
    fallback_end_date: datetime.date,
    today: datetime.date,
) -> Tuple[Dict[str, Dict[str, str]], datetime.date, datetime.date, int, List[str]]:
    windows: Dict[str, Dict[str, str]] = {}
    starts: List[datetime.date] = []
    ends: List[datetime.date] = []
    warnings: List[str] = []
    fallback_days = max((fallback_end_date - fallback_start_date).days + 1, 1)
    for platform in AD_DASHBOARD_PLATFORM_NAMES:
        raw = (platform_date_windows or {}).get(platform) or {}
        platform_start, platform_end, _days, platform_warnings = _coerce_ad_dashboard_window(
            days=fallback_days,
            date_from=raw.get('date_from') or fallback_start_date.isoformat(),
            date_to=raw.get('date_to') or fallback_end_date.isoformat(),
            today=today,
        )
        starts.append(platform_start)
        ends.append(platform_end)
        warnings.extend(platform_warnings)
        windows[platform] = {
            'date_from': platform_start.isoformat(),
            'date_to': platform_end.isoformat(),
        }
    fetch_start = min(starts) if starts else fallback_start_date
    fetch_end = max(ends) if ends else fallback_end_date
    fetch_days = max((fetch_end - fetch_start).days + 1, 1)
    return windows, fetch_start, fetch_end, fetch_days, list(dict.fromkeys(warnings))


def _row_matches_ad_platform_date_window(row: Dict[str, Any], platform_date_windows: Dict[str, Dict[str, str]]) -> bool:
    platform = str((row or {}).get('platform') or '').strip()
    window = platform_date_windows.get(platform) or {}
    row_date = _parse_dashboard_date((row or {}).get('date'))
    start_date = _parse_dashboard_date(window.get('date_from'))
    end_date = _parse_dashboard_date(window.get('date_to'))
    if not row_date or not start_date or not end_date:
        return True
    return start_date <= row_date <= end_date


def _filter_ad_rows_by_platform_date_windows(rows: List[Dict[str, Any]], platform_date_windows: Dict[str, Dict[str, str]]) -> List[Dict[str, Any]]:
    if not platform_date_windows:
        return list(rows)
    return [row for row in rows if _row_matches_ad_platform_date_window(row, platform_date_windows)]


def _ad_platform_filter_options(rows: List[Dict[str, Any]], limit: int = 200) -> Dict[str, Dict[str, List[Dict[str, Any]]]]:
    options: Dict[str, Dict[str, List[Dict[str, Any]]]] = {}
    for platform in AD_DASHBOARD_PLATFORM_NAMES:
        platform_rows = [row for row in rows if str((row or {}).get('platform') or '').strip().lower() == platform.lower()]
        options[platform] = _ad_filter_options(platform_rows, limit=limit, keys=AD_DASHBOARD_PLATFORM_FILTER_KEYS)
    return options


def _ad_funnel_steps(metrics: Dict[str, Any]) -> List[Dict[str, Any]]:
    onsite_value = metrics.get('onsite_registrations') or 0
    return [
        {'key': 'onsite_registrations', 'label': '站内注册', 'value': onsite_value, 'rate': 1.0 if float(onsite_value or 0) else 0},
        {'key': 'high_value_users', 'label': '高价值', 'value': metrics.get('high_value_users') or 0, 'rate': metrics.get('high_value_rate') or 0},
        {'key': 'im_entries', 'label': '自动报名人数', 'value': metrics.get('im_entries') or 0, 'rate': metrics.get('im_entry_rate') or 0},
        {'key': 'im_manual_reply_3', 'label': 'IM>=3', 'value': metrics.get('im_manual_reply_3') or 0, 'rate': metrics.get('im_reply_3_rate') or 0},
        {'key': 'guild_joins', 'label': '入会', 'value': metrics.get('guild_joins') or 0, 'rate': metrics.get('join_rate') or 0},
    ]


def _ad_platform_sections(rows: List[Dict[str, Any]], start_date: datetime.date, window_days: int, limit: int = 8) -> List[Dict[str, Any]]:
    sections: List[Dict[str, Any]] = []
    for platform in AD_DASHBOARD_PLATFORM_NAMES:
        platform_rows = [row for row in rows if str((row or {}).get('platform') or '').strip().lower() == platform.lower()]
        metrics = _ad_period_summary(platform_rows, start_date, start_date + timedelta(days=max(window_days - 1, 0)))
        if platform == 'Internal':
            metrics['natural_guild_joins'] = round(float(
                metrics.get('organic_guild_joins')
                or metrics.get('guild_joins')
                or 0.0
            ), 4)
        sections.append({
            'name': platform,
            'row_count': len(platform_rows),
            'metrics': metrics,
            'funnel': _ad_funnel_steps(metrics),
            'top_campaigns': _ad_top_breakdown(platform_rows, 'campaign', limit=limit),
            'top_ad_groups': _ad_top_breakdown(platform_rows, 'ad_group', limit=limit),
            'top_ads': _ad_top_breakdown(platform_rows, 'ad', limit=limit),
        })
    return sections


def _ad_reconciliation_rows(rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    reconciliation: List[Dict[str, Any]] = []
    for platform in AD_DASHBOARD_PLATFORM_NAMES:
        platform_rows = [row for row in rows if str((row or {}).get('platform') or '').strip().lower() == platform.lower()]
        af_rows = [row for row in platform_rows if str((row or {}).get('data_source') or '').strip().lower() == 'appsflyer']
        media_rows = [row for row in platform_rows if str((row or {}).get('data_source') or '').strip().lower() == platform.lower()]
        af_metrics = _ad_period_summary(af_rows, datetime.min.date(), datetime.max.date())
        media_metrics = _ad_period_summary(media_rows, datetime.min.date(), datetime.max.date())
        deltas: Dict[str, Dict[str, Any]] = {}
        for key in ('installs', 'registrations', 'high_value_users', 'im_manual_reply_3', 'guild_joins'):
            af_value = float(af_metrics.get(key) or 0.0)
            media_value = float(media_metrics.get(key) or 0.0)
            diff = media_value - af_value
            deltas[key] = {
                'af': round(af_value, 4),
                'media': round(media_value, 4),
                'diff': round(diff, 4),
                'diff_rate': round(diff / af_value, 4) if af_value else 0.0,
            }
        reconciliation.append({
            'platform': platform,
            'has_media_source': bool(media_rows),
            'af_row_count': len(af_rows),
            'media_row_count': len(media_rows),
            'af': af_metrics,
            'media': media_metrics,
            'deltas': deltas,
        })
    return reconciliation


def _ad_row_has_paid_detail_signal(row: Dict[str, Any]) -> bool:
    return float((row or {}).get('cost') or 0.0) > 0.0


def _ad_row_has_true_join_detail_signal(row: Dict[str, Any]) -> bool:
    for key in ('guild_joins', 'promotion_guild_joins', 'organic_guild_joins'):
        if float((row or {}).get(key) or 0.0) > 0.0:
            return True
    return False


def _ad_detail_rows(rows: List[Dict[str, Any]], limit: int = 500) -> List[Dict[str, Any]]:
    def _natural_detail_key(row: Dict[str, Any]) -> Tuple[str, str, str, str, str]:
        return (
            str((row or {}).get('platform') or '').strip() or 'Unknown',
            normalize_country_label((row or {}).get('country')) or 'Unknown',
            str((row or {}).get('campaign') or '').strip() or '未命名',
            str((row or {}).get('ad_group') or '').strip() or '未命名',
            str((row or {}).get('ad') or '').strip() or '未命名',
        )

    buckets: Dict[Tuple[str, str, str, str, str, str, str], Dict[str, Any]] = {}
    paid_bucket_keys_by_natural: Dict[Tuple[str, str, str, str, str], List[Tuple[str, str, str, str, str, str, str]]] = defaultdict(list)
    for row in rows:
        if not _ad_row_has_paid_detail_signal(row):
            continue
        natural_key = _natural_detail_key(row)
        key = (
            natural_key[0],
            _ad_account_label_or_missing((row or {}).get('app_id')),
            *natural_key[1:],
            str((row or {}).get('source_type') or '').strip() or '推广量',
        )
        bucket = buckets.setdefault(key, {
            'platform': key[0],
            'account': key[1],
            'account_id': _normalize_ad_account_id_candidate((row or {}).get('account_id')),
            'country': key[2],
            'campaign': key[3],
            'ad_group': key[4],
            'ad': key[5],
            'source_type': key[6],
            'rows': [],
        })
        bucket['rows'].append(row)
        if key not in paid_bucket_keys_by_natural[natural_key]:
            paid_bucket_keys_by_natural[natural_key].append(key)
    for row in rows:
        if _ad_row_has_paid_detail_signal(row):
            continue
        natural_key = _natural_detail_key(row)
        candidate_keys = paid_bucket_keys_by_natural.get(natural_key) or []
        if not candidate_keys:
            if not _ad_row_has_true_join_detail_signal(row):
                continue
            row_account = _ad_account_label_or_missing((row or {}).get('app_id'))
            key = (
                natural_key[0],
                row_account,
                *natural_key[1:],
                str((row or {}).get('source_type') or '').strip() or '推广量',
            )
            bucket = buckets.setdefault(key, {
                'platform': key[0],
                'account': key[1],
                'account_id': _normalize_ad_account_id_candidate((row or {}).get('account_id')),
                'country': key[2],
                'campaign': key[3],
                'ad_group': key[4],
                'ad': key[5],
                'source_type': key[6],
                'rows': [],
            })
            bucket['rows'].append(row)
            continue
        row_account = _ad_account_label_or_missing((row or {}).get('app_id'))
        matched_key = next((key for key in candidate_keys if key[1] == row_account), None)
        if matched_key is None and (_ad_is_placeholder_account(row_account) or len(candidate_keys) == 1):
            matched_key = candidate_keys[0]
        if matched_key is not None:
            buckets[matched_key]['rows'].append(row)
    detail_rows: List[Dict[str, Any]] = []
    for bucket in buckets.values():
        metrics = _aggregate_ad_metrics(bucket.get('rows') or [])
        detail_rows.append({
            'platform': bucket['platform'],
            'account': bucket['account'],
            # Keep the human account label for the stable observation identity,
            # but expose the authoritative numeric Meta account separately.
            'account_id': bucket.get('account_id') or '',
            'account_identity': bucket['account'],
            'country': bucket['country'],
            'campaign': bucket['campaign'],
            'ad_group': bucket['ad_group'],
            'ad': bucket['ad'],
            'source_type': bucket['source_type'],
            **metrics,
        })
    def _detail_sort_text(value: Any) -> str:
        text = str(value or '').strip()
        if not text or text in {'未命名', 'Unknown'}:
            return '\uffff'
        return text.casefold()

    detail_rows.sort(
        key=lambda item: (
            _detail_sort_text(item.get('account')),
            _detail_sort_text(item.get('country')),
            _detail_sort_text(item.get('campaign')),
            _detail_sort_text(item.get('ad_group')),
            _detail_sort_text(item.get('ad')),
            _detail_sort_text(item.get('source_type')),
            -float(item.get('cost') or 0.0),
        ),
    )
    return detail_rows[:limit]


def _ad_platform_period_overview(
    rows: List[Dict[str, Any]],
    end_date: datetime.date,
    *,
    all_rows: Optional[List[Dict[str, Any]]] = None,
    target_app: str = 'all',
) -> Dict[str, List[Dict[str, Any]]]:
    periods = {
        'day': (end_date, end_date),
        'week': (end_date - timedelta(days=6), end_date),
        'month': (end_date - timedelta(days=29), end_date),
    }
    overview: Dict[str, List[Dict[str, Any]]] = {}
    for period_key, (start_date, period_end) in periods.items():
        period_rows = []
        for platform in AD_DASHBOARD_OVERVIEW_PLATFORM_NAMES:
            platform_rows = [
                row for row in rows
                if str((row or {}).get('platform') or '').strip().lower() == platform.lower()
            ]
            visible_rows = [
                row for row in platform_rows
                if (row_date := _parse_dashboard_date((row or {}).get('date'))) and start_date <= row_date <= period_end
            ]
            source_counts: Dict[str, int] = {}
            for row in visible_rows:
                source_label = str((row or {}).get('data_source') or '').strip() or 'Unknown'
                source_counts[source_label] = source_counts.get(source_label, 0) + 1
            period_metrics = _ad_period_summary(platform_rows, start_date, period_end)
            row_payload = {
                'platform': platform,
                'row_count': len(visible_rows),
                'source_summary': ' / '.join(f'{name} {count}' for name, count in sorted(source_counts.items())) or '-',
                'country_breakdown': _ad_period_country_breakdown(platform, visible_rows),
                **_ad_period_overview_attribution_summary(platform, period_metrics),
            }
            period_rows.append(row_payload)
        overview[period_key] = period_rows
    return overview


def _ad_dashboard_latest_complete_utc_date(now: Optional[datetime] = None) -> datetime.date:
    """Expose the latest published UTC data day only after the BJ 09:20 refresh."""
    current = now or datetime.now(timezone.utc)
    if current.tzinfo is None:
        current = current.replace(tzinfo=timezone.utc)
    current_bj = current.astimezone(ZoneInfo('Asia/Shanghai'))
    cutoff_bj = current_bj.replace(hour=9, minute=20, second=0, microsecond=0)
    days_back = 1 if current_bj >= cutoff_bj else 2
    return current_bj.date() - timedelta(days=days_back)


def _coerce_ad_dashboard_window(
    *,
    days: int,
    date_from: Optional[str],
    date_to: Optional[str],
    today: datetime.date,
    max_days: int = 120,
) -> Tuple[datetime.date, datetime.date, int, List[str]]:
    warnings: List[str] = []
    requested_days = min(max(int(days or 30), 1), max_days)
    parsed_to = _parse_dashboard_date(date_to) or today
    parsed_from = _parse_dashboard_date(date_from)
    if parsed_to > today:
        parsed_to = today
        warnings.append(f'广告看板只展示最新完整 UTC 日期，已截断到 {today.isoformat()}。')
    if parsed_from is not None and parsed_from > today:
        parsed_from = today
        if not any('最新完整 UTC 日期' in item for item in warnings):
            warnings.append(f'广告看板只展示最新完整 UTC 日期，已截断到 {today.isoformat()}。')
    if parsed_from is None:
        parsed_from = parsed_to
    if parsed_from > parsed_to:
        parsed_from, parsed_to = parsed_to, parsed_from
    actual_days = (parsed_to - parsed_from).days + 1
    if actual_days > max_days:
        parsed_from = parsed_to - timedelta(days=max_days - 1)
        actual_days = max_days
        warnings.append(f'为避免平台 API 查询过重，单次最多展示 {max_days} 天，已自动截取到最近 {max_days} 天。')
    return parsed_from, parsed_to, actual_days, warnings


def _ad_dashboard_cache_window(now_ts: Optional[float] = None, timezone_name: str = 'Asia/Shanghai') -> Tuple[float, float]:
    try:
        tz = ZoneInfo(str(timezone_name or 'Asia/Shanghai').strip() or 'Asia/Shanghai')
    except Exception:
        tz = ZoneInfo('Asia/Shanghai')
    now = datetime.fromtimestamp(float(now_ts or time.time()), tz)
    boundary = now.replace(hour=9, minute=20, second=0, microsecond=0)
    if now < boundary:
        boundary = boundary - timedelta(days=1)
    next_boundary = boundary + timedelta(days=1)
    return boundary.timestamp(), next_boundary.timestamp()


def _fetch_appsflyer_partner_by_date_rows(
    *,
    token: str,
    app_id: str,
    base_url: str,
    timezone_name: str,
    from_date: datetime.date,
    to_date: datetime.date,
    session: Any,
) -> List[Dict[str, Any]]:
    url = f'{str(base_url or "https://hq1.appsflyer.com").rstrip("/")}/api/agg-data/export/app/{quote(str(app_id), safe="")}/partners_by_date_report/v5'
    params = {'from': from_date.isoformat(), 'to': to_date.isoformat()}
    if timezone_name:
        params['timezone'] = timezone_name
    response = session.get(
        url,
        params=params,
        headers={'Authorization': f'Bearer {token}'},
        timeout=20.0,
    )
    if getattr(response, 'status_code', 200) == 403 and 'limit reached for partners-daily-report' in str(getattr(response, 'text', '') or '').lower():
        response = session.get(
            f'{str(base_url or "https://hq1.appsflyer.com").rstrip("/")}/api/agg-data/export/app/{quote(str(app_id), safe="")}/partners_report/v5',
            params=params,
            headers={'Authorization': f'Bearer {token}'},
            timeout=20.0,
        )
    response.raise_for_status()
    text = _dashboard_response_text(response)
    if not text.strip():
        return []
    parsed_rows: List[Dict[str, Any]] = []
    reader = csv.DictReader(io.StringIO(text))
    normalized_field_names = {_normalized_csv_key(name) for name in (reader.fieldnames or [])}
    has_date_field = any(
        _normalized_csv_key(candidate) in normalized_field_names
        for candidate in ['date', 'day', 'install date', 'install time']
    )
    is_multi_day_request = from_date != to_date
    for raw in reader:
        row_date = _parse_dashboard_date(_csv_lookup(raw, ['date', 'day', 'install date', 'install time']))
        if row_date is None and is_multi_day_request and not has_date_field:
            continue
        row_date = row_date or to_date
        if row_date is None:
            continue
        media_source = _csv_lookup(raw, [
            'media source',
            'media source (pid)',
            'media_source',
            'media_source_pid',
            'pid',
            'partner',
            'source',
            'network',
            'ad network',
        ])
        channel = _csv_lookup(raw, ['channel', 'af channel', 'af_channel', 'site id', 'site_id'])
        campaign = _csv_lookup(raw, ['campaign', 'campaign name', 'campaign (c)', 'campaign c', 'campaign_id', 'campaign id', 'c'])
        ad_group = _csv_lookup(raw, [
            'adset',
            'ad set',
            'ad set name',
            'adset name',
            'adset_id',
            'adset id',
            'adgroup',
            'ad group',
            'ad group name',
            'adgroup name',
            'adgroup_id',
            'adgroup id',
        ])
        ad_name = _csv_lookup(raw, [
            'ad',
            'ad name',
            'ad_name',
            'ad id',
            'ad_id',
            'creative',
            'creative name',
            'creative_id',
            'creative id',
        ])
        ad_account = _csv_lookup(raw, [
            'ad account',
            'ad account name',
            'account',
            'account name',
            'account_name',
            'af_ad_account',
            'af_ad_account_name',
            'advertiser',
            'advertiser name',
            'agency',
            '广告账户',
            '广告账户名称',
        ])
        normalized_media_source = media_source or channel or 'Organic'
        country = _csv_lookup(raw, ['country', 'country code', 'country_code', 'geo', 'region'])
        country_label = normalize_country_label(country) if country else 'Unknown'
        af_complete_registration = _dashboard_float(_csv_event_lookup(raw, 'af_complete_registration'))
        high_value_users = _dashboard_float(_csv_event_lookup(raw, 'l1_task_high_value'))
        im_link_clicks = _dashboard_float(_csv_event_lookup(raw, 'im_link_click'))
        im_first_replies = _dashboard_float(_csv_event_lookup(raw, 'im_user_first_manual_reply'))
        im_step2_triggers = _dashboard_float(_csv_event_lookup(raw, 'im_apply_step2_triggered'))
        im_manual_reply_3 = _dashboard_float(_csv_event_lookup(raw, 'im_user_manual_reply_count_3'))
        guild_joins = _dashboard_float(_csv_event_lookup(raw, 'join_guild'))
        installs = _dashboard_float(_csv_lookup(raw, ['installs', 'install', 'total installs', 'meta安装', 'af安装', 'app激活', '安装', '激活']))
        account_label = _ad_account_label_or_missing(ad_account)
        parsed_rows.append({
            'date': row_date.isoformat(),
            'data_source': 'AppsFlyer',
            'app_id': account_label,
            'appsflyer_app_id': app_id,
            'platform': _ad_platform_label(normalized_media_source, channel, campaign, ad_group, ad_name),
            'country': country_label,
            'country_attribution_status': (
                'appsflyer_report_country'
                if country
                else 'unresolved_waiting_meta_delivery_country'
            ),
            'country_attribution_source': (
                'appsflyer_partner_report'
                if country
                else 'meta_insights_country_breakdown_pending'
            ),
            'media_source': normalized_media_source,
            'campaign': campaign or '未命名',
            'ad_group': ad_group,
            'ad': ad_name,
            'source_type': '自然量' if _ad_is_organic_media_source(normalized_media_source) else '推广量',
            'cost': _dashboard_float(_csv_lookup(raw, ['cost', 'af cost', 'cost usd', 'total cost', '花费', '消耗', '总消耗'])),
            'installs': installs,
            'af_installs': installs,
            'registrations': af_complete_registration or _dashboard_float(_csv_lookup(raw, ['registrations', 'registration', 'af complete registration', 'complete registration', 'af registration', '注册', 'af注册', '站内注册'])),
            'af_registrations': af_complete_registration or _dashboard_float(_csv_lookup(raw, ['registrations', 'registration', 'af complete registration', 'complete registration', 'af registration', 'af注册'])),
            'onsite_registrations': _dashboard_float(_csv_lookup(raw, ['onsite registrations', 'onsite_registrations', 'site registrations', '站内注册', '注册'])),
            'high_value_users': high_value_users or _dashboard_float(_csv_lookup(raw, ['high value users', 'high_value_users', 'high value', 'l1 high value', '高价值', '高价值人数'])),
            'im_entries': _dashboard_float(_csv_lookup(raw, [
                'auto signup message sent users',
                'auto_signup_message_sent_users',
                'auto signup users',
                'auto_signup_users',
                'auto registration message sent users',
                '自动报名消息发送人数',
                '自动报名消息发送',
                '自动报名人数',
                'enter im',
                'entered im',
                'im entries',
                'im_entries',
                '进入im',
                '进入 im',
                '进入IM',
            ])) or im_link_clicks,
            'im_first_replies': im_first_replies or _dashboard_float(_csv_lookup(raw, ['first manual reply', 'first_manual_reply'])),
            'im_step2_triggers': im_step2_triggers or _dashboard_float(_csv_lookup(raw, ['apply step2', 'step2 triggered', 'r2 sent'])),
            'im_manual_reply_3': im_manual_reply_3 or _dashboard_float(_csv_lookup(raw, ['manual reply count 3', 'reply count 3', 'im>=3', 'im 3'])),
            'im_link_clicks': im_link_clicks or _dashboard_float(_csv_lookup(raw, ['im link click', 'link click'])),
            'guild_joins': 0.0,
            'af_guild_joins': guild_joins or _dashboard_float(_csv_lookup(raw, ['join guild', 'joingroup', 'join group', 'join_group', 'guild joins', 'joins', '当日入会全部新增', '当日入会', '入会', '入会人数', '加入公会', '加入公会人数'])),
            'purchases': _dashboard_float(_csv_lookup(raw, ['purchases', 'purchase', 'af purchase', 'paying users', 'payers'])),
            'revenue': _dashboard_float(_csv_lookup(raw, ['revenue', 'af revenue', 'event revenue', 'total revenue'])),
            'clicks': _dashboard_float(_csv_lookup(raw, ['clicks', 'click'])),
            'link_clicks': _dashboard_float(_csv_lookup(raw, ['link clicks', 'link_clicks', 'clicks', 'click'])),
            'impressions': _dashboard_float(_csv_lookup(raw, ['impressions', 'impression'])),
        })
    return parsed_rows


def _redact_dashboard_error_text(value: Any) -> str:
    text = str(value or '').strip()
    if not text:
        return ''
    text = re.sub(r'Bearer\s+[A-Za-z0-9._~+/=-]+', 'Bearer [redacted]', text, flags=re.IGNORECASE)
    text = re.sub(r'(access[_-]?token=)[^&\s]+', r'\1[redacted]', text, flags=re.IGNORECASE)
    text = re.sub(r'(/token/)[^/?#\s]+', r'\1[redacted]', text, flags=re.IGNORECASE)
    text = re.sub(r'(authorization=)[^&\s]+', r'\1[redacted]', text, flags=re.IGNORECASE)
    text = re.sub(r'https?://[^\s]+', '[redacted-url]', text)
    return text[:240]


def _dashboard_error_message(exc: Exception) -> str:
    response = getattr(exc, 'response', None)
    if response is not None:
        status_code = getattr(response, 'status_code', None)
        reason = _redact_dashboard_error_text(getattr(response, 'reason', '') or '')
        body = _redact_dashboard_error_text(getattr(response, 'text', '') or '')
        parts = [f'HTTP {status_code}' if status_code else 'HTTP error']
        if reason:
            parts.append(reason)
        if body:
            parts.append(body)
        return '；'.join(parts)
    return _redact_dashboard_error_text(exc)


def _fetch_appsflyer_in_app_event_rows(
    *,
    token: str,
    app_id: str,
    base_url: str,
    timezone_name: str,
    from_date: datetime.date,
    to_date: datetime.date,
    session: Any,
) -> List[Dict[str, Any]]:
    url = f'{str(base_url or "https://hq1.appsflyer.com").rstrip("/")}/api/raw-data/export/app/{quote(str(app_id), safe="")}/in_app_events_report/v5'
    params = {'from': from_date.isoformat(), 'to': to_date.isoformat()}
    if timezone_name:
        params['timezone'] = timezone_name
    response = session.get(
        url,
        params=params,
        headers={'Authorization': f'Bearer {token}'},
        timeout=30.0,
    )
    response.raise_for_status()
    text = _dashboard_response_text(response)
    if not text.strip():
        return []
    buckets: Dict[Tuple[str, str, str, str, str, str, str], Dict[str, Any]] = {}
    reader = csv.DictReader(io.StringIO(text))
    for raw in reader:
        event_name = _csv_lookup(raw, ['event name', 'event_name', 'event'])
        if not event_name:
            continue
        row_date = _parse_dashboard_date(_csv_lookup(raw, ['event time', 'event_time', 'date', 'day', 'install time']))
        if row_date is None:
            continue
        media_source = _csv_lookup(raw, [
            'media source',
            'media_source',
            'media source (pid)',
            'media_source_pid',
            'pid',
            'partner',
            'source',
            'network',
            'ad network',
        ])
        channel = _csv_lookup(raw, ['channel', 'af channel', 'af_channel', 'site id', 'site_id'])
        normalized_media_source = media_source or channel or 'Organic'
        campaign = _csv_lookup(raw, ['campaign', 'campaign name', 'campaign (c)', 'campaign c', 'campaign_id', 'campaign id', 'c', 'af_campaign']) or '未命名'
        ad_group = _csv_lookup(raw, [
            'adset',
            'ad set',
            'ad set name',
            'adset name',
            'adgroup',
            'ad group',
            'ad group name',
            'adgroup name',
            'adset id',
            'adgroup id',
            'af_adset',
            'af_adset_id',
        ])
        ad_name = _csv_lookup(raw, ['ad', 'ad name', 'ad_name', 'ad id', 'ad_id', 'creative', 'creative name', 'creative id', 'af_ad', 'af_ad_id'])
        ad_account = _csv_lookup(raw, [
            'ad account',
            'ad account name',
            'account',
            'account name',
            'account_name',
            'af_ad_account',
            'af_ad_account_name',
            'advertiser',
            'advertiser name',
            'agency',
            '广告账户',
            '广告账户名称',
        ])
        country = _csv_lookup(raw, ['country', 'country code', 'country_code', 'geo', 'region'])
        country_label = normalize_country_label(country) if country else 'Unknown'
        platform = _ad_platform_label(normalized_media_source, channel, campaign, ad_group, ad_name)
        account_label = _ad_account_label_or_missing(ad_account)
        key = (
            row_date.isoformat(),
            account_label,
            platform,
            country_label,
            normalized_media_source,
            campaign,
            ad_group,
            ad_name,
        )
        bucket = buckets.setdefault(key, {
            'date': row_date.isoformat(),
            'data_source': 'AppsFlyer',
            'app_id': account_label,
            'appsflyer_app_id': app_id,
            'platform': platform,
            'country': key[3],
            'country_attribution_status': (
                'appsflyer_report_country'
                if country
                else 'unresolved_waiting_meta_delivery_country'
            ),
            'country_attribution_source': (
                'appsflyer_raw_event_report'
                if country
                else 'meta_insights_country_breakdown_pending'
            ),
            'media_source': normalized_media_source,
            'campaign': campaign,
            'ad_group': ad_group,
            'ad': ad_name,
            'source_type': '推广量' if normalized_media_source.lower() not in {'organic', 'none'} else '自然量',
            **_empty_ad_metrics(),
        })
        normalized_event = event_name.strip().lower()
        normalized_event_key = re.sub(r'[^a-z0-9]+', '_', normalized_event).strip('_')
        normalized_event_compact = re.sub(r'[^a-z0-9]+', '', normalized_event)
        if normalized_event_key in {'af_complete_registration', 'complete_registration'} or normalized_event_compact == 'afcompleteregistration':
            bucket['af_registrations'] = float(bucket.get('af_registrations') or 0.0) + 1.0
            bucket['registrations'] = float(bucket.get('registrations') or 0.0) + 1.0
        elif normalized_event_key == 'l1_task_high_value':
            bucket['high_value_users'] = float(bucket.get('high_value_users') or 0.0) + 1.0
        elif normalized_event_key in {'enter_im', 'im_enter', 'im_entry', 'im_user_first_manual_reply'}:
            bucket['im_entries'] = float(bucket.get('im_entries') or 0.0) + 1.0
            if normalized_event_key == 'im_user_first_manual_reply':
                bucket['im_first_replies'] = float(bucket.get('im_first_replies') or 0.0) + 1.0
        elif normalized_event_key == 'im_apply_step2_triggered':
            bucket['im_step2_triggers'] = float(bucket.get('im_step2_triggers') or 0.0) + 1.0
        elif normalized_event_key == 'im_user_manual_reply_count_3':
            bucket['im_manual_reply_3'] = float(bucket.get('im_manual_reply_3') or 0.0) + 1.0
        elif normalized_event_key == 'im_link_click':
            bucket['im_link_clicks'] = float(bucket.get('im_link_clicks') or 0.0) + 1.0
        elif (
            normalized_event_key in {'join_guild', 'join_group', 'joingroup', 'guild_join', 'guild_joins'}
            or normalized_event_compact in {'joinguild', 'joingroup', 'guildjoin', 'guildjoins'}
            or normalized_event in {'入会', '加入公会'}
        ):
            bucket['af_guild_joins'] = float(bucket.get('af_guild_joins') or 0.0) + 1.0
    return list(buckets.values())


def _fetch_appsflyer_app_ids(*, token: str, base_url: str, session: Any) -> List[str]:
    response = session.get(
        f'{str(base_url or "https://hq1.appsflyer.com").rstrip("/")}/api/mng/apps',
        params={'limit': 1000, 'offset': 0},
        headers={'Authorization': f'Bearer {token}'},
        timeout=20.0,
    )
    response.raise_for_status()
    payload = response.json()
    app_ids: List[str] = []
    for row in payload.get('data') or []:
        app_id = str((row or {}).get('id') or '').strip()
        if app_id:
            app_ids.append(app_id)
    return _parse_config_list(app_ids)


def _dashboard_iso_start_of_day(value: datetime.date, timezone_name: str = 'UTC') -> str:
    try:
        tz = ZoneInfo(str(timezone_name or 'UTC').strip() or 'UTC')
    except Exception:
        tz = timezone.utc
    start = datetime.combine(value, datetime.min.time(), tzinfo=tz)
    return start.astimezone(timezone.utc).isoformat().replace('+00:00', 'Z')


def _dashboard_event_date(value: Any, timezone_name: str) -> Optional[datetime.date]:
    text = str(value or '').strip()
    if not text:
        return None
    normalized = text.replace('Z', '+00:00')
    try:
        dt = datetime.fromisoformat(normalized)
    except Exception:
        return _parse_dashboard_date(text)
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    try:
        tz = ZoneInfo(str(timezone_name or 'UTC').strip() or 'UTC')
    except Exception:
        tz = timezone.utc
    return dt.astimezone(tz).date()


def _fetch_bind_success_event_rows(
    *,
    token: str,
    base_url: str,
    from_date: datetime.date,
    to_date: datetime.date,
    project: str = 'TUGAO',
    timezone_name: str = 'UTC',
    session: Any,
    page_size: int = 500,
    max_pages: int = 20,
) -> List[Dict[str, Any]]:
    normalized_token = str(token or '').strip()
    if not normalized_token:
        return []
    url = f'{str(base_url or "https://servertest.timetrade.club").rstrip("/")}/api/v1/analytics/bind-success-events'
    params: Dict[str, Any] = {
        'start_time': _dashboard_iso_start_of_day(from_date, timezone_name),
        'end_time': _dashboard_iso_start_of_day(to_date + timedelta(days=1), timezone_name),
        'page_size': min(max(int(page_size or 500), 1), 500),
    }
    normalized_project = str(project or '').strip()
    if normalized_project:
        params['project'] = normalized_project
    rows: List[Dict[str, Any]] = []
    cursor = ''
    for _ in range(max(int(max_pages or 1), 1)):
        request_params = dict(params)
        if cursor:
            request_params['cursor'] = cursor
        response = session.get(
            url,
            params=request_params,
            headers={'Authorization': f'Bearer {normalized_token}'},
            timeout=30.0,
        )
        response.raise_for_status()
        payload = response.json()
        for raw in payload.get('data') or []:
            if not isinstance(raw, dict):
                continue
            if str(raw.get('bind_status') or '').strip().lower() != 'success':
                continue
            row_date = (
                _dashboard_event_date(raw.get('bind_success_time'), timezone_name)
                or _dashboard_event_date(raw.get('updated_at'), timezone_name)
                or _parse_dashboard_date(raw.get('business_date_jakarta'))
                or to_date
            )
            media_source = str(raw.get('media_source') or '').strip() or 'Internal'
            is_organic_bind = _ad_is_organic_media_source(media_source)
            campaign = str(raw.get('campaign_name') or raw.get('campaign_id') or '').strip() or '未命名'
            ad_group = str(raw.get('adset_name') or raw.get('adset_id') or '').strip()
            ad_name = str(raw.get('ad_name') or raw.get('ad_id') or '').strip()
            raw_country = str(raw.get('country') or '').strip()
            country = normalize_country_label(raw_country) if raw_country else 'Unknown'
            platform = _ad_platform_label(media_source, '', campaign, ad_group, ad_name)
            rows.append({
                'date': row_date.isoformat(),
                'data_source': 'BindSuccess',
                'app_id': _ad_missing_account_label(),
                'appsflyer_app_id': str(raw.get('project') or normalized_project or 'TUGAO').strip() or 'TUGAO',
                'platform': platform,
                'country': country,
                'country_attribution_status': (
                    'bind_success_country'
                    if raw_country
                    else 'unresolved_waiting_meta_delivery_country'
                ),
                'country_attribution_source': (
                    'bind_success_api'
                    if raw_country
                    else 'meta_insights_country_breakdown_pending'
                ),
                'media_source': media_source,
                'campaign': campaign,
                'ad_group': ad_group,
                'ad': ad_name,
                'source_type': '自然量入会' if is_organic_bind else '推广入会',
                'guild_joins': 1.0,
                'promotion_guild_joins': 0.0 if is_organic_bind else 1.0,
                'organic_guild_joins': 1.0 if is_organic_bind else 0.0,
                'bind_success_has_wa': bool(raw.get('has_wa')),
            })
        cursor = str(payload.get('next_cursor') or '').strip()
        if not payload.get('has_more') or not cursor:
            break
    return rows


def _normalize_meta_api_version(value: str) -> str:
    normalized = str(value or '').strip().lower()
    if not normalized:
        return 'v25.0'
    return normalized if normalized.startswith('v') else f'v{normalized}'


def _normalize_meta_ad_account_id(value: str) -> str:
    normalized = str(value or '').strip()
    if not normalized:
        return ''
    return normalized if normalized.startswith('act_') else f'act_{normalized}'


def _meta_action_value(actions: Any, candidates: List[str]) -> float:
    wanted = {str(candidate or '').strip().lower() for candidate in candidates if str(candidate or '').strip()}
    total = 0.0
    for row in actions or []:
        if not isinstance(row, dict):
            continue
        action_type = str(row.get('action_type') or '').strip().lower()
        if action_type in wanted:
            total += _dashboard_float(row.get('value'))
    return total


def _meta_first_action_value(actions: Any, candidates: List[str]) -> float:
    for candidate in candidates:
        value = _meta_action_value(actions, [candidate])
        if value:
            return value
    return 0.0


def _fetch_meta_ad_accounts(*, token: str, api_version: str, base_url: str, session: Any) -> List[str]:
    url = f'{str(base_url or "https://graph.facebook.com").rstrip("/")}/{_normalize_meta_api_version(api_version)}/me/adaccounts'
    account_ids: List[str] = []
    params = {
        'fields': 'id,account_id,name,account_status,currency,timezone_name',
        'limit': 200,
    }
    next_url = url
    while next_url:
        response = session.get(
            next_url,
            params=params if next_url == url else None,
            headers={'Authorization': f'Bearer {token}'},
            timeout=20.0,
        )
        response.raise_for_status()
        payload = response.json()
        for row in payload.get('data') or []:
            account_id = str((row or {}).get('id') or (row or {}).get('account_id') or '').strip()
            normalized_account_id = _normalize_meta_ad_account_id(account_id)
            if normalized_account_id:
                account_ids.append(normalized_account_id)
        next_url = str(((payload.get('paging') or {}).get('next')) or '').strip()
        params = None
    return _parse_config_list(account_ids)


def _fetch_meta_ad_account_timezone(*, token: str, ad_account_id: str, api_version: str, base_url: str, session: Any) -> str:
    normalized_account_id = _normalize_meta_ad_account_id(ad_account_id)
    if not normalized_account_id:
        return 'UTC'
    url = f'{str(base_url or "https://graph.facebook.com").rstrip("/")}/{_normalize_meta_api_version(api_version)}/{quote(normalized_account_id, safe="")}'
    response = session.get(
        url,
        params={'fields': 'timezone_name'},
        headers={'Authorization': f'Bearer {token}'},
        timeout=20.0,
    )
    response.raise_for_status()
    payload = response.json()
    timezone_name = str((payload or {}).get('timezone_name') or '').strip()
    return timezone_name or 'UTC'


def _meta_hour_start(value: Any) -> Optional[int]:
    text = str(value or '').strip()
    if not text:
        return None
    match = re.search(r'(\d{1,2}):\d{2}', text)
    if not match:
        return None
    hour = int(match.group(1))
    return hour if 0 <= hour <= 23 else None


def _meta_row_utc_date(raw: Dict[str, Any], account_timezone: str) -> Optional[datetime.date]:
    row_date = _parse_dashboard_date(raw.get('date_start'))
    if row_date is None:
        return None
    hour = _meta_hour_start(raw.get('hourly_stats_aggregated_by_advertiser_time_zone'))
    if hour is None:
        return row_date
    try:
        account_tz = ZoneInfo(str(account_timezone or 'UTC').strip() or 'UTC')
    except Exception:
        account_tz = timezone.utc
    account_dt = datetime.combine(row_date, datetime.min.time(), tzinfo=account_tz) + timedelta(hours=hour)
    return account_dt.astimezone(timezone.utc).date()


def _meta_local_day_overlapping_utc_dates(value: Any, account_timezone: str) -> List[datetime.date]:
    local_date = _parse_dashboard_date(value)
    if local_date is None:
        return []
    try:
        account_tz = ZoneInfo(str(account_timezone or 'UTC').strip() or 'UTC')
    except Exception:
        account_tz = timezone.utc
    local_start = datetime.combine(local_date, datetime.min.time(), tzinfo=account_tz)
    utc_start = local_start.astimezone(timezone.utc)
    utc_end = (local_start + timedelta(days=1)).astimezone(timezone.utc)
    final_date = (utc_end - timedelta(microseconds=1)).date()
    dates: List[datetime.date] = []
    cursor = utc_start.date()
    while cursor <= final_date:
        dates.append(cursor)
        cursor += timedelta(days=1)
    return dates


def _meta_insight_params(
    *,
    from_date: datetime.date,
    to_date: datetime.date,
    breakdowns: str,
    include_actions: bool = True,
) -> Dict[str, Any]:
    fields = [
            'date_start',
            'account_id',
            'account_name',
            'campaign_id',
            'campaign_name',
            'adset_id',
            'adset_name',
            'ad_id',
            'ad_name',
            'spend',
            'impressions',
            'reach',
            'frequency',
            'cpm',
            'cpc',
            'ctr',
            'clicks',
            'inline_link_clicks',
    ]
    if include_actions:
        fields.extend([
            'actions',
            'action_values',
            'conversions',
        ])
    return {
        'fields': ','.join(fields),
        'time_increment': 1,
        'level': 'ad',
        'breakdowns': breakdowns,
        'limit': 500,
        'time_range': json.dumps({'since': from_date.isoformat(), 'until': to_date.isoformat()}),
    }


def _fetch_meta_insight_rows(
    *,
    token: str,
    ad_account_id: str,
    api_version: str,
    base_url: str,
    from_date: datetime.date,
    to_date: datetime.date,
    session: Any,
    account_timezone: str = 'UTC',
    hourly: bool = False,
    include_actions: bool = True,
    suppress_media_metrics: bool = False,
) -> List[Dict[str, Any]]:
    normalized_account_id = _normalize_meta_ad_account_id(ad_account_id)
    if not normalized_account_id:
        return []
    url = f'{str(base_url or "https://graph.facebook.com").rstrip("/")}/{_normalize_meta_api_version(api_version)}/{quote(normalized_account_id, safe="")}/insights'
    breakdowns = 'hourly_stats_aggregated_by_advertiser_time_zone' if hourly else 'country'
    country_evidence_only = bool(not hourly and suppress_media_metrics)
    fetch_from = from_date - timedelta(days=1) if hourly or country_evidence_only else from_date
    fetch_to = to_date + timedelta(days=1) if hourly or country_evidence_only else to_date
    params: Optional[Dict[str, Any]] = _meta_insight_params(
        from_date=fetch_from,
        to_date=fetch_to,
        breakdowns=breakdowns,
        include_actions=include_actions,
    )
    parsed_rows: List[Dict[str, Any]] = []
    next_url = url
    retried_without_breakdowns = False
    while next_url:
        response = session.get(
            next_url,
            params=params if next_url == url else None,
            headers={'Authorization': f'Bearer {token}'},
            timeout=30.0,
        )
        if (hourly or country_evidence_only) and response.status_code >= 400:
            response.raise_for_status()
        if response.status_code >= 400 and params and params.get('breakdowns') and not retried_without_breakdowns:
            retried_without_breakdowns = True
            params = dict(params)
            params.pop('breakdowns', None)
            next_url = url
            response = session.get(
                next_url,
                params=params,
                headers={'Authorization': f'Bearer {token}'},
                timeout=30.0,
            )
        response.raise_for_status()
        payload = response.json()
        for raw in payload.get('data') or []:
            if not isinstance(raw, dict):
                continue
            row_date = _meta_row_utc_date(raw, account_timezone) if hourly else _parse_dashboard_date(raw.get('date_start'))
            if row_date is None:
                continue
            evidence_dates = (
                _meta_local_day_overlapping_utc_dates(raw.get('date_start'), account_timezone)
                if country_evidence_only
                else []
            )
            if (row_date < from_date or row_date > to_date) and not any(
                from_date <= evidence_date <= to_date for evidence_date in evidence_dates
            ):
                continue
            actions = (raw.get('actions') or []) if include_actions else []
            action_values = (raw.get('action_values') or []) if include_actions else []
            conversions = (raw.get('conversions') or []) if include_actions else []
            account_label = str(raw.get('account_name') or raw.get('account_id') or normalized_account_id).strip()
            campaign_label = str(raw.get('campaign_name') or raw.get('campaign_id') or '').strip()
            ad_group_label = str(raw.get('adset_name') or raw.get('adset_id') or '').strip()
            ad_label = str(raw.get('ad_name') or raw.get('ad_id') or '').strip()
            raw_country = str(raw.get('country') or '').strip()
            country_label = normalize_country_label(raw_country) if raw_country else 'Unknown'
            meta_installs = _meta_first_action_value(actions, [
                'mobile_app_install',
                'app_install',
                'omni_app_install',
            ])
            meta_registrations = _meta_first_action_value(actions, [
                'fb_mobile_complete_registration',
                'app_custom_event.fb_mobile_complete_registration',
                'complete_registration',
                'omni_complete_registration',
                'offsite_conversion.fb_pixel_complete_registration',
            ])
            meta_guild_joins = _meta_first_action_value(actions, [
                'app_custom_event.join_guild',
                'app_custom_event.joingroup',
                'app_custom_event.join_group',
                'offsite_conversion.custom.join_guild',
                'offsite_conversion.custom.joingroup',
                'offsite_conversion.custom.join_group',
                'join_guild',
                'joingroup',
                'join_group',
                'app_custom_event.SubmitApplication',
                'app_custom_event.submit_application',
                'offsite_conversion.custom.SubmitApplication',
                'offsite_conversion.custom.submit_application',
                'SubmitApplication',
                'submit_application',
            ])
            if not meta_guild_joins:
                meta_guild_joins = _meta_first_action_value(conversions, [
                    'submit_application_mobile_app',
                    'submit_application_total',
                    'submit_application',
                    'app_custom_event.join_guild',
                    'app_custom_event.joingroup',
                    'app_custom_event.join_group',
                    'offsite_conversion.custom.join_guild',
                    'offsite_conversion.custom.joingroup',
                    'offsite_conversion.custom.join_group',
                    'join_guild',
                    'joingroup',
                    'join_group',
                    'app_custom_event.SubmitApplication',
                    'app_custom_event.submit_application',
                    'offsite_conversion.custom.SubmitApplication',
                    'offsite_conversion.custom.submit_application',
                    'SubmitApplication',
                ])
            parsed_row = {
                'date': row_date.isoformat(),
                'data_source': 'Meta',
                'platform': 'Meta',
                'account_id': normalized_account_id,
                'target_app': _ad_dashboard_target_app_from_account(
                    platform='Meta',
                    account_id=normalized_account_id,
                    account_label=account_label,
                ),
                'app_id': account_label or normalized_account_id,
                'country': country_label,
                'country_attribution_status': (
                    'meta_delivery_country'
                    if raw_country
                    else 'unresolved_waiting_meta_delivery_country'
                ),
                'country_attribution_source': (
                    'meta_insights_country_breakdown'
                    if raw_country
                    else 'meta_hourly_insights'
                ),
                'media_source': 'Meta',
                'campaign': campaign_label or '未命名',
                'ad_group': ad_group_label,
                'ad': ad_label,
                'campaign_id': str(raw.get('campaign_id') or '').strip(),
                'adset_id': str(raw.get('adset_id') or '').strip(),
                'ad_id': str(raw.get('ad_id') or '').strip(),
                'ad_name': ad_label,
                'source_type': '推广量',
                'cost': 0.0 if suppress_media_metrics else _dashboard_float(raw.get('spend')),
                'installs': meta_installs,
                'meta_installs': meta_installs,
                'registrations': meta_registrations,
                'meta_registrations': meta_registrations,
                'high_value_users': 0.0,
                'im_first_replies': 0.0,
                'im_step2_triggers': 0.0,
                'im_manual_reply_3': 0.0,
                'im_link_clicks': 0.0,
                'guild_joins': 0.0,
                'meta_guild_joins': meta_guild_joins,
                'purchases': 0.0,
                'revenue': 0.0,
                'clicks': 0.0 if suppress_media_metrics else _dashboard_float(raw.get('clicks')),
                'link_clicks': 0.0 if suppress_media_metrics else _dashboard_float(raw.get('inline_link_clicks') or raw.get('clicks')),
                'impressions': 0.0 if suppress_media_metrics else _dashboard_float(raw.get('impressions')),
                'reach': 0.0 if suppress_media_metrics else _dashboard_float(raw.get('reach')),
            }
            if from_date <= row_date <= to_date:
                parsed_rows.append(parsed_row)
            if country_evidence_only and raw_country:
                for evidence_date in evidence_dates:
                    if evidence_date == row_date or evidence_date < from_date or evidence_date > to_date:
                        continue
                    evidence_row = dict(parsed_row)
                    evidence_row['date'] = evidence_date.isoformat()
                    evidence_row.update(_empty_ad_metrics())
                    evidence_row['country_attribution_grain'] = 'meta_local_day_utc_overlap'
                    parsed_rows.append(evidence_row)
        next_url = str(((payload.get('paging') or {}).get('next')) or '').strip()
        params = None
    return parsed_rows


def build_ad_data_dashboard_snapshot(
    *,
    token: str,
    app_ids: List[str],
    timezone_name: str = 'UTC',
    base_url: str = 'https://hq1.appsflyer.com',
    session: Any = requests,
    meta_token: str = '',
    meta_ad_account_ids: Optional[List[str]] = None,
    meta_api_version: str = 'v25.0',
    meta_base_url: str = 'https://graph.facebook.com',
    meta_session: Any = None,
    bind_success_token: str = '',
    bind_success_base_url: str = 'https://servertest.timetrade.club',
    bind_success_project: str = 'TUGAO',
    bind_success_session: Any = None,
    days: int = 30,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    filters: Optional[Dict[str, Any]] = None,
    platform_filters: Optional[Dict[str, Any]] = None,
    platform_date_windows: Optional[Dict[str, Any]] = None,
    target_app: str = 'all',
    top_limit: int = 8,
    include_fact_rows: bool = False,
) -> Dict[str, Any]:
    normalized_token = str(token or '').strip()
    normalized_app_ids = _parse_config_list(app_ids)
    normalized_meta_token = str(meta_token or '').strip()
    normalized_meta_ad_account_ids = _parse_config_list(meta_ad_account_ids or [])
    meta_account_access_policy = MetaAdAccountAccessPolicy.from_environment()
    meta_account_access_decisions = [
        meta_account_access_policy.configured(account_id)
        for account_id in normalized_meta_ad_account_ids
    ]
    normalized_meta_api_version = _normalize_meta_api_version(meta_api_version)
    normalized_bind_success_token = str(bind_success_token or '').strip()
    normalized_bind_success_base_url = str(bind_success_base_url or 'https://servertest.timetrade.club').strip() or 'https://servertest.timetrade.club'
    normalized_bind_success_project = str(bind_success_project or 'TUGAO').strip() or 'TUGAO'
    normalized_filters = _normalize_ad_filters(filters)
    normalized_platform_filters = _normalize_ad_platform_filters(platform_filters)
    normalized_target_app = _normalize_ad_dashboard_target_app(target_app) or 'all'
    normalized_top_limit = min(max(int(top_limit or 8), 3), 25)
    meta_session = meta_session or session
    bind_success_session = bind_success_session or session
    try:
        tz = ZoneInfo(str(timezone_name or 'Asia/Shanghai').strip())
    except Exception:
        tz = timezone.utc
        timezone_name = 'UTC'
    today = _ad_dashboard_latest_complete_utc_date()
    start_date, end_date, window_days, window_warnings = _coerce_ad_dashboard_window(
        days=days,
        date_from=date_from,
        date_to=date_to,
        today=today,
    )
    normalized_platform_date_windows, platform_start_date, platform_end_date, platform_window_days, platform_window_warnings = _normalize_ad_platform_date_windows(
        platform_date_windows,
        fallback_start_date=start_date,
        fallback_end_date=end_date,
        today=today,
    )
    window_warnings.extend(platform_window_warnings)
    start_date = platform_start_date
    end_date = platform_end_date
    window_days = platform_window_days
    overview_start_date = min(start_date, end_date - timedelta(days=29))
    fetch_start_date = overview_start_date
    fetch_end_date = end_date
    appsflyer_rows: List[Dict[str, Any]] = []
    meta_rows: List[Dict[str, Any]] = []
    bind_success_rows: List[Dict[str, Any]] = []
    errors: List[Dict[str, str]] = []
    appsflyer_partner_report = {'ok': 0, 'failed': 0, 'row_count': 0}
    appsflyer_event_report = {'ok': 0, 'failed': 0, 'row_count': 0}
    appsflyer_event_metric_source = {
        'raw': 0,
        'partner_aggregate': 0,
        'missing': 0,
    }
    meta_report = {
        'hourly_ok': 0,
        'hourly_failed': 0,
        'daily_actions_ok': 0,
        'daily_actions_failed': 0,
        'daily_fallback': 0,
        'timezone_lookup_failed': 0,
        'rate_limited': 0,
        'account_access': access_summary(meta_account_access_decisions),
    }
    bind_success_report = {'ok': 0, 'failed': 0, 'row_count': 0}
    marketing_diagnostics_token = str(os.getenv('MARKETING_DIAGNOSTICS_API_TOKEN') or os.getenv('BI_MARKETING_DIAGNOSTICS_API_TOKEN') or '').strip()
    if marketing_diagnostics_token and not include_fact_rows:
        try:
            from app.marketing_diagnostics_api import DEFAULT_MARKETING_DIAGNOSTICS_DAILY_URL, MarketingDiagnosticsDailyClient

            page_size = int(os.getenv('MARKETING_DIAGNOSTICS_API_PAGE_SIZE') or 500)
            marketing_client = MarketingDiagnosticsDailyClient(
                token=marketing_diagnostics_token,
                base_url=(
                    os.getenv('MARKETING_DIAGNOSTICS_DAILY_API_BASE_URL')
                    or os.getenv('BI_MARKETING_DIAGNOSTICS_DAILY_API_BASE_URL')
                    or DEFAULT_MARKETING_DIAGNOSTICS_DAILY_URL
                ),
                session=session,
                page_size=page_size,
            )
            marketing_result = marketing_client.fetch(
                start_date=start_date,
                end_date=end_date,
                datasets=['ad_daily', 'natural_im_funnel_daily'],
            )
            if marketing_result.rows:
                payload = build_ad_data_dashboard_snapshot_from_rows(
                    marketing_result.rows,
                    timezone_name=str(timezone_name or 'UTC'),
                    days=window_days,
                    date_from=start_date.isoformat(),
                    date_to=end_date.isoformat(),
                    filters=normalized_filters,
                    platform_filters=normalized_platform_filters,
                    platform_date_windows=normalized_platform_date_windows,
                    target_app=normalized_target_app,
                    top_limit=normalized_top_limit,
                    sources={
                        'marketing_diagnostics': {
                            'configured': True,
                            'row_count': len(marketing_result.rows),
                            'raw_row_count': marketing_result.raw_row_count,
                            'pages': marketing_result.pages,
                            'datasets': marketing_result.datasets,
                            'label': 'Marketing Diagnostics',
                        },
                    },
                    insights=window_warnings + ['TimeTrade marketing-diagnostics daily 已作为本窗口主事实行。'],
                    errors=[],
                )
                payload['source'] = 'marketing_diagnostics_daily_api'
                payload['configured'] = True
                if include_fact_rows:
                    payload['_fact_rows'] = _ad_materialize_fact_rows(marketing_result.rows)
                return payload
        except Exception as exc:
            errors.append({'source': 'Marketing Diagnostics', 'app_id': 'daily', 'message': _dashboard_error_message(exc)})
    app_ids_source = 'configured'
    app_ids_status = 'not_configured' if not normalized_token else ('configured' if normalized_app_ids else 'pending')
    if normalized_token and not normalized_app_ids:
        try:
            normalized_app_ids = _fetch_appsflyer_app_ids(
                token=normalized_token,
                base_url=base_url,
                session=session,
            )
            app_ids_source = 'app_list_api'
            app_ids_status = 'ok'
        except Exception as exc:
            app_ids_status = 'failed'
            errors.append({'source': 'AppsFlyer App List', 'app_id': 'app-list', 'message': _dashboard_error_message(exc)})
    elif normalized_token and normalized_app_ids:
        app_ids_status = 'configured'
    appsflyer_configured = bool(normalized_token and normalized_app_ids)
    if appsflyer_configured:
        for app_id in normalized_app_ids:
            partner_rows: List[Dict[str, Any]] = []
            try:
                partner_rows = _fetch_appsflyer_partner_by_date_rows(
                    token=normalized_token,
                    app_id=app_id,
                    base_url=base_url,
                    timezone_name=str(timezone_name or ''),
                    from_date=fetch_start_date,
                    to_date=fetch_end_date,
                    session=session,
                )
                appsflyer_rows.extend(partner_rows)
                appsflyer_partner_report['ok'] += 1
                appsflyer_partner_report['row_count'] += len(partner_rows)
            except Exception as exc:
                appsflyer_partner_report['failed'] += 1
                errors.append({'source': 'AppsFlyer Partner Report', 'app_id': app_id, 'message': _dashboard_error_message(exc)})
            partner_has_events = any(
                float((row or {}).get(key) or 0.0)
                for row in partner_rows if isinstance(row, dict)
                for key in ('af_registrations', 'high_value_users', 'im_entries', 'im_manual_reply_3', 'guild_joins')
            )
            if partner_has_events:
                appsflyer_event_metric_source['partner_aggregate'] += 1
            else:
                try:
                    event_rows = _fetch_appsflyer_in_app_event_rows(
                        token=normalized_token,
                        app_id=app_id,
                        base_url=base_url,
                        timezone_name=str(timezone_name or ''),
                        from_date=fetch_start_date,
                        to_date=fetch_end_date,
                        session=session,
                    )
                    appsflyer_rows.extend(event_rows)
                    appsflyer_event_report['ok'] += 1
                    appsflyer_event_report['row_count'] += len(event_rows)
                    if event_rows:
                        appsflyer_event_metric_source['raw'] += 1
                    else:
                        appsflyer_event_metric_source['missing'] += 1
                except Exception as exc:
                    appsflyer_event_report['failed'] += 1
                    appsflyer_event_metric_source['missing'] += 1
                    errors.append({'source': 'AppsFlyer Raw Events', 'app_id': app_id, 'message': _dashboard_error_message(exc)})
    meta_account_ids_source = 'configured'
    if normalized_meta_token and not normalized_meta_ad_account_ids:
        try:
            normalized_meta_ad_account_ids = _fetch_meta_ad_accounts(
                token=normalized_meta_token,
                api_version=normalized_meta_api_version,
                base_url=meta_base_url,
                session=meta_session,
            )
            meta_account_ids_source = 'account_list_api'
        except Exception as exc:
            errors.append({'source': 'Meta', 'app_id': 'ad-account-list', 'message': _dashboard_error_message(exc)})
    meta_configured = bool(normalized_meta_token and normalized_meta_ad_account_ids)
    if meta_configured:
        for ad_account_id in normalized_meta_ad_account_ids:
            configured_access = meta_account_access_policy.configured(ad_account_id)
            if not configured_access.should_sync:
                continue
            account_timezone = 'UTC'
            try:
                account_timezone = _fetch_meta_ad_account_timezone(
                    token=normalized_meta_token,
                    ad_account_id=ad_account_id,
                    api_version=normalized_meta_api_version,
                    base_url=meta_base_url,
                    session=meta_session,
                )
            except Exception:
                meta_report['timezone_lookup_failed'] += 1
            try:
                account_rows = _fetch_meta_insight_rows(
                    token=normalized_meta_token,
                    ad_account_id=ad_account_id,
                    api_version=normalized_meta_api_version,
                    base_url=meta_base_url,
                    from_date=fetch_start_date,
                    to_date=fetch_end_date,
                    session=meta_session,
                    account_timezone=account_timezone,
                    hourly=True,
                    include_actions=False,
                )
                meta_rows.extend(account_rows)
                meta_report['hourly_ok'] += 1
                try:
                    meta_rows.extend(_fetch_meta_insight_rows(
                        token=normalized_meta_token,
                        ad_account_id=ad_account_id,
                        api_version=normalized_meta_api_version,
                        base_url=meta_base_url,
                        from_date=fetch_start_date,
                        to_date=fetch_end_date,
                        session=meta_session,
                        account_timezone=account_timezone,
                        hourly=False,
                        include_actions=True,
                        suppress_media_metrics=True,
                    ))
                    meta_report['daily_actions_ok'] += 1
                except MetaRateLimitBlocked:
                    meta_report['daily_actions_failed'] += 1
                    meta_report['rate_limited'] += 1
                except Exception:
                    meta_report['daily_actions_failed'] += 1
            except MetaRateLimitBlocked:
                meta_report['hourly_failed'] += 1
                meta_report['rate_limited'] += 1
            except Exception as hourly_exc:
                meta_report['hourly_failed'] += 1
                try:
                    meta_rows.extend(_fetch_meta_insight_rows(
                        token=normalized_meta_token,
                        ad_account_id=ad_account_id,
                        api_version=normalized_meta_api_version,
                        base_url=meta_base_url,
                        from_date=fetch_start_date,
                        to_date=fetch_end_date,
                        session=meta_session,
                        account_timezone=account_timezone,
                        hourly=False,
                    ))
                    meta_report['daily_fallback'] += 1
                except Exception as exc:
                    access_decision = classify_meta_exception(
                        meta_account_access_policy, ad_account_id, exc,
                    )
                    if access_decision is not None:
                        meta_account_access_decisions = [
                            access_decision if row.account_id == access_decision.account_id else row
                            for row in meta_account_access_decisions
                        ]
                        meta_report['account_access'] = access_summary(meta_account_access_decisions)
                    if access_decision is None or access_decision.should_alert:
                        errors.append({'source': 'Meta', 'app_id': ad_account_id, 'message': _dashboard_error_message(exc)})
    bind_success_configured = bool(normalized_bind_success_token)
    if bind_success_configured:
        try:
            bind_success_rows = _fetch_bind_success_event_rows(
                token=normalized_bind_success_token,
                base_url=normalized_bind_success_base_url,
                from_date=fetch_start_date,
                to_date=fetch_end_date,
                project=normalized_bind_success_project,
                timezone_name=str(timezone_name or ''),
                session=bind_success_session,
            )
            bind_success_report['ok'] = 1
            bind_success_report['row_count'] = len(bind_success_rows)
        except Exception as exc:
            bind_success_report['failed'] = 1
            errors.append({'source': 'BindSuccess', 'app_id': normalized_bind_success_project, 'message': _dashboard_error_message(exc)})
    raw_dashboard_rows = _ad_enrich_countries_from_meta_delivery(
        appsflyer_rows + meta_rows + bind_success_rows
    )
    all_app_rows = _ad_dashboard_apply_target_app_filter(raw_dashboard_rows, 'all')
    all_rows = _ad_dashboard_apply_target_app_filter(raw_dashboard_rows, normalized_target_app)
    globally_filtered_rows = _filter_ad_rows(all_rows, normalized_filters)
    all_app_globally_filtered_rows = _filter_ad_rows(all_app_rows, normalized_filters)
    selected_globally_filtered_rows = _ad_rows_in_date_range(globally_filtered_rows, start_date, end_date)
    platform_date_filtered_rows = _filter_ad_rows_by_platform_date_windows(selected_globally_filtered_rows, normalized_platform_date_windows)
    rows = _filter_ad_rows_by_platform_filters(platform_date_filtered_rows, normalized_platform_filters)
    # Platform filters are detail-panel filters. The top day/week/month overview
    # stays scoped only by global date/app filters so drilling into Meta country
    # or account does not rewrite the dashboard-level totals.
    period_rows = globally_filtered_rows
    period_all_app_rows = all_app_globally_filtered_rows
    configured = bool(appsflyer_configured or meta_configured or bind_success_configured)
    fact_dates = {
        str((row or {}).get('date') or '').strip()
        for row in rows
        if str((row or {}).get('date') or '').strip()
    }
    daily_rows: List[Dict[str, Any]] = []
    for offset in range(window_days):
        day = start_date + timedelta(days=offset)
        metrics = _ad_period_summary(rows, day, day)
        daily_rows.append({'date': day.isoformat(), 'has_fact_data': day.isoformat() in fact_dates, **metrics})
    day_summary = _ad_period_summary(period_rows, end_date, end_date)
    week_summary = _ad_period_summary(period_rows, end_date - timedelta(days=6), end_date)
    month_summary = _ad_period_summary(period_rows, end_date - timedelta(days=29), end_date)
    selected_summary = _ad_period_summary(rows, start_date, end_date)
    insights = list(window_warnings)
    if not configured:
        insights.append('等待配置 AppsFlyer 或广告平台 API token。')
    elif errors and not rows:
        insights.append('广告数据拉取失败，请检查 token、账户 ID 和接口权限。')
    elif not rows and all_rows:
        insights.append('当前筛选条件下暂无数据，请放宽账户、渠道、广告系列或广告筛选。')
    elif not any(float(row.get('cost') or 0) or float(row.get('installs') or 0) for row in rows):
        insights.append('当前筛选窗口暂无广告消耗或安装数据。')
    else:
        top_media = _ad_top_breakdown(rows, 'media_source', limit=1)
        if top_media:
            insights.append(f"本期消耗最高渠道：{top_media[0]['name']}，安装 {int(top_media[0].get('installs') or 0)}。")
        if appsflyer_rows and not meta_rows and any(float(row.get('installs') or 0) for row in rows) and not any(float(row.get('cost') or 0) for row in rows):
            insights.append('AF 当前返回安装/点击/展示，但成本为 0；请确认 AppsFlyer cost 集成或广告平台 token 是否已接入。')
        if appsflyer_event_metric_source.get('partner_aggregate'):
            insights.append('AF 后链路已使用聚合报表事件列 fallback；可用于趋势和日报，但 raw 明细权限恢复前不能做用户级事件追溯。')
        if appsflyer_partner_report['row_count'] and appsflyer_event_report['failed'] and appsflyer_event_metric_source.get('missing'):
            insights.append('AF 基础聚合报表已返回安装/点击/展示，但 raw in-app events 拉取失败；未随聚合报表返回的高价值、自动报名人数、IM>=3 等后链路指标会显示为 0 或缺失。')
        elif appsflyer_partner_report['row_count'] and not appsflyer_event_report['row_count'] and not any(float(row.get(key) or 0) for row in rows for key in ('af_registrations', 'im_entries', 'im_manual_reply_3', 'guild_joins')):
            insights.append('AF 聚合报表未返回后链路事件列，且 raw in-app events 没有事件行；请确认 AppsFlyer 事件名和 raw data 权限。')
        if meta_configured and meta_rows:
            if meta_report.get('hourly_ok'):
                insights.append(f"Meta 已按小时数据聚合到 UTC+0/GMT 日窗：{meta_report.get('hourly_ok')} 个广告账户。")
            if meta_report.get('daily_fallback'):
                insights.append(f"Meta 有 {meta_report.get('daily_fallback')} 个广告账户小时数据不可用，已 fallback 到日级 insight；该部分不是严格 08:00 切日。")
            else:
                insights.append(f"Meta 已接入 {len(normalized_meta_ad_account_ids)} 个广告账户，可与 AF 归因数据做口径对比。")
        if bind_success_rows:
            insights.append(f"真实入会成功事件已接入 {len(bind_success_rows)} 行，用于补齐入会口径。")
        if float(week_summary.get('cost') or 0) and float(week_summary.get('installs') or 0):
            insights.append(f"近 7 天 CPI：{week_summary.get('cpi')}。")
    payload = {
        'configured': configured,
        'token_configured': bool(normalized_token),
        'app_ids_configured': len(normalized_app_ids),
        'app_ids': normalized_app_ids,
        'app_ids_source': app_ids_source,
        'meta_token_configured': bool(normalized_meta_token),
        'meta_ad_account_ids_configured': len(normalized_meta_ad_account_ids),
        'meta_ad_account_ids': normalized_meta_ad_account_ids,
        'meta_account_ids_source': meta_account_ids_source,
        'bind_success_token_configured': bool(normalized_bind_success_token),
        'sources': {
            'appsflyer': {
                'configured': appsflyer_configured,
                'token_configured': bool(normalized_token),
                'account_count': len(normalized_app_ids),
                'row_count': len(appsflyer_rows),
                'label': 'AppsFlyer',
                'app_list_status': app_ids_status,
                'app_ids_source': app_ids_source,
                'partner_report': appsflyer_partner_report,
                'raw_event_report': appsflyer_event_report,
                'event_metric_source': appsflyer_event_metric_source,
            },
            'meta': {
                'configured': meta_configured,
                'token_configured': bool(normalized_meta_token),
                'account_count': len(normalized_meta_ad_account_ids),
                'row_count': len(meta_rows),
                'label': 'Meta',
                'report': meta_report,
                'granularity': 'hourly_utc_window' if meta_report.get('hourly_ok') and not meta_report.get('daily_fallback') else ('mixed_hourly_daily_fallback' if meta_report.get('hourly_ok') else 'daily_fallback'),
            },
            'bind_success': {
                'configured': bind_success_configured,
                'token_configured': bool(normalized_bind_success_token),
                'account_count': 1 if bind_success_configured else 0,
                'row_count': len(bind_success_rows),
                'label': 'Bind Success',
                'project': normalized_bind_success_project,
                'base_url_configured': bool(normalized_bind_success_base_url),
                'report': bind_success_report,
            },
            'google': {
                'configured': False,
                'token_configured': False,
                'account_count': 0,
                'row_count': len([row for row in appsflyer_rows if str((row or {}).get('platform') or '').strip().lower() == 'google']),
                'label': 'Google Ads',
            },
            'tiktok': {
                'configured': False,
                'token_configured': False,
                'account_count': 0,
                'row_count': len([row for row in appsflyer_rows if str((row or {}).get('platform') or '').strip().lower() == 'tiktok']),
                'label': 'TikTok',
            },
        },
        'source': 'appsflyer_pull_api,meta_marketing_api,bind_success_events_api',
        'timezone': str(timezone_name or 'UTC'),
        'generated_at': datetime.now(timezone.utc).isoformat(),
        'window_days': window_days,
        'date_start': start_date.isoformat(),
        'date_end': end_date.isoformat(),
        'raw_row_count': len(selected_globally_filtered_rows),
        'fetched_row_count': len(all_rows),
        'row_count': len(rows),
        'applied_filters': {
            'date_from': start_date.isoformat(),
            'date_to': end_date.isoformat(),
            'top_limit': normalized_top_limit,
            'target_app': normalized_target_app,
            **{key: normalized_filters.get(key, []) for key in AD_DASHBOARD_FILTER_KEYS},
            'platform_filters': normalized_platform_filters,
            'platform_date_windows': normalized_platform_date_windows,
        },
        'target_app_options': _ad_dashboard_target_app_options(),
        'filter_options': _ad_filter_options(platform_date_filtered_rows, keys=AD_DASHBOARD_GLOBAL_FILTER_KEYS),
        'platform_filter_options': _ad_platform_filter_options(platform_date_filtered_rows),
        'periods': {
            'selected': selected_summary,
            'day': day_summary,
            'week': week_summary,
            'month': month_summary,
        },
        'daily': daily_rows,
        'period_platform_overview': _ad_platform_period_overview(
            period_rows,
            end_date,
            all_rows=period_all_app_rows,
            target_app=normalized_target_app,
        ),
        'platform_cards': _ad_platform_cards(rows),
        'platform_sections': _ad_platform_sections(rows, start_date, window_days, limit=normalized_top_limit),
        'platform_detail_rows': {
            platform: _ad_detail_rows([
                row for row in rows
                if str((row or {}).get('platform') or '').strip().lower() == platform.lower()
            ], limit=500)
            for platform in AD_DASHBOARD_PLATFORM_NAMES
        },
        'reconciliation': _ad_reconciliation_rows(globally_filtered_rows),
        'event_mappings': AD_DASHBOARD_EVENT_MAPPINGS,
        'top_media_sources': _ad_top_breakdown(rows, 'media_source', limit=normalized_top_limit),
        'top_campaigns': _ad_top_breakdown(rows, 'campaign', limit=normalized_top_limit),
        'insights': insights,
        'errors': errors,
    }
    if include_fact_rows:
        payload['_fact_rows'] = _ad_materialize_fact_rows(all_rows)
    return payload


AD_DASHBOARD_FACT_COLUMNS = tuple(_empty_ad_metrics().keys())


def ensure_ad_dashboard_fact_tables(conn: sqlite3.Connection) -> None:
    metric_columns = ',\n                    '.join(f'{key} REAL NOT NULL DEFAULT 0' for key in AD_DASHBOARD_FACT_COLUMNS)
    conn.execute(f"""
                CREATE TABLE IF NOT EXISTS ad_dashboard_fact_rows (
                    row_id TEXT PRIMARY KEY,
                    date TEXT NOT NULL,
                    data_source TEXT NOT NULL,
                    platform TEXT NOT NULL,
                    app_id TEXT NOT NULL DEFAULT '',
                    appsflyer_app_id TEXT NOT NULL DEFAULT '',
                    target_app TEXT NOT NULL DEFAULT 'inactive',
                    account_id TEXT NOT NULL DEFAULT '',
                    account_name TEXT NOT NULL DEFAULT '',
                    country TEXT NOT NULL DEFAULT '',
                    media_source TEXT NOT NULL DEFAULT '',
                    campaign TEXT NOT NULL DEFAULT '',
                    campaign_id TEXT NOT NULL DEFAULT '',
                    adset_id TEXT NOT NULL DEFAULT '',
                    ad_id TEXT NOT NULL DEFAULT '',
                    ad_group TEXT NOT NULL DEFAULT '',
                    ad TEXT NOT NULL DEFAULT '',
                    source_type TEXT NOT NULL DEFAULT '',
                    row_count INTEGER NOT NULL DEFAULT 0,
                    {metric_columns},
                    payload_json TEXT NOT NULL DEFAULT '{{}}',
                    updated_at TEXT NOT NULL
                )
                """)
    conn.execute("""
                CREATE TABLE IF NOT EXISTS ad_dashboard_sync_state (
                    source TEXT NOT NULL,
                    date TEXT NOT NULL,
                    status TEXT NOT NULL,
                    row_count INTEGER NOT NULL DEFAULT 0,
                    error_message TEXT NOT NULL DEFAULT '',
                    updated_at TEXT NOT NULL,
                    PRIMARY KEY (source, date)
                )
                """)
    conn.execute("CREATE INDEX IF NOT EXISTS idx_ad_dashboard_fact_date_platform ON ad_dashboard_fact_rows(date, platform)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_ad_dashboard_fact_dims ON ad_dashboard_fact_rows(platform, country, app_id, campaign, ad_group, ad)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_ad_dashboard_sync_date ON ad_dashboard_sync_state(date, status)")
    try:
        existing_columns = {str(row[1]) for row in conn.execute("PRAGMA table_info(ad_dashboard_fact_rows)").fetchall()}
        if 'target_app' not in existing_columns:
            conn.execute("ALTER TABLE ad_dashboard_fact_rows ADD COLUMN target_app TEXT NOT NULL DEFAULT 'inactive'")
        lineage_columns_added = False
        for key in ('account_id', 'account_name', 'campaign_id', 'adset_id', 'ad_id'):
            if key not in existing_columns:
                conn.execute(f"ALTER TABLE ad_dashboard_fact_rows ADD COLUMN {key} TEXT NOT NULL DEFAULT ''")
                lineage_columns_added = True
        for key in AD_DASHBOARD_FACT_COLUMNS:
            if key not in existing_columns:
                conn.execute(f"ALTER TABLE ad_dashboard_fact_rows ADD COLUMN {key} REAL NOT NULL DEFAULT 0")
        if lineage_columns_added:
            conn.execute("""
            UPDATE ad_dashboard_fact_rows
            SET account_id = COALESCE(NULLIF(account_id, ''),
                    CASE WHEN json_valid(payload_json) THEN NULLIF(json_extract(payload_json, '$.account_id'), '') END,
                    CASE WHEN json_valid(payload_json) THEN NULLIF(json_extract(payload_json, '$.ad_account_id'), '') END, ''),
                account_name = COALESCE(NULLIF(account_name, ''),
                    CASE WHEN json_valid(payload_json) THEN NULLIF(json_extract(payload_json, '$.account_name'), '') END,
                    NULLIF(app_id, ''), ''),
                campaign_id = COALESCE(NULLIF(campaign_id, ''),
                    CASE WHEN json_valid(payload_json) THEN NULLIF(json_extract(payload_json, '$.campaign_id'), '') END, ''),
                adset_id = COALESCE(NULLIF(adset_id, ''),
                    CASE WHEN json_valid(payload_json) THEN NULLIF(json_extract(payload_json, '$.adset_id'), '') END, ''),
                ad_id = COALESCE(NULLIF(ad_id, ''),
                    CASE WHEN json_valid(payload_json) THEN NULLIF(json_extract(payload_json, '$.ad_id'), '') END, '')
            """)
        def _backfill_target_app(target_app: str, account_ids: set[str], aliases: set[str] | None = None) -> None:
            variants = sorted(
                {str(value).strip() for value in account_ids if str(value).strip()}
                | {f'act_{value}' for value in account_ids if str(value).strip()}
                | {str(value).strip() for value in (aliases or set()) if str(value).strip()}
            )
            if not variants:
                return
            placeholders = ','.join('?' for _ in variants)
            conn.execute(
                f"""
                UPDATE ad_dashboard_fact_rows
                SET target_app = ?
                WHERE app_id IN ({placeholders})
                   OR appsflyer_app_id IN ({placeholders})
                """,
                [target_app, *variants, *variants],
            )

        _backfill_target_app('linky', AD_DASHBOARD_TARGET_APP_LINKY_META_ACCOUNTS, AD_DASHBOARD_TARGET_APP_LINKY_ALIASES)
        _backfill_target_app('timo', AD_DASHBOARD_TARGET_APP_TIMO_META_ACCOUNTS, AD_DASHBOARD_TARGET_APP_TIMO_ALIASES)
        conn.execute(
            """
            UPDATE ad_dashboard_fact_rows
            SET target_app = 'inactive'
            WHERE target_app IS NULL
               OR target_app = ''
               OR target_app NOT IN ('linky', 'timo', 'inactive')
            """
        )
        conn.execute("CREATE INDEX IF NOT EXISTS idx_ad_dashboard_fact_target_app ON ad_dashboard_fact_rows(target_app, platform, date)")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_ad_dashboard_fact_lineage ON ad_dashboard_fact_rows(date, platform, account_id, campaign_id, campaign)")
    except Exception:
        pass


def _ad_dashboard_fact_insert_sql() -> str:
    columns = [
        'row_id', 'date', 'data_source', 'platform', 'app_id', 'appsflyer_app_id',
        'target_app', 'account_id', 'account_name', 'country', 'media_source',
        'campaign', 'campaign_id', 'adset_id', 'ad_id', 'ad_group', 'ad', 'source_type',
        'row_count',
        *AD_DASHBOARD_FACT_COLUMNS,
        'payload_json', 'updated_at',
    ]
    placeholders = ','.join('?' for _ in columns)
    update_columns = [column for column in columns if column != 'row_id']
    update_clause = ','.join(f'{column}=excluded.{column}' for column in update_columns)
    return f"INSERT INTO ad_dashboard_fact_rows ({','.join(columns)}) VALUES ({placeholders}) ON CONFLICT(row_id) DO UPDATE SET {update_clause}"


def upsert_ad_dashboard_fact_rows(
    conn: sqlite3.Connection,
    rows: List[Dict[str, Any]],
    *,
    synced_at: Optional[str] = None,
) -> int:
    ensure_ad_dashboard_fact_tables(conn)
    now = synced_at or datetime.now(timezone.utc).isoformat()
    materialized = _ad_materialize_fact_rows(rows)
    sql = _ad_dashboard_fact_insert_sql()
    count = 0
    for row in materialized:
        stored = {key: row.get(key) for key in [
            'date', 'data_source', 'platform', 'app_id', 'appsflyer_app_id',
            'target_app', 'account_id', 'account_name', 'country', 'media_source',
            'campaign', 'campaign_id', 'adset_id', 'ad_id', 'ad_group', 'ad', 'source_type',
        ]}
        stored['account_id'] = row.get('account_id') or row.get('ad_account_id') or ''
        stored['account_name'] = row.get('account_name') or row.get('app_id') or ''
        stored['campaign_id'] = row.get('campaign_id') or ''
        stored['adset_id'] = row.get('adset_id') or ''
        stored['ad_id'] = row.get('ad_id') or ''
        stored['external_app'] = row.get('external_app')
        stored['target_app'] = _ad_dashboard_row_target_app(stored)
        stored['row_count'] = int(row.get('row_count') or 0)
        for key in AD_DASHBOARD_FACT_COLUMNS:
            stored[key] = float(row.get(key) or 0.0)
        payload = {key: value for key, value in row.items() if key not in {'payload_json'}}
        values = [
            _ad_fact_row_id(row),
            stored['date'],
            stored['data_source'],
            stored['platform'],
            stored['app_id'],
            stored['appsflyer_app_id'],
            stored['target_app'],
            stored['account_id'],
            stored['account_name'],
            stored['country'],
            stored['media_source'],
            stored['campaign'],
            stored['campaign_id'],
            stored['adset_id'],
            stored['ad_id'],
            stored['ad_group'],
            stored['ad'],
            stored['source_type'],
            stored['row_count'],
            *[stored[key] for key in AD_DASHBOARD_FACT_COLUMNS],
            json.dumps(payload, ensure_ascii=False, sort_keys=True),
            now,
        ]
        conn.execute(sql, values)
        count += 1
    return count


def upsert_ad_creative_performance_daily_rows(
    conn: sqlite3.Connection,
    rows: List[Dict[str, Any]],
) -> int:
    ensure_creative_intelligence_tables(conn)
    sql = """
        INSERT INTO ad_creative_performance_daily (
            report_date_london, asset_id, creative_id, ad_id, adset_id, campaign_id, country, project,
            spend, impressions, clicks, ctr, cpm, installs, cpi, af_model_join_events, tugao_real_bind_count,
            real_bind_cpa, af_to_real_bind_rate, data_quality_status, attribution_level, creative_grain,
            is_dynamic_creative, grain_warning
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(report_date_london, asset_id, ad_id) DO UPDATE SET
            creative_id=excluded.creative_id,
            adset_id=excluded.adset_id,
            campaign_id=excluded.campaign_id,
            country=excluded.country,
            project=excluded.project,
            spend=excluded.spend,
            impressions=excluded.impressions,
            clicks=excluded.clicks,
            ctr=excluded.ctr,
            cpm=excluded.cpm,
            installs=excluded.installs,
            cpi=excluded.cpi,
            af_model_join_events=excluded.af_model_join_events,
            tugao_real_bind_count=excluded.tugao_real_bind_count,
            real_bind_cpa=excluded.real_bind_cpa,
            af_to_real_bind_rate=excluded.af_to_real_bind_rate,
            data_quality_status=excluded.data_quality_status,
            attribution_level=excluded.attribution_level,
            creative_grain=excluded.creative_grain,
            is_dynamic_creative=excluded.is_dynamic_creative,
            grain_warning=excluded.grain_warning
    """
    count = 0
    for row in rows or []:
        report_date = str(row.get('report_date_london') or row.get('date') or '').strip()
        ad_id = str(row.get('ad_id') or '').strip()
        if not report_date or not ad_id:
            continue
        spend = float(row.get('spend') if row.get('spend') is not None else row.get('cost') or 0.0)
        impressions = float(row.get('impressions') or 0.0)
        clicks = float(row.get('clicks') or 0.0)
        installs = float(row.get('installs') or row.get('meta_installs') or 0.0)
        real_binds = int(float(row.get('tugao_real_bind_count') or row.get('real_bind_count') or 0))
        af_joins = float(row.get('af_model_join_events') or row.get('af_guild_joins') or 0.0)
        asset_id = str(row.get('asset_id') or '').strip() or f"meta_ad_{hashlib.sha1(ad_id.encode('utf-8')).hexdigest()[:16]}"
        dynamic = bool(int(row.get('is_dynamic_creative') or 0)) if str(row.get('is_dynamic_creative') or '').strip() else False
        grain = str(row.get('creative_grain') or ('dynamic' if dynamic else 'ad')).strip() or 'ad'
        conn.execute(sql, (
            report_date,
            asset_id,
            str(row.get('creative_id') or '').strip(),
            ad_id,
            str(row.get('adset_id') or '').strip(),
            str(row.get('campaign_id') or '').strip(),
            str(row.get('country') or '').strip(),
            str(row.get('project') or '').strip(),
            round(spend, 4),
            round(impressions, 4),
            round(clicks, 4),
            round(float(row.get('ctr') or (clicks / impressions if impressions else 0.0)), 6),
            round(float(row.get('cpm') or (spend / impressions * 1000 if impressions else 0.0)), 4),
            round(installs, 4),
            round(spend / installs, 4) if installs else None,
            round(af_joins, 4),
            real_binds,
            round(spend / real_binds, 4) if real_binds else None,
            round(real_binds / af_joins, 4) if af_joins else None,
            str(row.get('data_quality_status') or 'media_only_ad_id').strip() or 'media_only_ad_id',
            str(row.get('attribution_level') or '广告级').strip() or '广告级',
            grain,
            1 if dynamic else 0,
            str(row.get('grain_warning') or '当前为 Meta 广告级投放表现；真实入会等后链路需等待 ad_id 归因接入。').strip(),
        ))
        count += 1
    return count


def _creative_perf_match_text(value: Any) -> str:
    return re.sub(r'\s+', ' ', str(value or '').strip()).casefold()


def _creative_perf_fact_key(row: Dict[str, Any]) -> Tuple[str, str, str, str]:
    return (
        _creative_perf_match_text(row.get('date')),
        _creative_perf_match_text(row.get('campaign')),
        _creative_perf_match_text(row.get('ad_group')),
        _creative_perf_match_text(row.get('ad')),
    )


def _creative_perf_account_matches(meta: Dict[str, Any], fact: Dict[str, Any]) -> bool:
    meta_account = normalize_meta_ad_account_id(meta.get('account_id'))
    fact_account = normalize_meta_ad_account_id(
        fact.get('account_id')
        or fact.get('ad_account_id')
        or fact.get('app_id')
        or fact.get('appsflyer_app_id')
    )
    return not meta_account or not fact_account or meta_account == fact_account


def _creative_perf_fact_rows_for_meta(
    conn: sqlite3.Connection,
    meta_rows: List[Dict[str, Any]],
) -> Dict[Tuple[str, str, str, str], List[Dict[str, Any]]]:
    dates = sorted({str(row.get('date') or '').strip() for row in meta_rows or [] if str(row.get('date') or '').strip()})
    if not dates:
        return {}
    rows = conn.execute(
        """
        SELECT date, data_source, platform, app_id, appsflyer_app_id, target_app, country,
               campaign, ad_group, ad, af_guild_joins, guild_joins, bind_success_users,
               onsite_registrations, high_value_users, im_entries, payload_json
        FROM ad_dashboard_fact_rows
        WHERE date >= ? AND date <= ?
          AND platform = 'Meta'
          AND data_source IN ('AppsFlyer', 'TugaoFunnel')
        """,
        (min(dates), max(dates)),
    ).fetchall()
    by_key: Dict[Tuple[str, str, str, str], List[Dict[str, Any]]] = {}
    for row in rows:
        item = dict(row)
        try:
            payload = json.loads(str(item.pop('payload_json') or '{}'))
        except Exception:
            payload = {}
        if isinstance(payload, dict):
            item['account_id'] = payload.get('account_id') or payload.get('ad_account_id') or item.get('account_id') or ''
            item['ad_account_id'] = payload.get('ad_account_id') or payload.get('account_id') or ''
        by_key.setdefault(_creative_perf_fact_key(item), []).append(item)
    return by_key


def _creative_perf_downstream_for_meta_row(
    meta: Dict[str, Any],
    fact_rows_by_key: Dict[Tuple[str, str, str, str], List[Dict[str, Any]]],
) -> Dict[str, Any]:
    candidates = [
        row
        for row in fact_rows_by_key.get(_creative_perf_fact_key(meta), [])
        if _creative_perf_account_matches(meta, row)
    ]
    by_source: Dict[str, List[Dict[str, Any]]] = {}
    for row in candidates:
        by_source.setdefault(str(row.get('data_source') or ''), []).append(row)
    result: Dict[str, Any] = {
        'af_model_join_events': 0.0,
        'tugao_real_bind_count': 0,
        'data_quality_status': 'media_only_ad_id',
        'attribution_warning': '当前仅有 Meta ad_id 媒体指标，AF/Tugao 后链路未匹配到唯一广告文本行。',
    }
    af_rows = by_source.get('AppsFlyer') or []
    tugao_rows = by_source.get('TugaoFunnel') or []
    ambiguous_sources: List[str] = []
    if len(af_rows) == 1:
        result['af_model_join_events'] = float(af_rows[0].get('af_guild_joins') or 0.0)
    elif len(af_rows) > 1:
        ambiguous_sources.append('AppsFlyer')
    if len(tugao_rows) == 1:
        result['tugao_real_bind_count'] = int(float(tugao_rows[0].get('guild_joins') or tugao_rows[0].get('bind_success_users') or 0.0))
    elif len(tugao_rows) > 1:
        ambiguous_sources.append('TugaoFunnel')
    if ambiguous_sources:
        result['data_quality_status'] = 'ambiguous_downstream_text_match'
        result['attribution_warning'] = f"AF/Tugao 文本归因匹配到多行：{', '.join(ambiguous_sources)}，未写入歧义来源。"
    elif af_rows or tugao_rows:
        result['data_quality_status'] = 'ad_id_with_downstream_text_match'
        result['attribution_warning'] = 'Meta ad_id 媒体指标已按日期/账户/campaign/ad_group/ad 文本唯一匹配 AF/Tugao 后链路。'
    return result


def build_ad_creative_performance_rows_from_meta_rows(
    conn: sqlite3.Connection,
    meta_rows: List[Dict[str, Any]],
) -> List[Dict[str, Any]]:
    ensure_creative_intelligence_tables(conn)
    ad_ids = sorted({str(row.get('ad_id') or '').strip() for row in meta_rows or [] if str(row.get('ad_id') or '').strip()})
    asset_by_ad: Dict[str, sqlite3.Row] = {}
    if ad_ids:
        placeholders = ','.join('?' for _ in ad_ids)
        for row in conn.execute(
            f"""
            SELECT asset_id, ad_id, creative_id, adset_id, campaign_id, country, project, is_dynamic_creative
            FROM ad_creative_asset
            WHERE ad_id IN ({placeholders})
            """,
            tuple(ad_ids),
        ).fetchall():
            asset_by_ad[str(row['ad_id'] or '')] = row
    downstream_by_key = _creative_perf_fact_rows_for_meta(conn, meta_rows)
    rows: List[Dict[str, Any]] = []
    for meta in meta_rows or []:
        ad_id = str(meta.get('ad_id') or '').strip()
        if not ad_id:
            continue
        asset = asset_by_ad.get(ad_id)
        is_dynamic = bool(asset and int(asset['is_dynamic_creative'] or 0))
        downstream = _creative_perf_downstream_for_meta_row(meta, downstream_by_key)
        rows.append({
            'date': meta.get('date'),
            'asset_id': (asset['asset_id'] if asset else '') or f"meta_ad_{hashlib.sha1(ad_id.encode('utf-8')).hexdigest()[:16]}",
            'creative_id': (asset['creative_id'] if asset else '') or str(meta.get('creative_id') or ''),
            'ad_id': ad_id,
            'adset_id': (asset['adset_id'] if asset else '') or str(meta.get('adset_id') or ''),
            'campaign_id': (asset['campaign_id'] if asset else '') or str(meta.get('campaign_id') or ''),
            'country': (asset['country'] if asset else '') or str(meta.get('country') or ''),
            'project': (asset['project'] if asset else '') or str(meta.get('project') or meta.get('target_app') or ''),
            'spend': meta.get('cost') or 0.0,
            'impressions': meta.get('impressions') or 0.0,
            'clicks': meta.get('clicks') or 0.0,
            'ctr': meta.get('ctr') or 0.0,
            'cpm': meta.get('cpm') or 0.0,
            'installs': meta.get('meta_installs') or meta.get('installs') or 0.0,
            'af_model_join_events': downstream.get('af_model_join_events') or 0.0,
            'tugao_real_bind_count': downstream.get('tugao_real_bind_count') or 0,
            'data_quality_status': downstream.get('data_quality_status') or 'media_only_ad_id',
            'creative_grain': 'dynamic' if is_dynamic else 'ad',
            'is_dynamic_creative': 1 if is_dynamic else 0,
            'grain_warning': downstream.get('attribution_warning') or '',
        })
    return rows


def replace_ad_dashboard_fact_rows_for_dates(
    conn: sqlite3.Connection,
    rows: List[Dict[str, Any]],
    *,
    start_date: datetime.date,
    end_date: datetime.date,
    synced_at: Optional[str] = None,
) -> int:
    """Merge facts while preserving settled rows from inaccessible accounts."""
    ensure_ad_dashboard_fact_tables(conn)
    return upsert_ad_dashboard_fact_rows(conn, rows, synced_at=synced_at)


def mark_ad_dashboard_sync_state(
    conn: sqlite3.Connection,
    *,
    source: str,
    start_date: datetime.date,
    end_date: datetime.date,
    status: str,
    row_count: int = 0,
    error_message: str = '',
    synced_at: Optional[str] = None,
) -> None:
    ensure_ad_dashboard_fact_tables(conn)
    now = synced_at or datetime.now(timezone.utc).isoformat()
    cursor = start_date
    while cursor <= end_date:
        conn.execute(
            """
            INSERT INTO ad_dashboard_sync_state(source, date, status, row_count, error_message, updated_at)
            VALUES(?, ?, ?, ?, ?, ?)
            ON CONFLICT(source, date) DO UPDATE SET
                status=excluded.status,
                row_count=excluded.row_count,
                error_message=excluded.error_message,
                updated_at=excluded.updated_at
            """,
            (source, cursor.isoformat(), status, int(row_count or 0), str(error_message or ''), now),
        )
        cursor += timedelta(days=1)


def read_ad_dashboard_fact_rows(
    conn: sqlite3.Connection,
    *,
    start_date: datetime.date,
    end_date: datetime.date,
) -> List[Dict[str, Any]]:
    ensure_ad_dashboard_fact_tables(conn)
    columns = [
        'date', 'data_source', 'platform', 'app_id', 'appsflyer_app_id',
        'target_app', 'account_id', 'account_name', 'country', 'media_source',
        'campaign', 'campaign_id', 'adset_id', 'ad_id', 'ad_group', 'ad', 'source_type',
        'row_count',
        *AD_DASHBOARD_FACT_COLUMNS,
        'payload_json',
    ]
    result: List[Dict[str, Any]] = []
    for row in conn.execute(
        f"SELECT {','.join(columns)} FROM ad_dashboard_fact_rows WHERE date >= ? AND date <= ? ORDER BY date, platform, cost DESC",
        (start_date.isoformat(), end_date.isoformat()),
    ).fetchall():
        item = dict(row)
        payload: Dict[str, Any] = {}
        try:
            payload = json.loads(str(item.get('payload_json') or '{}'))
        except Exception:
            payload = {}
        if isinstance(payload, dict):
            item['account_id'] = payload.get('account_id') or payload.get('ad_account_id') or item.get('account_id') or ''
            item['account_name'] = item.get('account_name') or payload.get('account_name') or item.get('app_id') or ''
            item['campaign_id'] = item.get('campaign_id') or payload.get('campaign_id') or ''
            item['adset_id'] = item.get('adset_id') or payload.get('adset_id') or ''
            item['ad_id'] = item.get('ad_id') or payload.get('ad_id') or ''
            item['ad_account_id'] = payload.get('ad_account_id') or payload.get('account_id') or item.get('ad_account_id') or ''
            item['external_app'] = payload.get('external_app') or item.get('external_app') or ''
            for metadata_key in (
                'historical_recovery',
                'historical_source_created_at',
                'country_attribution_status',
                'country_attribution_source',
                'country_attribution_grain',
            ):
                if metadata_key in payload:
                    item[metadata_key] = payload.get(metadata_key)
        for key in AD_DASHBOARD_FACT_COLUMNS:
            item[key] = float(item.get(key) or 0.0)
        item['row_count'] = int(item.get('row_count') or 0)
        item['app_id'] = _normalize_ad_fact_account_value(item)
        item['target_app'] = _ad_dashboard_row_target_app(item)
        item.pop('payload_json', None)
        result.append(item)
    return _ad_enrich_unknown_countries(
        _ad_enrich_countries_from_meta_delivery(result)
    )


def ad_dashboard_fact_dates_complete(
    conn: sqlite3.Connection,
    *,
    start_date: datetime.date,
    end_date: datetime.date,
) -> bool:
    ensure_ad_dashboard_fact_tables(conn)
    expected = max((end_date - start_date).days + 1, 0)
    if expected <= 0:
        return False
    rows = conn.execute(
        "SELECT COUNT(DISTINCT date) FROM ad_dashboard_fact_rows WHERE date >= ? AND date <= ?",
        (start_date.isoformat(), end_date.isoformat()),
    ).fetchone()
    return int(rows[0] if rows else 0) >= expected


def ad_dashboard_fact_rows_completeness(
    rows: List[Dict[str, Any]],
    *,
    start_date: datetime.date,
    end_date: datetime.date,
    appsflyer_required: bool = True,
) -> Dict[str, Any]:
    expected_dates = []
    cursor = start_date
    while cursor <= end_date:
        expected_dates.append(cursor.isoformat())
        cursor += timedelta(days=1)
    rows_by_date: Dict[str, List[Dict[str, Any]]] = {day: [] for day in expected_dates}
    for row in rows or []:
        row_date = _parse_dashboard_date((row or {}).get('date'))
        if row_date and start_date <= row_date <= end_date:
            rows_by_date.setdefault(row_date.isoformat(), []).append(row)
    missing_dates = [day for day in expected_dates if not rows_by_date.get(day)]
    missing_appsflyer: List[Dict[str, str]] = []
    unresolved_meta_country: List[Dict[str, str]] = []
    missing_meta_lineage: List[Dict[str, str]] = []
    for day in expected_dates:
        for row in rows_by_date.get(day) or []:
            if str((row or {}).get('data_source') or '').strip().lower() != 'meta':
                continue
            has_delivery = any(
                float((row or {}).get(key) or 0.0)
                for key in ('cost', 'impressions', 'clicks', 'link_clicks', 'meta_installs', 'installs')
            )
            account_id = str((row or {}).get('account_id') or (row or {}).get('ad_account_id') or '').strip()
            campaign = str((row or {}).get('campaign') or '').strip()
            campaign_id = str((row or {}).get('campaign_id') or '').strip()
            if has_delivery and (not account_id or not campaign or not campaign_id):
                missing_meta_lineage.append({
                    'date': day,
                    'account_id': account_id,
                    'campaign': campaign,
                    'campaign_id': campaign_id,
                })
            status = str((row or {}).get('country_attribution_status') or '').strip()
            if status not in {'unresolved_waiting_meta_delivery_country', 'meta_delivery_country_ambiguous'}:
                continue
            if not _ad_country_is_unknown((row or {}).get('country')):
                continue
            if not any(float((row or {}).get(key) or 0.0) for key in ('cost', 'impressions', 'clicks', 'link_clicks')):
                continue
            unresolved_meta_country.append({'date': day, 'account_id': str((row or {}).get('account_id') or (row or {}).get('ad_account_id') or ''), 'campaign': str((row or {}).get('campaign') or '')})
    if appsflyer_required:
        media_sources = {'meta', 'google', 'tiktok'}
        for day in expected_dates:
            day_rows = rows_by_date.get(day) or []
            platforms_with_media = {
                str((row or {}).get('platform') or '').strip()
                for row in day_rows
                if str((row or {}).get('platform') or '').strip().lower() != 'internal'
                and str((row or {}).get('data_source') or '').strip().lower() in media_sources
                and (
                    float((row or {}).get('cost') or 0.0)
                    or float((row or {}).get('meta_installs') or 0.0)
                    or float((row or {}).get('installs') or 0.0)
                )
            }
            platforms_with_appsflyer = {
                str((row or {}).get('platform') or '').strip()
                for row in day_rows
                if str((row or {}).get('data_source') or '').strip().lower() == 'appsflyer'
            }
            for platform in sorted(platforms_with_media - platforms_with_appsflyer):
                missing_appsflyer.append({'date': day, 'platform': platform})
    complete = not missing_dates and not missing_appsflyer and not unresolved_meta_country and not missing_meta_lineage
    reason_parts = []
    if missing_dates:
        reason_parts.append('missing_dates=' + ','.join(missing_dates[:5]))
    if missing_appsflyer:
        sample = ','.join(f"{item['date']}:{item['platform']}" for item in missing_appsflyer[:5])
        reason_parts.append('missing_appsflyer=' + sample)
    if unresolved_meta_country:
        sample = ','.join(f"{item['date']}:{item['account_id']}:{item['campaign']}" for item in unresolved_meta_country[:5])
        reason_parts.append('unresolved_meta_country=' + sample)
    if missing_meta_lineage:
        sample = ','.join(
            f"{item['date']}:{item['account_id'] or '-'}:{item['campaign_id'] or item['campaign'] or '-'}"
            for item in missing_meta_lineage[:5]
        )
        reason_parts.append('missing_meta_lineage=' + sample)
    return {
        'complete': complete,
        'missing_dates': missing_dates,
        'missing_appsflyer': missing_appsflyer,
        'unresolved_meta_country': unresolved_meta_country,
        'missing_meta_lineage': missing_meta_lineage,
        'status': 'ok' if complete else 'partial',
        'error_message': '; '.join(reason_parts),
    }


def ad_dashboard_sync_error_user_message(error_message: str) -> str:
    text = str(error_message or '').strip()
    if not text:
        return '本地事实表存在缺口，已按本地事实表展示；请后台补齐后刷新。'
    if 'missing_dates=' in text:
        raw = text.split('missing_dates=', 1)[1].split(';', 1)[0].strip()
        dates = [item.strip() for item in raw.split(',') if item.strip()]
        if dates:
            return f"本地事实表缺少 {', '.join(dates)}，已跳过不完整缓存并尝试实时读取；后台补数任务会自动回填。"
    if 'missing_appsflyer=' not in text:
        return text
    raw = text.split('missing_appsflyer=', 1)[1].split(';', 1)[0].strip()
    items = [item.strip() for item in raw.split(',') if item.strip()]
    platforms = []
    dates = []
    for item in items:
        if ':' in item:
            day, platform = item.split(':', 1)
            dates.append(day.strip())
            platforms.append(platform.strip())
        else:
            platforms.append(item)
    unique_platforms = ', '.join(sorted({platform for platform in platforms if platform}))
    unique_dates = ', '.join(sorted({day for day in dates if day}))
    if unique_platforms and unique_dates:
        return f'上游 AppsFlyer 暂未返回 {unique_dates} 的 {unique_platforms} 归因数据；系统已显示 Meta 和 Tugao 本地数据，待后台自动重试补齐。'
    if unique_platforms:
        return f'上游 AppsFlyer 暂未返回 {unique_platforms} 归因数据；系统已显示已有本地数据，待后台自动重试补齐。'
    return '上游 AppsFlyer 暂未返回归因数据；系统已显示已有本地数据，待后台自动重试补齐。'


def build_ad_data_dashboard_snapshot_from_rows(
    rows: List[Dict[str, Any]],
    *,
    timezone_name: str = 'UTC',
    days: int = 30,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    filters: Optional[Dict[str, Any]] = None,
    platform_filters: Optional[Dict[str, Any]] = None,
    platform_date_windows: Optional[Dict[str, Any]] = None,
    target_app: str = 'all',
    top_limit: int = 8,
    sources: Optional[Dict[str, Any]] = None,
    insights: Optional[List[str]] = None,
    errors: Optional[List[Dict[str, str]]] = None,
) -> Dict[str, Any]:
    normalized_filters = _normalize_ad_filters(filters)
    normalized_platform_filters = _normalize_ad_platform_filters(platform_filters)
    normalized_target_app = _normalize_ad_dashboard_target_app(target_app) or 'all'
    normalized_top_limit = min(max(int(top_limit or 8), 3), 25)
    today = _ad_dashboard_latest_complete_utc_date()
    start_date, end_date, window_days, window_warnings = _coerce_ad_dashboard_window(
        days=days,
        date_from=date_from,
        date_to=date_to,
        today=today,
    )
    normalized_platform_date_windows, platform_start_date, platform_end_date, platform_window_days, platform_window_warnings = _normalize_ad_platform_date_windows(
        platform_date_windows,
        fallback_start_date=start_date,
        fallback_end_date=end_date,
        today=today,
    )
    window_warnings.extend(platform_window_warnings)
    start_date = platform_start_date
    end_date = platform_end_date
    window_days = platform_window_days
    country_enriched_rows = _ad_enrich_countries_from_meta_delivery(list(rows or []))
    all_app_rows = _ad_dashboard_enrich_account_from_peer_rows(
        _ad_dashboard_apply_target_app_filter(country_enriched_rows, 'all')
    )
    all_rows = _ad_dashboard_enrich_account_from_peer_rows(
        _ad_dashboard_apply_target_app_filter(country_enriched_rows, normalized_target_app)
    )
    globally_filtered_rows = _filter_ad_rows(all_rows, normalized_filters)
    all_app_globally_filtered_rows = _filter_ad_rows(all_app_rows, normalized_filters)
    selected_globally_filtered_rows = _ad_rows_in_date_range(globally_filtered_rows, start_date, end_date)
    platform_date_filtered_rows = _filter_ad_rows_by_platform_date_windows(selected_globally_filtered_rows, normalized_platform_date_windows)
    visible_rows = _filter_ad_rows_by_platform_filters(platform_date_filtered_rows, normalized_platform_filters)
    # Platform filters are detail-panel filters. The top day/week/month overview
    # stays scoped only by global date/app filters so drilling into Meta country
    # or account does not rewrite the dashboard-level totals.
    period_rows = globally_filtered_rows
    period_all_app_rows = all_app_globally_filtered_rows
    fact_dates = {
        str((row or {}).get('date') or '').strip()
        for row in visible_rows
        if str((row or {}).get('date') or '').strip()
    }
    daily_rows = []
    for offset in range(window_days):
        day = start_date + timedelta(days=offset)
        daily_rows.append({
            'date': day.isoformat(),
            'has_fact_data': day.isoformat() in fact_dates,
            **_ad_period_summary(visible_rows, day, day),
        })
    day_summary = _ad_period_summary(period_rows, end_date, end_date)
    week_summary = _ad_period_summary(period_rows, end_date - timedelta(days=6), end_date)
    month_summary = _ad_period_summary(period_rows, end_date - timedelta(days=29), end_date)
    selected_summary = _ad_period_summary(visible_rows, start_date, end_date)
    source_rows = Counter(str((row or {}).get('data_source') or '').strip().lower() for row in all_rows)
    default_sources = {
        'appsflyer': {'configured': True, 'row_count': int(source_rows.get('appsflyer') or 0), 'label': 'AppsFlyer'},
        'meta': {'configured': True, 'row_count': int(source_rows.get('meta') or 0), 'label': 'Meta', 'granularity': 'local_fact_rows'},
        'bind_success': {'configured': True, 'row_count': int(source_rows.get('bindsuccess') or 0), 'label': 'Bind Success'},
        'google': {'configured': False, 'row_count': len([row for row in all_rows if str((row or {}).get('platform') or '').strip().lower() == 'google']), 'label': 'Google Ads'},
        'tiktok': {'configured': False, 'row_count': len([row for row in all_rows if str((row or {}).get('platform') or '').strip().lower() == 'tiktok']), 'label': 'TikTok'},
    }
    source_payload = {**default_sources, **(sources or {})}
    insight_rows = [*window_warnings, *(insights or [])]
    if not visible_rows and all_rows:
        insight_rows.append('当前筛选条件下暂无本地明细数据，请放宽筛选或手动刷新补数。')
    return {
        'configured': bool(all_rows),
        'source': 'local_ad_dashboard_fact_rows',
        'timezone': str(timezone_name or 'UTC'),
        'generated_at': datetime.now(timezone.utc).isoformat(),
        'window_days': window_days,
        'date_start': start_date.isoformat(),
        'date_end': end_date.isoformat(),
        'raw_row_count': len(selected_globally_filtered_rows),
        'fetched_row_count': len(all_rows),
        'row_count': len(visible_rows),
        'sources': source_payload,
        'applied_filters': {
            'date_from': start_date.isoformat(),
            'date_to': end_date.isoformat(),
            'top_limit': normalized_top_limit,
            'target_app': normalized_target_app,
            **{key: normalized_filters.get(key, []) for key in AD_DASHBOARD_FILTER_KEYS},
            'platform_filters': normalized_platform_filters,
            'platform_date_windows': normalized_platform_date_windows,
        },
        'target_app_options': _ad_dashboard_target_app_options(),
        'filter_options': _ad_filter_options(platform_date_filtered_rows, keys=AD_DASHBOARD_GLOBAL_FILTER_KEYS),
        'platform_filter_options': _ad_platform_filter_options(platform_date_filtered_rows),
        'periods': {'selected': selected_summary, 'day': day_summary, 'week': week_summary, 'month': month_summary},
        'daily': daily_rows,
        'period_platform_overview': _ad_platform_period_overview(
            period_rows,
            end_date,
            all_rows=period_all_app_rows,
            target_app=normalized_target_app,
        ),
        'platform_cards': _ad_platform_cards(visible_rows),
        'platform_sections': _ad_platform_sections(visible_rows, start_date, window_days, limit=normalized_top_limit),
        'platform_detail_rows': {
            platform: _ad_detail_rows([
                row for row in visible_rows
                if str((row or {}).get('platform') or '').strip().lower() == platform.lower()
            ], limit=500)
            for platform in AD_DASHBOARD_PLATFORM_NAMES
        },
        'reconciliation': _ad_reconciliation_rows(globally_filtered_rows),
        'event_mappings': AD_DASHBOARD_EVENT_MAPPINGS,
        'top_media_sources': _ad_top_breakdown(visible_rows, 'media_source', limit=normalized_top_limit),
        'top_campaigns': _ad_top_breakdown(visible_rows, 'campaign', limit=normalized_top_limit),
        'insights': insight_rows,
        'errors': errors or [],
    }


def ad_dashboard_fact_window_for_context(context: Dict[str, Any]) -> Tuple[datetime.date, datetime.date]:
    today = _ad_dashboard_latest_complete_utc_date()
    start_date, end_date, _, _ = _coerce_ad_dashboard_window(
        days=int((context or {}).get('days') or 30),
        date_from=(context or {}).get('date_from'),
        date_to=(context or {}).get('date_to'),
        today=today,
    )
    _, platform_start_date, platform_end_date, _, _ = _normalize_ad_platform_date_windows(
        (context or {}).get('platform_date_windows') or {},
        fallback_start_date=start_date,
        fallback_end_date=end_date,
        today=today,
    )
    start_date = platform_start_date
    end_date = platform_end_date
    return min(start_date, end_date - timedelta(days=29)), end_date


def normalize_country_label(value: Any) -> str:
    raw = str(value or '').strip()
    if not raw:
        return ''
    return COUNTRY_LABEL_ALIASES.get(raw.lower(), raw)


def infer_country_context(*values: Any) -> str:
    direct_candidates = [normalize_country_label(value) for value in values if str(value or '').strip()]
    for candidate in direct_candidates:
        if candidate in PHONE_LOCALIZED_NUMBER_RULES or candidate in set(PHONE_PREFIX_COUNTRY_MAP.values()):
            return candidate
    inferred = _ad_country_label(*values)
    return '' if inferred == 'Unknown' else normalize_country_label(inferred)


def countries_match(user_country: Any, guild_country: Any) -> bool:
    normalized_user = normalize_country_label(user_country).lower()
    raw_guild_country = (
        ','.join(str(item or '').strip() for item in guild_country if str(item or '').strip())
        if isinstance(guild_country, (list, tuple, set))
        else str(guild_country or '').strip()
    )
    if not normalized_user or not raw_guild_country:
        return True
    guild_country_options = [
        normalize_country_label(item).lower()
        for item in re.split(r'[,，;；/|、]+', raw_guild_country)
        if str(item or '').strip()
    ]
    if not guild_country_options:
        normalized_guild = normalize_country_label(guild_country).lower()
        return not normalized_guild or normalized_user == normalized_guild
    return normalized_user in guild_country_options


SPANISH_LATAM_COMPAT_COUNTRIES: Tuple[str, ...] = ('Mexico', 'Colombia', 'Venezuela', 'Chile')


def normalize_country_options(value: Any) -> List[str]:
    if isinstance(value, (list, tuple, set)):
        raw_items = list(value)
    else:
        raw = str(value or '').strip()
        if not raw:
            raw_items = []
        elif raw.startswith('['):
            try:
                decoded = json.loads(raw)
            except (TypeError, ValueError, json.JSONDecodeError):
                decoded = raw
            raw_items = list(decoded) if isinstance(decoded, list) else re.split(r'[,，;；/|、]+', str(decoded))
        else:
            raw_items = re.split(r'[,，;；/|、]+', raw)
    normalized: List[str] = []
    seen: set[str] = set()
    for item in raw_items:
        country = normalize_country_label(item)
        key = country.casefold()
        if not country or key in seen:
            continue
        seen.add(key)
        normalized.append(country)
    return normalized


def guild_country_contract(row: Optional[Dict[str, Any]]) -> Dict[str, Any]:
    source = dict(row or {})
    legacy_country = str(source.get('country') or '').strip()
    legacy_options = normalize_country_options(legacy_country)
    guild_country = normalize_country_label(source.get('guild_country'))
    if not guild_country:
        guild_country = 'Mexico' if set(legacy_options) == set(SPANISH_LATAM_COMPAT_COUNTRIES) else normalize_country_label(legacy_country)
    eligible = normalize_country_options(source.get('eligible_user_countries'))
    if not eligible:
        eligible = legacy_options or ([guild_country] if guild_country else [])
    if guild_country and guild_country not in eligible:
        eligible.insert(0, guild_country)
    routing_region = str(source.get('routing_region') or '').strip()
    if not routing_region and len(eligible) > 1 and set(eligible).issubset(set(SPANISH_LATAM_COMPAT_COUNTRIES)):
        routing_region = 'ES_LATAM'
    return {
        'guild_country': guild_country,
        'eligible_user_countries': eligible,
        'routing_region': routing_region,
    }


def _phone_localization_rule(*, country: Any = '', area_code: Any = 0) -> Optional[Dict[str, Any]]:
    normalized_country = normalize_country_label(country)
    normalized_area_code = str(area_code or '').strip()
    if not normalized_country and normalized_area_code:
        normalized_country = PHONE_PREFIX_COUNTRY_MAP.get(normalized_area_code, '')
    if normalized_country in PHONE_LOCALIZED_NUMBER_RULES:
        return PHONE_LOCALIZED_NUMBER_RULES[normalized_country]
    for rule_country, rule in PHONE_LOCALIZED_NUMBER_RULES.items():
        if normalized_area_code and normalized_area_code == str(rule.get('country_code') or ''):
            return rule
        if normalized_country and normalized_country == rule_country:
            return rule
    return None


def _localized_phone_body(digits: str, rule: Optional[Dict[str, Any]]) -> str:
    body = ''.join(ch for ch in str(digits or '') if ch.isdigit())
    if not body or not rule:
        return body
    country_code = str(rule.get('country_code') or '').strip()
    national_lengths = {int(item) for item in (rule.get('national_lengths') or set())}
    local_lengths = {int(item) for item in (rule.get('local_lengths') or set())}
    trunk_prefixes = tuple(str(item) for item in (rule.get('trunk_prefixes') or ()))
    leading_digits = tuple(str(item) for item in (rule.get('leading_digits') or ()))

    def acceptable(value: str) -> bool:
        if national_lengths and len(value) not in national_lengths:
            return False
        if leading_digits and not value.startswith(leading_digits):
            return False
        return True

    if country_code and body.startswith(country_code) and len(body) > len(country_code):
        candidate = body[len(country_code):]
        if acceptable(candidate):
            return candidate
    for trunk in trunk_prefixes:
        if body.startswith(trunk) and len(body) in local_lengths:
            candidate = body[len(trunk):]
            if acceptable(candidate):
                return candidate
    return body


def _localized_phone_body_matches_rule(body: str, rule: Optional[Dict[str, Any]]) -> bool:
    normalized = ''.join(ch for ch in str(body or '') if ch.isdigit())
    if not normalized or not rule:
        return False
    national_lengths = {int(item) for item in (rule.get('national_lengths') or set())}
    leading_digits = tuple(str(item) for item in (rule.get('leading_digits') or ()))
    if national_lengths and len(normalized) not in national_lengths:
        return False
    if leading_digits and not normalized.startswith(leading_digits):
        return False
    return True


def localized_phone_match_keys(*, phone: Any, area_code: Any = 0, country: Any = '') -> set[str]:
    raw = str(phone or '').strip()
    digits = ''.join(ch for ch in raw if ch.isdigit())
    keys: set[str] = set()
    if digits:
        keys.add(digits)
    rule = _phone_localization_rule(country=country, area_code=area_code)
    has_specific_context = bool(rule)
    rules = [rule] if rule else list(PHONE_LOCALIZED_NUMBER_RULES.values())
    for candidate_rule in [item for item in rules if item]:
        country_code = str(candidate_rule.get('country_code') or '').strip()
        trunk_prefixes = tuple(str(item) for item in (candidate_rule.get('trunk_prefixes') or ()))
        local_lengths = {int(item) for item in (candidate_rule.get('local_lengths') or set())}
        body = _localized_phone_body(digits, candidate_rule)
        if not body:
            continue
        has_country_prefix = bool(country_code and digits.startswith(country_code) and len(digits) > len(country_code))
        has_local_trunk = any(digits.startswith(trunk) and len(digits) in local_lengths for trunk in trunk_prefixes)
        if not (
            has_specific_context
            or has_country_prefix
            or has_local_trunk
            or _localized_phone_body_matches_rule(body, candidate_rule)
        ):
            continue
        keys.add(body)
        if country_code:
            keys.add(f'{country_code}{body}')
            keys.add(f'+{country_code}{body}')
            keys.add(f'+{country_code} {body}')
        for trunk in trunk_prefixes:
            keys.add(f'{trunk}{body}')
    return {item for item in keys if item}



IGNORED_HISTORY_LEAD_STATUSES = frozenset({
    'archived_test_residue',
    'console_cleared_test_data',
})

GLOBAL_PHONE_PATTERN = re.compile(r'^\+(\d{1,3})(?:[ \-()]|\d){6,}$')
PHONE_CANDIDATE_PATTERN = re.compile(r'(\+?\d[\d \-().]{8,}\d)')
GROUP_VALUE_PATTERN = re.compile(r'^[A-Za-z]+-\d+$', flags=re.IGNORECASE)
GROUP_CANDIDATE_WITHOUT_DASH_PATTERN = re.compile(r'^[A-Za-z]+\d+$', flags=re.IGNORECASE)
REGISTRATION_GROUP_LABEL_PATTERN = re.compile(r'(?:^|\n|\b)(?:注册群组|注册群|group|registration_group)\s*[:：]\s*([^\n]+)', flags=re.IGNORECASE)
OTHER_CHANNEL_REGISTRATION_GROUP = '其他渠道'
EXTERNAL_APP_KNOWN_GUILD_APP_MAP = {
    'agency mx somente': 'timo',
    'royal latam': 'timo',
    '22000408': 'timo',
    'lvmy210446316420ie3d': 'timo',
    'timo001': 'timo',
    'royal id': 'timo',
    '11003905': 'timo',
    'agency of br somente': 'timo',
    'royal br': 'timo',
    '22000448': 'timo',
    'carote': 'linky',
    'permata': 'linky',
    'piso': 'linky',
    'sampanye': 'linky',
}
CJK_TEXT_PATTERN = re.compile(r'[\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff]')


def contains_cjk_text(value: Optional[str]) -> bool:
    return bool(CJK_TEXT_PATTERN.search(str(value or '')))


def is_blank_intake_field_value(value: Optional[str]) -> bool:
    text = str(value or '').strip()
    return not text or text.lower() in {'-', '—', '无', 'none', 'null', 'n/a', 'na'}


def normalize_registration_group_candidate(value: Optional[str]) -> Optional[str]:
    text = str(value or '').strip()
    if not text:
        return None
    if text == OTHER_CHANNEL_REGISTRATION_GROUP:
        return text
    # Do not treat labeled empty/place-holder fields or Chinese customer notes as group names.
    if re.match(r'^(?:phone|mobile|id|uid|account_id|code|invite\s*code|app|agency|guild|dept|group|registration_group|公会|邀请码|注册群组|注册群)\s*[:：]', text, flags=re.IGNORECASE):
        return None
    if is_blank_intake_field_value(text):
        return None
    # All real registration-group names are non-Chinese; do not classify Chinese customer notes as group names.
    if contains_cjk_text(text):
        return None
    # Strip common accidental trailing field fragments when a pasted value stays on one line.
    text = re.split(r'\s+(?:phone|mobile|id|uid|account_id|code|invite\s*code|app|agency|guild|dept|公会|邀请码)\s*[:：]', text, maxsplit=1, flags=re.IGNORECASE)[0].strip()
    text = text.strip(' ,;，；')
    if not text or text.strip().lower() in {'-', '—', '无', 'none', 'null', 'n/a', 'na'}:
        return None
    if contains_cjk_text(text):
        return None
    return text or None
PURE_DIGIT_ID_PATTERN = re.compile(r'^\d{6,12}$')
BARE_INVITE_CODE_PATTERN = re.compile(r'^(?:[A-Z]{6}|(?=.*[A-Z])[A-Z0-9]{6})$')
INVITE_CODE_CAPTURE_PATTERN = re.compile(r'(?:^|\b)(?:invite\s*code|personal\s*invite\s*code|code|个人邀请码|邀请码|kode\s+gabung\s+agensi|codigo\s+da\s+pessoa)\s*[:：是]?\s*([^\s"\'{}]{4,16})', flags=re.IGNORECASE)
INVITE_CODE_HOMOGLYPH_MAP = {
    'А': 'A', 'Β': 'B', 'В': 'B', 'С': 'C', 'Ε': 'E', 'Е': 'E', 'Η': 'H', 'Н': 'H', 'Ι': 'I', 'І': 'I',
    'Ј': 'J', 'Κ': 'K', 'К': 'K', 'М': 'M', 'Ν': 'N', 'О': 'O', 'Ο': 'O', 'Р': 'P', 'Ρ': 'P', 'Ѕ': 'S',
    'Т': 'T', 'Τ': 'T', 'Х': 'X', 'Χ': 'X', 'Υ': 'Y', 'Ү': 'Y', 'Ζ': 'Z',
    'а': 'A', 'β': 'B', 'в': 'B', 'с': 'C', 'ε': 'E', 'е': 'E', 'η': 'H', 'н': 'H', 'ι': 'I', 'і': 'I',
    'ј': 'J', 'κ': 'K', 'к': 'K', 'м': 'M', 'ո': 'N', 'ο': 'O', 'о': 'O', 'ρ': 'P', 'р': 'P', 'ѕ': 'S',
    'τ': 'T', 'т': 'T', 'χ': 'X', 'х': 'X', 'у': 'Y', 'γ': 'Y', 'ζ': 'Z',
}


def normalize_invite_code_candidate(raw: Optional[str]) -> Dict[str, Any]:
    raw_text = str(raw or '').strip()
    if not raw_text:
        return {
            'raw_input': None,
            'normalized': None,
            'has_homoglyphs': False,
            'unsupported_chars': [],
            'is_valid': False,
        }
    normalized_chars = []
    unsupported_chars = []
    has_homoglyphs = False
    for char in raw_text:
        upper_char = char.upper()
        if re.fullmatch(r'[A-Z0-9]', upper_char):
            normalized_chars.append(upper_char)
            continue
        mapped = INVITE_CODE_HOMOGLYPH_MAP.get(char) or INVITE_CODE_HOMOGLYPH_MAP.get(upper_char)
        if mapped:
            normalized_chars.append(mapped)
            has_homoglyphs = True
            continue
        unsupported_chars.append(char)
    normalized = ''.join(normalized_chars).upper() if normalized_chars else None
    has_confusable_characters = bool(normalized and any(char in {'0', 'O'} for char in normalized))
    is_valid = bool(normalized and BARE_INVITE_CODE_PATTERN.fullmatch(normalized) and not has_confusable_characters)
    return {
        'raw_input': raw_text,
        'normalized': normalized,
        'has_homoglyphs': has_homoglyphs,
        'has_confusable_characters': has_confusable_characters,
        'unsupported_chars': unsupported_chars,
        'is_valid': is_valid,
    }


def validate_invite_code_field(invite_code: Optional[str], *, invite_code_meta: Optional[Dict[str, Any]] = None) -> Optional[Dict[str, str]]:
    meta = invite_code_meta or normalize_invite_code_candidate(invite_code)
    normalized = str(meta.get('normalized') or '').strip().upper()
    if meta.get('unsupported_chars'):
        return {
            'reason': 'invalid_invite_code_format',
            'detail': 'invite_code contains unsupported non-Latin characters',
            'reply_text': 'Invalid Code. Use a 6-character personal code: letters or letters+digits, not all digits.',
        }
    if meta.get('has_confusable_characters'):
        return {
            'reason': 'invalid_code_confusable_characters',
            'detail': 'invite_code must not contain confusing characters 0 or O',
            'reply_text': 'Invalid Code. Do not use 0 or O. Use a 6-character personal code with letters or letters+digits.',
        }
    if normalized and not meta.get('is_valid'):
        return {
            'reason': 'invalid_invite_code_format',
            'detail': 'invite_code must be 6 English letters or letters+digits, not pure digits',
            'reply_text': 'Invalid Code. Use a 6-character personal code: letters or letters+digits, not all digits.',
        }
    return None


def extract_bare_multiline_candidates(text: str) -> Dict[str, Optional[str]]:
    # Treat tabs as field separators for pasted rows from spreadsheets/backoffice.
    # Do not let a phone candidate consume the adjacent ID cell, e.g.
    # "+62 877-6289-0159\t53321395".
    lines = [part.strip() for part in re.split(r'[\r\n\t]+', str(text or '')) if part.strip()]
    result: Dict[str, Optional[str]] = {
        'mobile_line': None,
        'registration_group_line': None,
        'account_id_line': None,
        'invite_code_line': None,
    }
    for line in lines:
        if result['mobile_line'] is None and GLOBAL_PHONE_PATTERN.fullmatch(line):
            result['mobile_line'] = line
            continue
        if result['registration_group_line'] is None and GROUP_VALUE_PATTERN.fullmatch(line):
            result['registration_group_line'] = line
            continue
        if result['account_id_line'] is None and PURE_DIGIT_ID_PATTERN.fullmatch(line):
            result['account_id_line'] = line
            continue
        if result['invite_code_line'] is None:
            invite_match = INVITE_CODE_CAPTURE_PATTERN.search(line)
            invite_meta = normalize_invite_code_candidate(invite_match.group(1) if invite_match else line)
            normalized = str(invite_meta.get('normalized') or '').strip().upper()
            if invite_meta.get('is_valid') and normalized.lower() not in {'linky', 'fumi'}:
                result['invite_code_line'] = normalized
                continue
    if result['registration_group_line'] is None:
        for line in lines:
            if line in {result.get('mobile_line'), result.get('account_id_line'), result.get('invite_code_line')}:
                continue
            if GLOBAL_PHONE_PATTERN.fullmatch(line) or PURE_DIGIT_ID_PATTERN.fullmatch(line):
                continue
            invite_meta = normalize_invite_code_candidate(line)
            normalized = str(invite_meta.get('normalized') or '').strip().upper()
            if INVITE_CODE_CAPTURE_PATTERN.search(line):
                continue
            if invite_meta.get('is_valid') and normalized.lower() not in {'linky', 'fumi'}:
                continue
            if re.fullmatch(r'(Linky|FUMI|Piso|Permata|Sampanye|Carote)', line, flags=re.IGNORECASE):
                continue
            candidate = normalize_registration_group_candidate(line)
            if candidate:
                result['registration_group_line'] = candidate
                break
    return result


def extract_invalid_group_candidate(text: str) -> Optional[str]:
    for line in [line.strip() for line in str(text or '').splitlines() if line.strip()]:
        if GROUP_CANDIDATE_WITHOUT_DASH_PATTERN.fullmatch(line):
            return line
    labeled_match = re.search(r'(?:注册群组|group)\s*[:：]?\s*([A-Za-z]+\d+)', str(text or ''), flags=re.IGNORECASE)
    if labeled_match:
        candidate = str(labeled_match.group(1) or '').strip()
        if GROUP_CANDIDATE_WITHOUT_DASH_PATTERN.fullmatch(candidate):
            return candidate
    return None


def normalize_phone_identity(*, mobile: str, area_code: int, country: str) -> tuple[str, int, str]:
    raw = str(mobile or '').strip()
    normalized_country = normalize_country_label(country)
    normalized_area_code = int(area_code or 0)

    international_match = re.fullmatch(r'\+(\d{1,3})\s+(\d{6,15})', raw)
    if international_match:
        prefix = international_match.group(1)
        body = international_match.group(2)
        normalized_area_code = int(prefix)
        normalized_country = PHONE_PREFIX_COUNTRY_MAP.get(prefix, normalized_country)
        return body, normalized_area_code, normalized_country

    explicit_prefix_match = re.fullmatch(r'\+(\d{1,3})[ \-().]+([\d \-().]{4,})', raw)
    if explicit_prefix_match:
        prefix = explicit_prefix_match.group(1)
        body = ''.join(ch for ch in explicit_prefix_match.group(2) if ch.isdigit())
        if body:
            normalized_area_code = int(prefix)
            normalized_country = PHONE_PREFIX_COUNTRY_MAP.get(prefix, normalized_country)
            return body, normalized_area_code, normalized_country

    if raw.startswith('+'):
        digits = '+' + ''.join(ch for ch in raw[1:] if ch.isdigit())
        for prefix in sorted(PHONE_PREFIX_COUNTRY_MAP.keys(), key=len, reverse=True):
            if digits.startswith(f'+{prefix}'):
                normalized_area_code = int(prefix)
                normalized_country = PHONE_PREFIX_COUNTRY_MAP[prefix]
                body = digits[len(prefix) + 1:]
                return body, normalized_area_code, normalized_country
        raw = ''.join(ch for ch in raw if ch.isdigit())
    else:
        raw = ''.join(ch for ch in raw if ch.isdigit())

    if normalized_area_code and not normalized_country:
        normalized_country = PHONE_PREFIX_COUNTRY_MAP.get(str(normalized_area_code), normalized_country)

    localized_rule = _phone_localization_rule(country=normalized_country, area_code=normalized_area_code)
    if localized_rule:
        country_code = str(localized_rule.get('country_code') or '').strip()
        localized_body = _localized_phone_body(raw, localized_rule)
        if localized_body:
            raw = localized_body
            if country_code and not normalized_area_code:
                normalized_area_code = int(country_code)
            if country_code and not normalized_country:
                normalized_country = PHONE_PREFIX_COUNTRY_MAP.get(country_code, normalized_country)

    return raw, normalized_area_code, normalized_country


def format_display_phone(phone: Optional[str], *, area_code: Optional[int] = None, country: Optional[str] = None) -> str:
    raw = str(phone or '').strip()
    if not raw or raw == '-':
        return '-'
    if re.search(r'[^\d\s+\-().]', raw):
        return raw
    digits_only = ''.join(ch for ch in raw if ch.isdigit())
    if raw.startswith('+'):
        normalized_mobile, normalized_area_code, _ = normalize_phone_identity(mobile=raw, area_code=int(area_code or 0), country=str(country or ''))
        if normalized_mobile and normalized_area_code:
            return f'+{normalized_area_code} {normalized_mobile}'
        return raw
    normalized_area_code = int(area_code or 0)
    normalized_country = infer_country_context(country) if country else ''
    if not normalized_area_code and normalized_country:
        normalized_mobile, normalized_area_code, _ = normalize_phone_identity(mobile=raw, area_code=0, country=normalized_country)
        if normalized_mobile and normalized_area_code:
            return f'+{normalized_area_code} {normalized_mobile}'
    if not normalized_area_code and digits_only:
        for prefix in sorted(PHONE_PREFIX_COUNTRY_MAP.keys(), key=len, reverse=True):
            if digits_only.startswith(prefix) and len(digits_only) > len(prefix):
                normalized_area_code = int(prefix)
                digits_only = digits_only[len(prefix):]
                break
    if normalized_area_code and digits_only:
        normalized_mobile, normalized_area_code, _ = normalize_phone_identity(
            mobile=digits_only,
            area_code=normalized_area_code,
            country=normalized_country,
        )
        return f'+{normalized_area_code} {normalized_mobile or digits_only}'
    return raw


EXTERNAL_APP_ID_ONLY_COUNTRY_PLACEHOLDER = 'Tugao'
EXTERNAL_APP_ID_ONLY_PHONE_PREFIX = 'Linky:'
EXTERNAL_APP_LEGACY_ID_ONLY_PHONE_PREFIXES = ('Tugao:',)


def is_external_app_id_only_phone(value: Any) -> bool:
    normalized_value = str(value or '').strip().casefold()
    supported_prefixes = (EXTERNAL_APP_ID_ONLY_PHONE_PREFIX, *EXTERNAL_APP_LEGACY_ID_ONLY_PHONE_PREFIXES)
    return any(normalized_value.startswith(prefix.casefold()) for prefix in supported_prefixes)


def make_external_app_id_only_phone(account_id: Any) -> str:
    compact_account_id = re.sub(r'[^0-9A-Za-z_.:-]+', '', str(account_id or '').strip()) or 'unknown'
    return f'{EXTERNAL_APP_ID_ONLY_PHONE_PREFIX}{compact_account_id}'


def validate_fast_intake_fields(*, mobile: Optional[str], app_name: Optional[str], account_id: Optional[str], country: Optional[str] = None) -> Optional[Dict[str, str]]:
    phone_text = str(mobile or '').strip()
    app_text = str(app_name or '').strip()
    account_text = str(account_id or '').strip()

    normalized_phone_text = format_display_phone(phone_text, country=country) if phone_text else ''
    if phone_text and not re.fullmatch(r'\+\d{1,3}\s\d{6,15}', normalized_phone_text):
        return {
            'reason': 'invalid_phone_format',
            'detail': 'mobile must use format +<country code> <number>',
            'reply_text': 'Invalid phone format. Use +<country code> <number>.',
        }

    if account_text and not account_text.isdigit():
        return {
            'reason': 'invalid_account_id_format',
            'detail': 'account_id must contain digits only',
            'reply_text': 'Invalid ID. Digits only.',
        }

    app_key = app_text.lower()
    app_id_lengths = {'linky': 8, 'fumi': 8, 'timo': 12}
    if app_key in app_id_lengths and account_text and not re.fullmatch(rf'\d{{{app_id_lengths[app_key]}}}', account_text):
        app_label = {'linky': 'Linky', 'fumi': 'FUMI', 'timo': 'Timo'}.get(app_key, app_text)
        return {
            'reason': 'invalid_account_id_format',
            'detail': f'{app_label} account_id must be exactly {app_id_lengths[app_key]} digits',
            'reply_text': f'Invalid ID. {app_label} requires exactly {app_id_lengths[app_key]} digits.',
        }

    return None


def parse_manual_cs_message(*, text: str, image_ocr_text: Optional[str] = None) -> Dict[str, Any]:
    text = str(text or '')
    image_ocr_text = str(image_ocr_text or '')
    combined = "\n".join(part for part in [text, image_ocr_text] if part).strip()
    normalized = combined.replace('：', ':')
    text_normalized = text.replace('：', ':')
    bare_candidates = extract_bare_multiline_candidates(text)
    ocr_normalized = normalize_native_ocr_fields(image_ocr_text) if image_ocr_text.strip() else {}
    country_context = infer_country_context(normalized, text_normalized, ocr_normalized.get('country'))

    mobile = None
    phone_candidate = None
    if bare_candidates.get('mobile_line'):
        phone_candidate = str(bare_candidates['mobile_line']).strip()
    else:
        phone_match = PHONE_CANDIDATE_PATTERN.search(normalized)
        if phone_match:
            phone_candidate = phone_match.group(1).strip()
    if phone_candidate:
        mobile, area_code, country = normalize_phone_identity(mobile=phone_candidate, area_code=0, country=country_context)
    else:
        area_code, country = 0, country_context

    text_account_id = None
    ocr_account_id = None
    labeled_patterns = [
        r'(?:^|\b)(?:id|uid|ywid|用户id|用户ID)\s*[:：是]?\s*(\d{6,})',
    ]
    for pattern in labeled_patterns:
        match = re.search(pattern, text.replace('：', ':'), flags=re.IGNORECASE)
        if match:
            text_account_id = match.group(1)
            break
    if image_ocr_text:
        match = re.search(r'(?:uid|id)\s*[:：]?\s*(\d{6,})', image_ocr_text, flags=re.IGNORECASE)
        if match:
            ocr_account_id = match.group(1)
    if not ocr_account_id:
        ocr_account_id = str(ocr_normalized.get('account_id') or '').strip() or None
    text_invite_code = None
    invite_code_meta = {
        'raw_input': None,
        'normalized': None,
        'has_homoglyphs': False,
        'unsupported_chars': [],
        'is_valid': False,
    }
    match = INVITE_CODE_CAPTURE_PATTERN.search(text.replace('：', ':'))
    if match:
        invite_code_meta = normalize_invite_code_candidate(str(match.group(1) or '').strip())
        if invite_code_meta.get('is_valid'):
            text_invite_code = str(invite_code_meta.get('normalized') or '').strip().upper() or None
    ocr_invite_meta = normalize_invite_code_candidate(
        str(
            ocr_normalized.get('person_code')
            or ocr_normalized.get('invite_code')
            or ocr_normalized.get('guild_invite_code')
            or ''
        ).strip().upper() or None
    )
    ocr_invite_code = str(ocr_invite_meta.get('normalized') or '').strip().upper() if ocr_invite_meta.get('is_valid') else None
    bare_invite_meta = normalize_invite_code_candidate(str(bare_candidates.get('invite_code_line') or '').strip().upper() or None)
    bare_invite_code = str(bare_invite_meta.get('normalized') or '').strip().upper() if bare_invite_meta.get('is_valid') else None
    invite_code = ocr_invite_code or text_invite_code or bare_invite_code
    selected_invite_meta = ocr_invite_meta if ocr_invite_code else (invite_code_meta if text_invite_code else bare_invite_meta)
    inferred_text_account_id = str(bare_candidates.get('account_id_line') or '').strip() or None
    account_id = ocr_account_id or text_account_id or inferred_text_account_id
    if not account_id:
        digit_runs = re.findall(r'\b\d{6,}\b', normalized)
        if mobile:
            digit_runs = [run for run in digit_runs if run != mobile and run != f"62{mobile}"]
        if digit_runs:
            account_id = digit_runs[-1]

    registration_group = None
    labeled_group_match = REGISTRATION_GROUP_LABEL_PATTERN.search(text_normalized)
    if labeled_group_match:
        registration_group = normalize_registration_group_candidate(labeled_group_match.group(1))
    group_patterns = [
        r'([A-Za-z]+-\d+)',
    ]
    for pattern in group_patterns:
        if registration_group:
            break
        match = re.search(pattern, text_normalized, flags=re.IGNORECASE)
        if match:
            registration_group = normalize_registration_group_candidate(match.group(1))
            break
    if not registration_group and bare_candidates.get('registration_group_line'):
        registration_group = str(bare_candidates['registration_group_line']).strip()

    app_name = None
    for app in ['Linky', 'FUMI']:
        if re.search(rf'\b{re.escape(app)}\b', text_normalized, flags=re.IGNORECASE):
            app_name = app
            break

    dept_name = None
    explicit_dept_patterns = [
        r'(?:公会|agency|guild|dept)\s*[:：]?\s*([A-Za-z]+(?:-\d+)?)',
    ]
    inferred_dept_patterns = [
        r'\b(Piso|Permata|Sampanye|Carote)\b(?!-\d)',
    ]
    for pattern in explicit_dept_patterns:
        match = re.search(pattern, text_normalized, flags=re.IGNORECASE)
        if match:
            dept_name = match.group(1)
            break
    if not dept_name:
        dept_name = str(ocr_normalized.get('guild_name') or ocr_normalized.get('agency_name') or '').strip() or None
    if not dept_name:
        for pattern in inferred_dept_patterns:
            match = re.search(pattern, text_normalized, flags=re.IGNORECASE)
            if match:
                dept_name = match.group(1)
                break

    conflicts = []
    if text_account_id and ocr_account_id and text_account_id != ocr_account_id:
        conflicts.append('account_id_conflict')

    missing_fields = [
            name for name, value in {
                'mobile': mobile,
                'account_id': account_id,
                'app_name': app_name,
                'dept_name': dept_name,
                'invite_code': invite_code,
            }.items() if not value
    ]

    score = 0.0
    weights = {
        'mobile': 0.2,
        'account_id': 0.2,
        'registration_group': 0.2,
        'app_name': 0.15,
        'dept_name': 0.15,
        'invite_code': 0.1,
    }
    values = {
        'mobile': mobile,
        'account_id': account_id,
        'registration_group': registration_group,
        'app_name': app_name,
        'dept_name': dept_name,
        'invite_code': invite_code,
    }
    for key, weight in weights.items():
        if values[key]:
            score += weight
    if conflicts:
        score -= 0.15
    confidence = max(0.0, min(round(score, 2), 1.0))

    return {
        'mobile': mobile,
        'area_code': area_code,
        'country': country,
        'account_id': account_id,
        'registration_group': registration_group,
        'app_name': app_name,
        'dept_name': dept_name,
        'invite_code': invite_code,
        'confidence': confidence,
        'missing_fields': missing_fields,
        'conflicts': conflicts,
        'evidence': {
            'text_used': bool(text.strip()),
            'image_ocr_used': bool(image_ocr_text.strip()),
            'text_account_id': text_account_id,
            'ocr_account_id': ocr_account_id,
            'text_invite_code': text_invite_code,
            'ocr_invite_code': ocr_invite_code,
            'invite_code_raw_input': selected_invite_meta.get('raw_input'),
            'invite_code_had_homoglyphs': bool(selected_invite_meta.get('has_homoglyphs')),
            'invite_code_unsupported_chars': list(selected_invite_meta.get('unsupported_chars') or []),
        },
        'invite_code_meta': selected_invite_meta,
        'raw_text': text,
        'raw_ocr_text': image_ocr_text,
    }


def extract_explicit_intake_fields(text: str) -> Dict[str, Optional[str]]:
    normalized = str(text or '').replace('：', ':')
    lines = [line.strip() for line in normalized.splitlines() if line.strip()]
    patterns = {
        'mobile': [r'(?:^|\n|\b)(?:phone|mobile)\s*:\s*([^\n]+)'],
        'account_id': [r'(?:^|\n|\b)(?:id|uid|account_id)\s*:\s*(\d{6,})'],
        'registration_group': [r'(?:^|\n|\b)(?:group|registration_group)\s*:\s*([^\n]+)'],
        'app_name': [r'(?:^|\n|\b)(?:app)\s*:\s*([^\n]+)'],
        'dept_name': [r'(?:^|\n|\b)(?:agency|guild|dept|公会)\s*:\s*([^\n]+)'],
        'invite_code': [r'(?:^|\n|\b)(?:invite\s*code|personal\s*invite\s*code|code|个人邀请码|邀请码)\s*:\s*([^\n]+)'],
    }
    result: Dict[str, Optional[str]] = {
        'mobile': None,
        'account_id': None,
        'registration_group': None,
        'app_name': None,
        'dept_name': None,
        'invite_code': None,
    }
    for key, pats in patterns.items():
        for pattern in pats:
            match = re.search(pattern, normalized, flags=re.IGNORECASE)
            if match:
                result[key] = str(match.group(1) or '').strip()
                break
    if result['app_name'] is None:
        for line in lines:
            if re.fullmatch(r'(Linky|FUMI)', line, flags=re.IGNORECASE):
                result['app_name'] = line
                break
    if result['dept_name'] is None:
        for line in lines:
            if re.fullmatch(r'(Piso|Permata|Sampanye|Carote)', line, flags=re.IGNORECASE):
                result['dept_name'] = line
                break
    return result


DEFAULT_DB_PATH = str(Path(__file__).resolve().parents[1] / "data" / "automation.db")



GROUP_ATMOSPHERE_PAGE_VERSION = '20260615-learning-extract-v1'















class LeadUpsertRequest(BaseModel):
    trace_id: str
    source_platform: str
    source_campaign: Optional[str] = None
    source_page_id: str
    country: str
    area_code: int
    mobile: str
    yw_id: Optional[str] = None
    app_name: Optional[str] = None
    dept_name: Optional[str] = None
    pendaftaran_group: Optional[str] = None
    inviter_id: Optional[str] = None
    occurred_at: Optional[str] = None
    parser_confidence: Optional[float] = None
    parser_missing_fields: list[str] = Field(default_factory=list)
    parser_conflicts: list[str] = Field(default_factory=list)
    parser_raw_text: Optional[str] = None
    parser_raw_ocr_text: Optional[str] = None
    parser_version: str = 'manual_cs_parser_v2'
    parser_status: str = 'unknown'
    review_reason_codes: list[str] = Field(default_factory=list)
    routing_decision: Optional[str] = None
    recommended_next_action: Optional[str] = None
    review_status: str = 'not_needed'


class EventCollectRequest(BaseModel):
    trace_id: str
    lead_id: Optional[str] = None
    event_type: str
    event_source: str
    event_value: Optional[str] = None
    page_id: Optional[str] = None
    session_id: Optional[str] = None
    operator_id: Optional[str] = None
    operator_name: Optional[str] = None
    raw_payload: Dict[str, Any] = Field(default_factory=dict)
    happened_at: Optional[str] = None


class TaskCreateRequest(BaseModel):
    lead_id: str
    task_type: str
    priority: str
    payload: Dict[str, Any] = Field(default_factory=dict)
    dedupe_key: str
    created_by: str
    created_at: str


class TaskResultRequest(BaseModel):
    status: str
    result_code: Optional[str] = None
    result_reason: Optional[str] = None
    toast_text: Optional[str] = None
    evidence_url: Optional[str] = None
    retry_count: int = 0
    executor_type: Optional[str] = None
    executor_id: Optional[str] = None
    finished_at: str
    raw_result: Dict[str, Any] = Field(default_factory=dict)


class CustomerSyncRequest(BaseModel):
    lead_id: str
    task_id: str
    yw_id: Optional[str] = None
    mobile: str
    area_code: int
    crm_patch: Dict[str, Any] = Field(default_factory=dict)
    sync_mode: str = "upsert"


class AccountSubmissionRequest(BaseModel):
    lead_id: str
    task_id: Optional[str] = None
    submission_type: str
    account_id: Optional[str] = None
    account_id_type: Optional[str] = None
    file_url: Optional[str] = None
    file_type: Optional[str] = None
    source_channel: Optional[str] = None
    expected_guild: Optional[str] = None
    route_snapshot: Dict[str, Any] = Field(default_factory=dict)
    source_bot_app_id: Optional[str] = None
    source_message_id: Optional[str] = None
    source_chat_id: Optional[str] = None
    submitted_by: Optional[str] = None
    submitted_at: str
    remark: Optional[str] = None


class ManualCsSubmissionRequest(BaseModel):
    mobile: str
    registration_group: str
    app_name: str
    dept_name: str
    country: Optional[str] = None
    invite_code: Optional[str] = None
    app_name_explicit: bool = False
    dept_name_explicit: bool = False
    submission_type: str
    account_id: Optional[str] = None
    file_url: Optional[str] = None
    file_type: Optional[str] = None
    image_ocr_text: Optional[str] = None
    submitted_by: str
    source_channel: str = "manual_cs_lark"
    source_bot_app_id: Optional[str] = None
    source_message_id: Optional[str] = None
    source_chat_id: Optional[str] = None
    remark: Optional[str] = None
    submitted_at: str




class OpsIntakeSubmitRequest(BaseModel):
    text: str
    profile_name: Optional[str] = None


class OpsIntakeParseRequest(BaseModel):
    text: str
    fields: Dict[str, Any] = Field(default_factory=dict)


class OpsBindFailedClearRequest(BaseModel):
    guild_name: Optional[str] = None
    date: Optional[str] = None
    submitted_by: Optional[str] = None
    scope: Optional[str] = None
    item_ids: list[str] = Field(default_factory=list)
    limit: int = Field(default=500, ge=1, le=2000)


class OpsIntakeResolveRequest(BaseModel):
    action: str = Field(default='resolved')
    reason: str = Field(default='')
    note: Optional[str] = None


class OpsIntakeGuildAssigneesRequest(BaseModel):
    user_ids: list[str] = Field(default_factory=list)


class OpsIntakeGuildHealthRefreshRequest(BaseModel):
    guild_names: list[str] = Field(default_factory=list)
    only_if_unknown_or_stale: bool = True


class OpsIntakeFeedbackDoneRequest(BaseModel):
    force: bool = False
    reason: Optional[str] = None


class OpsTimoIntakeSubmitRequest(BaseModel):
    guild_name: Optional[str] = None
    mobile: Optional[str] = None
    timo_id: Optional[str] = None
    group_name: Optional[str] = None
    app_name: Optional[str] = None
    source_text: Optional[str] = None
    source_channel: Optional[str] = None
    profile_name: Optional[str] = None
    auto_verify: bool = False


class OpsTimoIntakeVerifyRequest(BaseModel):
    force_crm_sync: bool = False


class OpsSugoIntakeVerifyRequest(BaseModel):
    guild_name: str
    sogo_id: str


class ExternalAppIntakeSubmissionRequest(BaseModel):
    source: str
    app: Optional[str] = None
    external_user_id: str
    external_session_id: Optional[str] = None
    external_message_id: Optional[str] = None
    customer_service_id: str
    customer_service_name: Optional[str] = None
    phone: Optional[str] = None
    country: Optional[str] = None
    linky_account_id: Optional[str] = None
    timo_id: Optional[str] = None
    guild: Optional[str] = None
    guild_id: Optional[str] = None
    guild_sid: Optional[str] = None
    code: Optional[str] = None
    group: Optional[str] = None
    raw_text: Optional[str] = None
    remark: Optional[str] = None


class ExternalAppPhoneBackfillRequest(BaseModel):
    source: str
    app: Optional[str] = None
    external_user_id: Optional[str] = None
    external_session_id: Optional[str] = None
    external_message_id: Optional[str] = None
    customer_service_id: Optional[str] = None
    customer_service_name: Optional[str] = None
    phone: str
    linky_account_id: str
    guild: str
    raw_text: Optional[str] = None
    remark: Optional[str] = None


class ExternalAppIntakeFeedbackActionRequest(BaseModel):
    customer_service_id: str
    customer_service_name: Optional[str] = None


class ManualReviewResolveRequest(BaseModel):
    decision: str
    reviewed_by: str
    review_note: Optional[str] = None
    account_id: Optional[str] = None
    app_name: Optional[str] = None
    dept_name: Optional[str] = None
    registration_group: Optional[str] = None
    submitted_at: str


class RecognitionResultRequest(BaseModel):
    status: str
    recognized_account_id: Optional[str] = None
    result_code: Optional[str] = None
    result_reason: Optional[str] = None
    finished_at: str
    raw_result: Dict[str, Any] = Field(default_factory=dict)


class BindCheckResultRequest(BaseModel):
    status: str
    result_code: Optional[str] = None
    result_reason: Optional[str] = None
    finished_at: str
    raw_result: Dict[str, Any] = Field(default_factory=dict)


class GroupJoinResultRequest(BaseModel):
    status: str
    result_code: Optional[str] = None
    result_reason: Optional[str] = None
    finished_at: str
    raw_result: Dict[str, Any] = Field(default_factory=dict)


class VoucherAttachRequest(BaseModel):
    image_path: str
    remark_suffix: Optional[str] = None


class RegistrationGroupApprovalBatchRequest(BaseModel):
    registration_group: str
    registration_group_name: Optional[str] = None
    approved_count: int
    approved_by: Optional[str] = None
    approved_by_name: Optional[str] = None
    source_platform: Optional[str] = None
    source_campaign: Optional[str] = None
    source_adset: Optional[str] = None
    source_ad: Optional[str] = None
    approved_at: str
    area: str = "Indonesia"
    remark: Optional[str] = None
    approval_run_id: Optional[str] = None


class RegistrationGroupApprovalDecisionRequest(BaseModel):
    registration_group: str
    decision: str = 'approve'
    decided_at: str
    decided_by: Optional[str] = None
    decided_by_name: Optional[str] = None
    source_platform: Optional[str] = None
    source_campaign: Optional[str] = None
    source_adset: Optional[str] = None
    source_ad: Optional[str] = None
    target_name_hint: Optional[str] = None
    target_phone_hint: Optional[str] = None
    approved_count: int = 1
    area: str = 'Indonesia'
    remark: Optional[str] = None
    force_immediate: bool = False
    expected_pending_count: Optional[int] = None
    expected_member_count: Optional[int] = None
    expected_requester_ids: Optional[List[str]] = None
    expected_requesters: Optional[List[Dict[str, Any]]] = None


class OfficialGroupApprovalCheckRequest(BaseModel):
    lead_id: str
    target_group: str
    checked_at: str
    checked_by: Optional[str] = None
    checked_by_name: Optional[str] = None
    source_platform: Optional[str] = None
    source_campaign: Optional[str] = None
    source_adset: Optional[str] = None
    source_ad: Optional[str] = None
    target_phone_hint: Optional[str] = None
    target_requester_id: Optional[str] = None
    target_requester_pending_hint: Optional[bool] = None
    remark: Optional[str] = None


class OfficialGroupApprovalDecisionRequest(BaseModel):
    lead_id: Optional[str] = None
    target_group: str
    decision: str = 'approve'
    decided_at: str
    decided_by: Optional[str] = None
    decided_by_name: Optional[str] = None
    source_platform: Optional[str] = None
    source_campaign: Optional[str] = None
    source_adset: Optional[str] = None
    source_ad: Optional[str] = None
    target_name_hint: Optional[str] = None
    target_phone_hint: Optional[str] = None
    target_requester_id: Optional[str] = None
    target_requester_pending_hint: Optional[bool] = None
    approval_runtime_account_key: Optional[str] = None
    approval_runtime_binding_index: Optional[int] = None
    approval_runtime_binding_target: Optional[str] = None
    remark: Optional[str] = None


class OfficialGroupApprovalRetryRequest(BaseModel):
    target_group: str
    decided_at: str
    decided_by: Optional[str] = None
    decided_by_name: Optional[str] = None
    source_platform: Optional[str] = None
    source_campaign: Optional[str] = None
    source_adset: Optional[str] = None
    source_ad: Optional[str] = None
    remark: Optional[str] = None


class OfficialGroupBatchRunRequest(BaseModel):
    decided_at: str
    decided_by: Optional[str] = None
    decided_by_name: Optional[str] = None
    source_platform: Optional[str] = None
    source_campaign: Optional[str] = None
    source_adset: Optional[str] = None
    source_ad: Optional[str] = None
    remark: Optional[str] = None
    limit_groups: int = 10
    limit_leads_per_group: Optional[int] = None
    account_key: Optional[str] = None
    binding_index: Optional[int] = None
    registration_group: Optional[str] = None
    target_group: Optional[str] = None
    allow_live_crm_phone_match: bool = True
    allow_crm_only_test_match: bool = True
    suppress_success_notifications: bool = False


class GroupApprovalCheckRequest(BaseModel):
    approval_scope: str
    registration_group: Optional[str] = None
    lead_id: Optional[str] = None
    target_group: Optional[str] = None
    checked_at: str
    checked_by: Optional[str] = None
    checked_by_name: Optional[str] = None
    source_platform: Optional[str] = None
    source_campaign: Optional[str] = None
    source_adset: Optional[str] = None
    source_ad: Optional[str] = None
    target_phone_hint: Optional[str] = None
    target_requester_id: Optional[str] = None
    remark: Optional[str] = None


class GroupApprovalDecisionRequest(BaseModel):
    approval_scope: str
    registration_group: Optional[str] = None
    lead_id: Optional[str] = None
    target_group: Optional[str] = None
    decision: str = 'approve'
    decided_at: str
    decided_by: Optional[str] = None
    decided_by_name: Optional[str] = None
    source_platform: Optional[str] = None
    source_campaign: Optional[str] = None
    source_adset: Optional[str] = None
    source_ad: Optional[str] = None
    target_name_hint: Optional[str] = None
    target_phone_hint: Optional[str] = None
    target_requester_id: Optional[str] = None
    approved_count: int = 1
    area: str = 'Indonesia'
    remark: Optional[str] = None
    force_immediate: bool = False
    expected_pending_count: Optional[int] = None
    expected_member_count: Optional[int] = None
    expected_requester_ids: Optional[List[str]] = None
    expected_requesters: Optional[List[Dict[str, Any]]] = None


class GroupApprovalBatchRunRequest(BaseModel):
    approval_scope: str
    decided_at: str
    decided_by: Optional[str] = None
    decided_by_name: Optional[str] = None
    source_platform: Optional[str] = None
    source_campaign: Optional[str] = None
    source_adset: Optional[str] = None
    source_ad: Optional[str] = None
    remark: Optional[str] = None
    limit_groups: int = 10
    limit_leads_per_group: Optional[int] = None
    allow_live_crm_phone_match: bool = True
    allow_crm_only_test_match: bool = True
    suppress_success_notifications: bool = False


class GroupApprovalExecutorWarmupRequest(BaseModel):
    approval_scope: str


class NotificationReadRequest(BaseModel):
    read_by: Optional[str] = None


class IntakeBotPresetUpdateRequest(BaseModel):
    app_id: Optional[str] = None
    app_secret: Optional[str] = None
    robot_name: Optional[str] = None
    default_app: str
    default_guild: str


class LocalIntakeBotGatewayActivationRequest(IntakeBotPresetUpdateRequest):
    pass


class GuildExecutorUpdateRequest(BaseModel):
    # Backend URLs are fixed by product policy; legacy fields are accepted but ignored.
    backend_url: Optional[str] = None
    guild_backend_url: Optional[str] = None
    login_username: Optional[str] = None
    password_secret_ref: Optional[str] = None
    guild_backend_token: Optional[str] = None
    oauth_token: Optional[str] = None
    oauth_token_secret: Optional[str] = None
    platform_backend_url: Optional[str] = None
    platform_backend_token: Optional[str] = None
    platform_authorization: Optional[str] = None
    cms_refresh_token: Optional[str] = None
    refresh_token: Optional[str] = None
    cms_guild_id: Optional[str] = None
    cms_guild_sid: Optional[str] = None
    country: Optional[str] = None
    guild_country: Optional[str] = None
    eligible_user_countries: List[str] = Field(default_factory=list)
    routing_region: Optional[str] = None
    proxy_url: Optional[str] = None
    proxy_region: Optional[str] = None
    proxy_type: Optional[str] = None
    enabled: bool = True
    browser_profile_key: Optional[str] = None
    bind_concurrency: int = 1
    request_timeout_seconds: int = 30
    notes: Optional[str] = None


class TimoGuildExecutorUpdateRequest(BaseModel):
    platform_backend_url: Optional[str] = None
    platform_authorization: Optional[str] = None
    cms_guild_id: Optional[str] = None
    cms_guild_sid: Optional[str] = None
    country: Optional[str] = None
    guild_country: Optional[str] = None
    eligible_user_countries: List[str] = Field(default_factory=list)
    routing_region: Optional[str] = None
    enabled: bool = True
    bind_concurrency: Optional[int] = None
    request_timeout_seconds: Optional[int] = None
    notes: Optional[str] = None


class SugoGuildExecutorUpdateRequest(BaseModel):
    login_username: Optional[str] = None
    platform_backend_url: Optional[str] = None
    platform_authorization: Optional[str] = None
    cms_refresh_token: Optional[str] = None
    refresh_token: Optional[str] = None
    cms_guild_id: Optional[str] = None
    cms_guild_sid: Optional[str] = None
    country: Optional[str] = None
    enabled: bool = True
    bind_concurrency: Optional[int] = None
    request_timeout_seconds: Optional[int] = None
    notes: Optional[str] = None


class ProductionOpsDaemonConfigUpdateRequest(BaseModel):
    enabled: bool = False
    registration_group: Optional[str] = None
    api_base_url: Optional[str] = None
    worker_base_url: Optional[str] = None
    interval_seconds: float = 20.0
    notify_chat_id: Optional[str] = None
    area: Optional[str] = None
    remark: Optional[str] = None
    approved_count: int = 1
    auto_recover_worker: bool = True


class ApprovalScheduleWindowRequest(BaseModel):
    start: str
    end: str


class ApprovalGroupBindingRequest(BaseModel):
    binding_id: Optional[str] = None
    link: Optional[str] = None
    group_name: Optional[str] = None
    area: Optional[str] = None
    notify_profile_name: Optional[str] = None
    enabled: Optional[bool] = True
    registration_group: Optional[str] = None
    group_id: Optional[str] = None
    provider_mode: Optional[str] = None
    registration_group_runtime: Optional[str] = None
    official_group_runtime: Optional[str] = None
    group_assistant_runtime: Optional[str] = None
    provider_capabilities: Dict[str, Any] = Field(default_factory=dict)
    baileys_enabled: Optional[bool] = None
    baileys_base_url: Optional[str] = None
    provider_base_url: Optional[str] = None
    baileys_token: Optional[str] = None
    provider_token: Optional[str] = None
    runtime_token: Optional[str] = None
    baileys_account_id: Optional[str] = None
    provider_account_id: Optional[str] = None
    account_id: Optional[str] = None
    approval_count_threshold: Optional[int] = None
    approval_timeout_minutes: Optional[int] = None
    auto_recover_worker: Optional[bool] = None
    schedule_windows: list[ApprovalScheduleWindowRequest] = []


class WhatsAppApprovalAccountUpdateRequest(BaseModel):
    account_name: str
    responsible_type: str
    assigned_customer_service_user_id: Optional[str] = None
    assigned_customer_service_user_ids: list[str] = []
    group_links: list[str] = []
    group_link_bindings: list[ApprovalGroupBindingRequest] = []
    area: Optional[str] = None
    notify_profile_name: Optional[str] = None
    provider_mode: Optional[str] = None
    registration_group_runtime: Optional[str] = None
    official_group_runtime: Optional[str] = None
    group_assistant_runtime: Optional[str] = None
    provider_capabilities: Dict[str, Any] = Field(default_factory=dict)
    baileys_enabled: Optional[bool] = None
    baileys_base_url: Optional[str] = None
    provider_base_url: Optional[str] = None
    baileys_token: Optional[str] = None
    provider_token: Optional[str] = None
    runtime_token: Optional[str] = None
    baileys_account_id: Optional[str] = None
    provider_account_id: Optional[str] = None
    account_id: Optional[str] = None
    approval_rule: Optional[str] = None
    approval_count_threshold: Optional[int] = None
    approval_timeout_minutes: Optional[int] = None
    auto_recover_worker: bool = True
    schedule_windows: list[ApprovalScheduleWindowRequest] = []
    enabled: bool = True
    notes: Optional[str] = None


class WhatsAppApprovalPairingCodeRequest(BaseModel):
    phone_number: str


class WhatsAppApprovalAreaOptionsUpdateRequest(BaseModel):
    options: list[str]


class McnRegionOptionUpdateRequest(BaseModel):
    code: str
    enabled: bool = True
    sort_order: Optional[int] = None


class McnRegionOptionsUpdateRequest(BaseModel):
    options: list[McnRegionOptionUpdateRequest] = Field(default_factory=list)


class SubmissionResubmitRequest(BaseModel):
    corrected_by: str
    submitted_at: str
    mobile: Optional[str] = None
    registration_group: Optional[str] = None
    invite_code: Optional[str] = None
    account_id: Optional[str] = None
    remark: Optional[str] = None


GUILD_BACKEND_BASE_URL = 'https://guild.linke.ai/guild'
PLATFORM_BACKEND_BASE_URL = 'https://cms.linke.ai/'

GUILD_EXECUTOR_PROXY_REGION_OPTIONS: list[dict[str, str]] = [
    {'value': '北京', 'label': '北京'},
    {'value': '上海', 'label': '上海'},
    {'value': '广州', 'label': '广州'},
    {'value': '深圳', 'label': '深圳'},
    {'value': '杭州', 'label': '杭州'},
    {'value': '南京', 'label': '南京'},
    {'value': '苏州', 'label': '苏州'},
    {'value': '成都', 'label': '成都'},
    {'value': '重庆', 'label': '重庆'},
    {'value': '武汉', 'label': '武汉'},
    {'value': '西安', 'label': '西安'},
    {'value': '郑州', 'label': '郑州'},
    {'value': '长沙', 'label': '长沙'},
    {'value': '厦门', 'label': '厦门'},
    {'value': '福州', 'label': '福州'},
]
GUILD_EXECUTOR_PROXY_REGION_VALUES: set[str] = {item['value'] for item in GUILD_EXECUTOR_PROXY_REGION_OPTIONS}
PROJECT_ROOT = Path(__file__).resolve().parents[1]
TIMO_GUILD_TASK_TIER_WAN: Tuple[int, ...] = (
    6, 12, 20, 30, 40, 60, 100, 140, 180, 240, 400, 500, 600, 700, 800,
    1000, 1200, 1400, 1600, 1800, 2000, 2400, 2800, 3200, 3600, 4000,
    4500, 5000, 5500, 6000, 6500, 7000, 7500, 8000, 8500, 9000, 9500,
    10000, 10500, 11000, 12000, 12500, 13000, 14000, 14500, 15000, 16000,
    16500, 17000, 17500, 18000, 18250, 18500, 19000, 19250, 19500, 20000,
    20500, 21000, 22000, 22500, 23000, 24000, 24500, 25000, 26000, 26500,
    27000, 28000, 30000,
)
TIMO_MEXICO_GUILD_TASK_REWARD_WAN: Dict[int, float] = {
    0: 0,
    6: 1.2,
    12: 1.4,
    20: 1.8,
    30: 2.2,
    40: 3,
    60: 5.4,
    100: 9,
    140: 9,
    180: 10,
    240: 13,
    400: 30,
    500: 24,
    600: 24,
    700: 24,
    800: 24,
    1000: 58,
    1200: 58,
    1400: 58,
    1600: 58,
    1800: 58,
    2000: 58,
    2400: 70,
    2800: 80,
    3200: 80,
    3600: 80,
    4000: 160,
    4500: 100,
    5000: 100,
    5500: 100,
    6000: 100,
    6500: 100,
    7000: 100,
    7500: 100,
    8000: 100,
    8500: 100,
    9000: 100,
    9500: 100,
    10000: 100,
    10500: 100,
    11000: 100,
    12000: 200,
    12500: 100,
    13000: 100,
    14000: 200,
    14500: 100,
    15000: 100,
    16000: 200,
    16500: 100,
    17000: 100,
    17500: 100,
    18000: 100,
    18250: 100,
    18500: 100,
    19000: 100,
    19250: 200,
    19500: 100,
    20000: 100,
    20500: 100,
    21000: 100,
    22000: 200,
    22500: 100,
    23000: 100,
    24000: 200,
    24500: 100,
    25000: 100,
    26000: 200,
    26500: 100,
    27000: 100,
    28000: 200,
    30000: 400,
}
TIMO_BRAZIL_GUILD_TASK_REWARD_WAN: Dict[int, float] = {
    0: 0,
    6: 1.2,
    12: 1.4,
    20: 1.8,
    30: 2.2,
    60: 5.4,
    100: 9,
    140: 9,
    180: 10,
    240: 13,
    400: 30,
    500: 24,
    600: 24,
    700: 24,
    800: 24,
    1000: 58,
    1200: 58,
    1400: 58,
    1600: 58,
    1800: 58,
    2000: 58,
    2400: 70,
    2800: 80,
    3200: 80,
    3600: 80,
    4000: 160,
    4500: 100,
    5000: 100,
    5500: 100,
    6000: 100,
    6500: 100,
    7000: 100,
    8000: 100,
    8500: 100,
    9000: 100,
    9500: 100,
    10000: 100,
    10500: 100,
    11000: 100,
    12000: 200,
    12500: 100,
    13000: 100,
    14000: 200,
    14500: 100,
    15000: 100,
    16000: 200,
    16500: 100,
    17000: 100,
    17500: 100,
    18000: 100,
    18250: 100,
    18500: 100,
    19000: 100,
    19250: 100,
    19500: 100,
    20000: 100,
    20500: 100,
    21000: 100,
    22000: 200,
    22500: 100,
    23000: 100,
    24000: 200,
    24500: 100,
    25000: 100,
    26000: 200,
    26500: 100,
    27000: 100,
    28000: 120,
    30000: 200,
    31000: 200,
    32000: 200,
    33000: 200,
    34000: 200,
    35000: 200,
    36000: 200,
    37000: 200,
    38000: 200,
    39000: 200,
    40000: 200,
    41000: 200,
    42000: 200,
    43000: 200,
    44000: 200,
    46000: 400,
    48000: 400,
    50000: 400,
}
TIMO_INDONESIA_GUILD_TASK_REWARD_WAN: Dict[int, int] = {
    0: 0,
    10: 1,
    20: 2,
    60: 10,
    100: 10,
    140: 10,
    200: 12,
    260: 13,
    320: 14,
    400: 23,
    500: 25,
    600: 30,
    800: 50,
    1000: 50,
    1200: 50,
    1400: 50,
    1600: 50,
    1800: 50,
    2000: 50,
    2500: 100,
    3000: 100,
    3500: 130,
    4000: 170,
    5000: 250,
    6000: 400,
}
TIMO_QUALITY_HOST_TASK_TIER_WAN: Tuple[int, ...] = (400,)
TIMO_REWARD_CLAIM_STATE_PATH = Path(
    os.getenv('TIMO_REWARD_CLAIM_STATE_PATH') or PROJECT_ROOT / 'data' / 'timo_guild_reward_claim_status.json'
)
TIMO_REWARD_CLAIM_HISTORY_PATH = Path(
    os.getenv('TIMO_REWARD_CLAIM_HISTORY_PATH') or PROJECT_ROOT / 'data' / 'timo_guild_reward_claim_history.jsonl'
)
_TIMO_REWARD_HISTORY_INDEX_LOCK = threading.Lock()
_TIMO_REWARD_HISTORY_INDEX: Dict[str, Any] = {
    'version': None,
    'payloads': (),
    'entries_by_guild': {},
}


def _load_timo_reward_history_index() -> Dict[str, Any]:
    """Parse reward history once per file version and index results by guild."""
    global _TIMO_REWARD_HISTORY_INDEX
    path = TIMO_REWARD_CLAIM_HISTORY_PATH
    try:
        stat = path.stat()
        version: Tuple[Any, ...] = (str(path), int(stat.st_mtime_ns), int(stat.st_size))
    except OSError:
        version = (str(path), 'missing')
    with _TIMO_REWARD_HISTORY_INDEX_LOCK:
        if _TIMO_REWARD_HISTORY_INDEX.get('version') == version:
            return _TIMO_REWARD_HISTORY_INDEX
        payloads: List[Dict[str, Any]] = []
        entries_by_guild: Dict[str, List[Dict[str, Any]]] = {}
        if version[-1] != 'missing':
            try:
                lines = path.read_text(encoding='utf-8').splitlines()
            except OSError:
                lines = []
            for raw_line in lines:
                line = raw_line.strip()
                if not line:
                    continue
                try:
                    payload = json.loads(line)
                except json.JSONDecodeError:
                    continue
                if not isinstance(payload, dict):
                    continue
                payloads.append(payload)
                payload_checked_at = str(payload.get('checked_at') or payload.get('checked_at_iso') or '').strip()
                for result in payload.get('results') or []:
                    if not isinstance(result, dict):
                        continue
                    names = {
                        str(result.get(key) or '').strip().lower()
                        for key in ('account', 'guild_name', 'guild', 'executor_name')
                        if str(result.get(key) or '').strip()
                    }
                    entry = {'payload_checked_at': payload_checked_at, 'result': result}
                    for name in names:
                        entries_by_guild.setdefault(name, []).append(entry)
        _TIMO_REWARD_HISTORY_INDEX = {
            'version': version,
            'payloads': tuple(payloads),
            'entries_by_guild': {name: tuple(entries) for name, entries in entries_by_guild.items()},
        }
        return _TIMO_REWARD_HISTORY_INDEX
TIMO_EXECUTOR_KEEPALIVE_STATUS_PATH = Path(
    os.getenv('TIMO_KEEPALIVE_STATE_PATH') or PROJECT_ROOT / 'data' / 'timo_executor_keepalive_status.json'
)
TIMO_KEEPALIVE_RUNNER_PATH = PROJECT_ROOT / 'scripts' / 'run_timo_keepalive.sh'
TIMO_EXECUTOR_KEEPALIVE_STALE_UNKNOWN_SECONDS = 15 * 60
TIMO_ANCHOR_EXPORT_CACHE_STATUS_PATHS = [
    PROJECT_ROOT / 'data' / 'timo_anchor_export_cache_daily_status.json',
    PROJECT_ROOT / 'data' / 'timo_anchor_export_cache_weekly_status.json',
    PROJECT_ROOT / 'data' / 'timo_anchor_export_cache_real_person_status.json',
    PROJECT_ROOT / 'data' / 'timo_anchor_export_cache_first_20k_diamonds_status.json',
]
LEGACY_SHARED_WEBJS_8787_HOSTS = {'127.0.0.1', 'localhost', '::1'}
LEGACY_SHARED_WEBJS_8787_PORT = 8787
PRODUCTION_OPS_DAEMON_LABEL = 'com.chauncey.mcn.production-ops-daemon'
PRODUCTION_OPS_DAEMON_LAUNCH_AGENT_PATH = Path.home() / 'Library' / 'LaunchAgents' / f'{PRODUCTION_OPS_DAEMON_LABEL}.plist'
PRODUCTION_OPS_DAEMON_ENV_PATH = PROJECT_ROOT / 'data' / 'production_ops_daemon.env'
PRODUCTION_OPS_DAEMON_STATUS_PATH = PROJECT_ROOT / 'data' / 'production_ops_daemon_status.json'
PRODUCTION_OPS_DAEMON_STATE_PATH = PROJECT_ROOT / 'data' / 'production_ops_daemon_state.json'
CMS_EXECUTOR_KEEPALIVE_STATUS_PATH = PROJECT_ROOT / 'data' / 'cms_executor_keepalive_status.json'
CMS_EXECUTOR_KEEPALIVE_STALE_UNKNOWN_SECONDS = 15 * 60
CMS_EXECUTOR_KEEPALIVE_SCRIPT_CANDIDATES = [
    PROJECT_ROOT / 'scripts' / 'cms_executor_keepalive.py',
    PROJECT_ROOT / 'scripts' / 'cms_keepalive.py',
]
PRODUCTION_OPS_DAEMON_INSTALL_SCRIPT = PROJECT_ROOT / 'scripts' / 'install_production_ops_daemon_launch_agent.sh'
PRODUCTION_OPS_DAEMON_UNINSTALL_SCRIPT = PROJECT_ROOT / 'scripts' / 'uninstall_production_ops_daemon_launch_agent.sh'
WHATSAPP_APPROVAL_DEFAULT_COUNT_THRESHOLD = 30
REGISTRATION_GROUP_APPROVAL_MAX_SINGLE_COUNT = 25
WHATSAPP_APPROVAL_DEFAULT_TIMEOUT_MINUTES = 30
WHATSAPP_APPROVAL_NOTIFY_PROFILE_BY_RESPONSIBLE_TYPE = {
    'registration_group': 'wa-approval-broadcast-02',
    'official_group': 'wa-approval-broadcast-03',
}
WHATSAPP_APPROVAL_WORKER_ROOT = PROJECT_ROOT / 'webjs-approval-worker'
WHATSAPP_APPROVAL_WORKER_AUTH_ACCOUNTS_DIR=WHATSAPP_APPROVAL_WORKER_ROOT / '.wwebjs_auth_accounts'
WHATSAPP_APPROVAL_WORKER_RUNTIME_DIR = PROJECT_ROOT / 'data' / 'whatsapp_approval_worker_runtimes'
WHATSAPP_APPROVAL_WORKER_LOG_DIR = PROJECT_ROOT / 'logs' / 'whatsapp_approval_workers'
WHATSAPP_APPROVAL_LOCALAUTH_RECOVERY_GRACE_SECONDS = 180.0
WHATSAPP_APPROVAL_WORKER_RESTART_SCRIPT = PROJECT_ROOT / 'scripts' / 'restart_registration_group_webjs_worker.sh'
OFFICIAL_GROUP_BRIDGE_DEFAULT_BASE_URL = 'http://127.0.0.1:55801'
OFFICIAL_GROUP_BRIDGE_START_SCRIPT = PROJECT_ROOT / 'scripts' / 'start_official_group_bridge.sh'
GROUP_ATMOSPHERE_MEDIA_UPLOAD_MAX_BYTES = 10 * 1024 * 1024
GROUP_ATMOSPHERE_CHAT_RECORD_JSON_MAX_BYTES = 30 * 1024 * 1024
GROUP_ATMOSPHERE_CHAT_RECORD_MAX_LINES = 50000
GROUP_ATMOSPHERE_LEGACY_DEFAULT_PHRASE_TYPE_KEYS = frozenset({
    'community_seed',
    'faq_helper',
    'newcomer_guide',
    'motivation_admin',
})


def _coerce_registration_group_single_approval_count(value: Any, *, pending_count: Optional[int] = None) -> int:
    try:
        count = max(1, int(value or 1))
    except Exception:
        count = 1
    count = min(count, REGISTRATION_GROUP_APPROVAL_MAX_SINGLE_COUNT)
    if pending_count is None:
        return count
    try:
        pending = max(0, int(pending_count))
    except Exception:
        return count
    if pending <= 0:
        return 0
    return min(count, pending)


def _whatsapp_approval_notify_profile_for_responsible_type(responsible_type: Any) -> str:
    return WHATSAPP_APPROVAL_NOTIFY_PROFILE_BY_RESPONSIBLE_TYPE.get(str(responsible_type or '').strip(), '')


def _legacy_shared_webjs_8787_allowed() -> bool:
    return str(os.getenv('ALLOW_LEGACY_WEBJS_8787') or '').strip().lower() in {'1', 'true', 'yes', 'on'}


def _is_legacy_shared_webjs_8787_url(raw_url: Any) -> bool:
    text = str(raw_url or '').strip().rstrip('/')
    if not text:
        return False
    candidate = text if '://' in text else f'http://{text}'
    try:
        parsed = urlparse(candidate)
    except Exception:
        return False
    host = str(parsed.hostname or '').strip().lower()
    try:
        port = parsed.port
    except ValueError:
        return False
    return host in LEGACY_SHARED_WEBJS_8787_HOSTS and port == LEGACY_SHARED_WEBJS_8787_PORT


def _sanitize_legacy_shared_webjs_worker_base_url(raw_url: Any) -> str:
    base_url = str(raw_url or '').strip().rstrip('/')
    if not base_url:
        return ''
    if _is_legacy_shared_webjs_8787_url(base_url) and not _legacy_shared_webjs_8787_allowed():
        return ''
    return base_url


def _coerce_positive_int(value: Any, default: int) -> int:
    try:
        normalized = int(value)
    except (TypeError, ValueError):
        return default
    return normalized if normalized > 0 else default


def _coerce_nonnegative_int(value: Any, default: int) -> int:
    if value is None or str(value).strip() == '':
        return default
    try:
        normalized = int(value)
    except (TypeError, ValueError):
        return default
    return normalized if normalized >= 0 else default


def _group_atmosphere_interval_seconds(seconds_value: Any, legacy_minutes_value: Any, default: int) -> int:
    # The DB/API legacy field is named *_minutes, but the ops page has always exposed seconds.
    selected = seconds_value if seconds_value is not None and str(seconds_value).strip() != '' else legacy_minutes_value
    return _coerce_nonnegative_int(selected, default)


def _group_atmosphere_mapping_interval_seconds(mapping: Dict[str, Any], seconds_key: str, legacy_minutes_key: str, default: int) -> int:
    source = mapping or {}
    return _group_atmosphere_interval_seconds(source.get(seconds_key), source.get(legacy_minutes_key), default)


def _legacy_approval_thresholds(rule: str) -> tuple[int, int]:
    normalized = str(rule or '').strip()
    if normalized == 'timeout_30m':
        return WHATSAPP_APPROVAL_DEFAULT_COUNT_THRESHOLD, WHATSAPP_APPROVAL_DEFAULT_TIMEOUT_MINUTES
    return WHATSAPP_APPROVAL_DEFAULT_COUNT_THRESHOLD, WHATSAPP_APPROVAL_DEFAULT_TIMEOUT_MINUTES


def _approval_condition_text(count_threshold: int, timeout_minutes: int) -> str:
    return f'满{count_threshold}人或{timeout_minutes}分钟'


def _normalize_schedule_windows_payload(items: list[Any]) -> list[dict[str, str]]:
    normalized: list[dict[str, str]] = []
    for item in items or []:
        if not isinstance(item, dict):
            continue
        start = str(item.get('start') or '').strip()
        end = str(item.get('end') or '').strip()
        if not re.fullmatch(r'\d{2}:\d{2}', start) or not re.fullmatch(r'\d{2}:\d{2}', end):
            continue
        normalized.append({'start': start, 'end': end})
    return normalized


MCN_INTERNAL_REGION_OPTIONS: list[dict[str, Any]] = [
    {'code': 'ID', 'value': 'Indonesia', 'label': 'Indonesia', 'label_zh': '印尼', 'phone_code': '62', 'language': 'id', 'enabled': True, 'sort_order': 10, 'aliases': ['Indonesia', '印尼', 'indo', 'ID']},
    {'code': 'BR', 'value': 'Brazil', 'label': 'Brazil', 'label_zh': '巴西', 'phone_code': '55', 'language': 'pt', 'enabled': True, 'sort_order': 20, 'aliases': ['Brazil', 'Brasil', '巴西', 'BR']},
    {'code': 'MX', 'value': 'Mexico', 'label': 'Mexico', 'label_zh': '墨西哥', 'phone_code': '52', 'language': 'es', 'enabled': True, 'sort_order': 30, 'aliases': ['Mexico', 'México', '墨西哥', 'MX']},
    {'code': 'VE', 'value': 'Venezuela', 'label': 'Venezuela', 'label_zh': '委内瑞拉', 'phone_code': '58', 'language': 'es', 'enabled': True, 'sort_order': 31, 'aliases': ['Venezuela', '委内瑞拉', 'VE']},
    {'code': 'CL', 'value': 'Chile', 'label': 'Chile', 'label_zh': '智利', 'phone_code': '56', 'language': 'es', 'enabled': True, 'sort_order': 32, 'aliases': ['Chile', '智利', 'CL']},
    {'code': 'CO', 'value': 'Colombia', 'label': 'Colombia', 'label_zh': '哥伦比亚', 'phone_code': '57', 'language': 'es', 'enabled': True, 'sort_order': 33, 'aliases': ['Colombia', '哥伦比亚', 'CO']},
    {'code': 'PH', 'value': 'Philippines', 'label': 'Philippines', 'label_zh': '菲律宾', 'phone_code': '63', 'language': 'en', 'enabled': False, 'sort_order': 40, 'aliases': ['Philippines', '菲律宾', 'PH']},
    {'code': 'MY', 'value': 'Malaysia', 'label': 'Malaysia', 'label_zh': '马来西亚', 'phone_code': '60', 'language': 'ms', 'enabled': False, 'sort_order': 50, 'aliases': ['Malaysia', '马来西亚', 'MY']},
    {'code': 'SG', 'value': 'Singapore', 'label': 'Singapore', 'label_zh': '新加坡', 'phone_code': '65', 'language': 'en', 'enabled': False, 'sort_order': 60, 'aliases': ['Singapore', '新加坡', 'SG']},
    {'code': 'HK', 'value': 'Hong Kong', 'label': 'Hong Kong', 'label_zh': '中国香港', 'phone_code': '852', 'language': 'zh', 'enabled': False, 'sort_order': 70, 'aliases': ['Hong Kong', '中国香港', '香港', 'HK']},
]


def _mcn_region_options(*, include_disabled: bool = False) -> list[dict[str, Any]]:
    rows = []
    for item in sorted(MCN_INTERNAL_REGION_OPTIONS, key=lambda row: int(row.get('sort_order') or 999)):
        if not include_disabled and not bool(item.get('enabled')):
            continue
        rows.append(dict(item))
    return rows


def _mcn_region_by_value(value: Any) -> Optional[dict[str, Any]]:
    normalized = str(value or '').strip().lower()
    if not normalized:
        return None
    for item in MCN_INTERNAL_REGION_OPTIONS:
        candidates = [item.get('code'), item.get('value'), item.get('label'), item.get('label_zh'), *(item.get('aliases') or [])]
        if any(str(candidate or '').strip().lower() == normalized for candidate in candidates):
            return dict(item)
    return None


def _enrich_mcn_region_option(value: Any) -> dict[str, Any]:
    raw = str(value or '').strip()
    region = _mcn_region_by_value(raw)
    if region:
        return {
            'value': str(region.get('value') or raw),
            'label': str(region.get('label') or region.get('value') or raw),
            'code': str(region.get('code') or ''),
            'label_zh': str(region.get('label_zh') or ''),
            'phone_code': str(region.get('phone_code') or ''),
            'language': str(region.get('language') or ''),
            'enabled': bool(region.get('enabled')),
        }
    return {'value': raw, 'label': raw, 'code': '', 'label_zh': raw, 'phone_code': '', 'language': '', 'enabled': True}


def _canonical_mcn_region_value(value: Any) -> str:
    raw = str(value or '').strip()
    region = _mcn_region_by_value(raw)
    if region:
        return str(region.get('value') or raw).strip()
    return raw


def _mcn_language_for_region(region: Any, *, default: str = '') -> str:
    matched = _mcn_region_by_value(region)
    if matched:
        return str(matched.get('language') or '').strip().lower()
    value = str(region or '').strip().lower()
    fallback = {
        '印尼': 'id',
        'indonesia': 'id',
        'indo': 'id',
        'id': 'id',
        '巴西': 'pt',
        'brazil': 'pt',
        'br': 'pt',
        '墨西哥': 'es',
        'mexico': 'es',
        'mx': 'es',
        '委内瑞拉': 'es',
        'venezuela': 'es',
        've': 'es',
        '智利': 'es',
        'chile': 'es',
        'cl': 'es',
        '哥伦比亚': 'es',
        'colombia': 'es',
        'co': 'es',
        '菲律宾': 'en',
        'philippines': 'en',
        'ph': 'en',
        '马来西亚': 'ms',
        'malaysia': 'ms',
        'my': 'ms',
        '新加坡': 'en',
        'singapore': 'en',
        'sg': 'en',
        '中国香港': 'zh',
        '香港': 'zh',
        'hong kong': 'zh',
        'hk': 'zh',
    }
    return fallback.get(value, str(default or '').strip().lower())


def _normalize_translation_source_language(language: Any = '', region: Any = '') -> str:
    value = str(language or '').strip().lower().replace('_', '-')
    language_aliases = {
        'indonesian': 'id',
        'bahasa indonesia': 'id',
        'indo': 'id',
        'portuguese': 'pt',
        'português': 'pt',
        'portugues': 'pt',
        'brazilian portuguese': 'pt',
        'spanish': 'es',
        'español': 'es',
        'espanol': 'es',
        'english': 'en',
        'malay': 'ms',
        'bahasa melayu': 'ms',
        'chinese': 'zh',
        '中文': 'zh',
        'auto': 'auto',
    }
    if value in language_aliases:
        return language_aliases[value]
    if re.fullmatch(r'[a-z]{2,3}(?:-[a-z]{2})?', value):
        return value.split('-', 1)[0]
    resolved = _mcn_language_for_region(region)
    return resolved or 'auto'


WHATSAPP_APPROVAL_DEFAULT_AREA_OPTIONS: list[dict[str, Any]] = [
    _enrich_mcn_region_option(item['value']) for item in _mcn_region_options(include_disabled=False)
]


def _shared_group_approval_target_label(*, approval_scope: str, result: Optional[Dict[str, Any]] = None, registration_group: Optional[str] = None, target_group: Optional[str] = None) -> str:
    payload = result if isinstance(result, dict) else {}
    details = payload.get('details') if isinstance(payload.get('details'), dict) else {}
    candidates = [
        details.get('target_group_label'),
        payload.get('target_group_label'),
        details.get('group_name'),
        details.get('target_group'),
        payload.get('group_name'),
        payload.get('target_group'),
        target_group,
        registration_group,
    ]
    for candidate in candidates:
        value = str(candidate or '').strip()
        if value:
            return value
    return str(approval_scope or '').strip()


def _with_shared_group_approval_result(result: Dict[str, Any], *, approval_scope: str, registration_group: Optional[str] = None, target_group: Optional[str] = None) -> Dict[str, Any]:
    normalized = dict(result or {})
    normalized['approval_scope'] = str(approval_scope or '').strip()
    normalized['target_group_label'] = _shared_group_approval_target_label(
        approval_scope=approval_scope,
        result=normalized,
        registration_group=registration_group,
        target_group=target_group,
    )
    return normalized


def _with_shared_group_approval_executor_result(result: Dict[str, Any], *, approval_scope: str, target_group: Optional[str] = None) -> Dict[str, Any]:
    normalized = dict(result or {})
    normalized['approval_scope'] = str(approval_scope or '').strip()
    if target_group is not None:
        normalized['target_group_label'] = str(target_group or '').strip()
    return normalized


def _normalize_area_options(options: list[str]) -> list[dict[str, Any]]:
    normalized: list[dict[str, Any]] = []
    seen: set[str] = set()
    for item in options:
        value = str(item or '').strip()
        if not value:
            continue
        enriched = _enrich_mcn_region_option(value)
        key = str(enriched.get('value') or value).strip()
        if key in seen:
            continue
        seen.add(key)
        normalized.append(enriched)
    return normalized


def _normalize_whatsapp_group_invite_link(value: Any) -> str:
    raw = str(value or '').strip()
    if not raw:
        return ''
    match = re.search(r'https://chat\.whatsapp\.com/([A-Za-z0-9_-]+)', raw)
    if match:
        return f'https://chat.whatsapp.com/{match.group(1)}'
    return raw


def _looks_like_whatsapp_group_jid(value: Any) -> bool:
    return bool(re.fullmatch(r'[^\s@]+@g\.us', str(value or '').strip()))


def _sanitize_whatsapp_group_jid(value: Any) -> str:
    candidate = str(value or '').strip()
    return candidate if _looks_like_whatsapp_group_jid(candidate) else ''


def _looks_like_whatsapp_invite_link(value: Any) -> bool:
    return str(value or '').strip().startswith('https://chat.whatsapp.com/')


def _fetch_whatsapp_invite_page_group_name(link: Any, *, timeout_seconds: float = 6.0) -> str:
    normalized_link = _normalize_whatsapp_group_invite_link(link)
    if not normalized_link.startswith('https://chat.whatsapp.com/'):
        return ''
    try:
        response = requests.get(
            normalized_link,
            headers={'User-Agent': 'Mozilla/5.0'},
            timeout=timeout_seconds,
        )
        response.raise_for_status()
    except Exception:
        return ''
    text = response.text or ''
    patterns = [
        r'<meta\s+property=["\']og:title["\']\s+content=["\']([^"\']+)["\']',
        r'<meta\s+content=["\']([^"\']+)["\']\s+property=["\']og:title["\']',
        r'<title[^>]*>(.*?)</title>',
    ]
    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE | re.DOTALL)
        if not match:
            continue
        title = html.unescape(re.sub(r'\s+', ' ', str(match.group(1) or '')).strip())
        if title and title != 'WhatsApp Group Invite' and not _looks_like_whatsapp_group_jid(title):
            return title
    return ''


def _whatsapp_approval_binding_config_fingerprint(binding: Dict[str, Any]) -> str:
    item = dict(binding or {})
    payload = {
        'link': _normalize_whatsapp_group_invite_link(item.get('link')),
        'area': str(item.get('area') or '').strip(),
        'notify_profile_name': str(item.get('notify_profile_name') or '').strip(),
        'enabled': False if item.get('enabled') is False else True,
        'approval_count_threshold': int(_coerce_positive_int(item.get('approval_count_threshold'), WHATSAPP_APPROVAL_DEFAULT_COUNT_THRESHOLD) or 0),
        'approval_timeout_minutes': int(_coerce_positive_int(item.get('approval_timeout_minutes'), WHATSAPP_APPROVAL_DEFAULT_TIMEOUT_MINUTES) or 0),
        'auto_recover_worker': bool(item.get('auto_recover_worker')) if item.get('auto_recover_worker') is not None else True,
        'schedule_windows': _normalize_schedule_windows_payload(item.get('schedule_windows') if isinstance(item.get('schedule_windows'), list) else []),
        'provider_mode': resolve_whatsapp_approval_provider_mode(binding=item, responsible_type=item.get('responsible_type')),
        'registration_group_runtime': str(item.get('registration_group_runtime') or '').strip().lower(),
        'official_group_runtime': str(item.get('official_group_runtime') or '').strip().lower(),
        'group_assistant_runtime': str(item.get('group_assistant_runtime') or '').strip().lower(),
        'baileys_base_url': str(item.get('baileys_base_url') or item.get('provider_base_url') or '').strip().rstrip('/'),
        'baileys_account_id': str(item.get('baileys_account_id') or item.get('provider_account_id') or item.get('account_id') or '').strip(),
        'provider_capabilities': item.get('provider_capabilities') if isinstance(item.get('provider_capabilities'), dict) else {},
        'baileys_enabled': bool(item.get('baileys_enabled')) if item.get('baileys_enabled') is not None else None,
    }
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(',', ':'))
    return hashlib.sha256(raw.encode('utf-8')).hexdigest()[:16]


def _new_whatsapp_approval_binding_id() -> str:
    return f"wabind_{uuid.uuid4().hex[:16]}"


def _whatsapp_approval_binding_lookup_keys(binding: Dict[str, Any]) -> list[tuple[str, str]]:
    item = dict(binding or {})
    keys: list[tuple[str, str]] = []
    binding_id = str(item.get('binding_id') or '').strip()
    if binding_id:
        keys.append(('binding_id', binding_id))
    link = _normalize_whatsapp_group_invite_link(item.get('link'))
    if link:
        keys.append(('link', link))
    raw_group_id = _sanitize_whatsapp_group_jid(item.get('group_id')) or _sanitize_whatsapp_group_jid(item.get('registration_group'))
    if raw_group_id:
        keys.append(('group_id', raw_group_id))
    area = str(item.get('area') or '').strip()
    if link and area:
        keys.append(('link_area', f"{link}::{area}"))
    return keys


def _whatsapp_approval_runtime_config_from_dict(item: Optional[Dict[str, Any]]) -> Dict[str, Any]:
    source = dict(item or {})
    account_id = first_baileys_account_id(source)
    base_url = str(source.get('baileys_base_url') or source.get('provider_base_url') or '').strip().rstrip('/')
    token = str(source.get('baileys_token') or source.get('provider_token') or source.get('runtime_token') or '').strip()
    config = {
        'provider_mode': str(source.get('provider_mode') or '').strip().lower(),
        'registration_group_runtime': str(source.get('registration_group_runtime') or '').strip().lower(),
        'official_group_runtime': str(source.get('official_group_runtime') or '').strip().lower(),
        'group_assistant_runtime': str(source.get('group_assistant_runtime') or '').strip().lower(),
        'baileys_base_url': base_url,
        'provider_base_url': str(source.get('provider_base_url') or base_url).strip().rstrip('/'),
        'baileys_token': token,
        'provider_token': str(source.get('provider_token') or token).strip(),
        'runtime_token': str(source.get('runtime_token') or token).strip(),
        'baileys_account_id': account_id,
        'provider_account_id': str(source.get('provider_account_id') or account_id).strip(),
        'account_id': str(source.get('account_id') or account_id).strip(),
        'provider_capabilities': dict(source.get('provider_capabilities') or {}) if isinstance(source.get('provider_capabilities'), dict) else {},
        'baileys_enabled': source.get('baileys_enabled') if source.get('baileys_enabled') is not None else None,
    }
    return config


def _merge_whatsapp_approval_runtime_configs(*configs: Optional[Dict[str, Any]]) -> Dict[str, Any]:
    merged: Dict[str, Any] = {}
    for config in configs:
        source = dict(config or {})
        for key, value in source.items():
            if key == 'provider_capabilities':
                if isinstance(value, dict) and value:
                    merged[key] = dict(value)
                continue
            if key == 'baileys_enabled':
                if value is not None:
                    merged[key] = value
                continue
            if str(value or '').strip():
                merged[key] = value
    return _whatsapp_approval_runtime_config_from_dict(merged)


def _explicit_whatsapp_runtime_mode(*sources: Optional[Dict[str, Any]]) -> str:
    for source in sources:
        if not isinstance(source, dict):
            continue
        for key in RUNTIME_MODE_KEYS:
            value = str(source.get(key) or '').strip().lower()
            if value:
                return value
    return ''


def _default_baileys_account_id_for_whatsapp_account(account_key: str) -> str:
    return default_baileys_account_id_for_account_key(account_key)


def _normalize_group_link_bindings(bindings: list[dict[str, Any]], *, responsible_type: str = '') -> list[dict[str, Any]]:
    normalized: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()
    for item in bindings or []:
        if not isinstance(item, dict):
            continue
        link = _normalize_whatsapp_group_invite_link(item.get('link'))
        area = str(item.get('area') or '').strip()
        if not link:
            continue
        key = (link, area)
        if key in seen:
            continue
        seen.add(key)
        raw_group_name = str(item.get('group_name') or '').strip()
        raw_registration_group = str(item.get('registration_group') or '').strip()
        raw_group_id = str(item.get('group_id') or '').strip()
        sanitized_group_id = _sanitize_whatsapp_group_jid(raw_group_id)
        if not sanitized_group_id and raw_group_id and not _looks_like_whatsapp_invite_link(raw_group_id):
            sanitized_group_id = raw_group_id
        sanitized_registration_group = '' if _looks_like_whatsapp_invite_link(raw_registration_group) else raw_registration_group
        if not sanitized_registration_group and sanitized_group_id:
            sanitized_registration_group = sanitized_group_id
        sanitized_group_name = '' if _looks_like_whatsapp_invite_link(raw_group_name) else raw_group_name
        row = {
            'binding_id': str(item.get('binding_id') or '').strip(),
            'link': link,
            'group_name': sanitized_group_name,
            'area': area,
            'notify_profile_name': str(item.get('notify_profile_name') or '').strip(),
            'enabled': False if item.get('enabled') is False else True,
            'registration_group': sanitized_registration_group,
            'group_id': sanitized_group_id,
            'approval_count_threshold': item.get('approval_count_threshold'),
            'approval_timeout_minutes': item.get('approval_timeout_minutes'),
            'auto_recover_worker': item.get('auto_recover_worker'),
            'schedule_windows': item.get('schedule_windows') if isinstance(item.get('schedule_windows'), list) else [],
            'provider_mode': str(item.get('provider_mode') or '').strip().lower(),
            'registration_group_runtime': str(item.get('registration_group_runtime') or '').strip().lower(),
            'official_group_runtime': str(item.get('official_group_runtime') or '').strip().lower(),
            'group_assistant_runtime': str(item.get('group_assistant_runtime') or '').strip().lower(),
            'baileys_base_url': str(item.get('baileys_base_url') or item.get('provider_base_url') or '').strip().rstrip('/'),
            'provider_base_url': str(item.get('provider_base_url') or item.get('baileys_base_url') or '').strip().rstrip('/'),
            'baileys_token': str(item.get('baileys_token') or item.get('provider_token') or item.get('runtime_token') or '').strip(),
            'provider_token': str(item.get('provider_token') or item.get('baileys_token') or item.get('runtime_token') or '').strip(),
            'runtime_token': str(item.get('runtime_token') or item.get('baileys_token') or item.get('provider_token') or '').strip(),
            'baileys_account_id': str(item.get('baileys_account_id') or item.get('provider_account_id') or item.get('account_id') or '').strip(),
            'provider_account_id': str(item.get('provider_account_id') or item.get('baileys_account_id') or item.get('account_id') or '').strip(),
            'account_id': str(item.get('account_id') or item.get('baileys_account_id') or item.get('provider_account_id') or '').strip(),
            'provider_capabilities': dict(item.get('provider_capabilities') or {}) if isinstance(item.get('provider_capabilities'), dict) else {},
            'baileys_enabled': item.get('baileys_enabled'),
        }
        for key in (
            'identity_status', 'identity_rebuild_reason', 'identity_resolved_at', 'identity_resolved_by',
            'last_probe_status', 'last_probe_reason', 'last_probe_at', 'last_probe_had_group_id',
            'last_probe_had_group_name', 'last_probe_self_participant_found', 'last_probe_self_is_admin',
            'last_probe_can_manage_membership_requests', 'last_probe_member_count', 'runtime_probe_group_id',
            'runtime_probe_group_name', 'queue_status', 'queue_confidence', 'previous_verified_group_id',
            'previous_verified_group_name', 'previous_verified_registration_group',
            'shadow_pending_count', 'shadow_requester_ids', 'shadow_reason_code', 'shadow_checked_at',
        ):
            if key in item:
                row[key] = item.get(key)
        normalized.append(row)
    return normalized


def _preferred_group_binding(bindings: list[dict[str, Any]]) -> dict[str, Any]:
    rows = [dict(item or {}) for item in (bindings or []) if isinstance(item, dict)]
    if not rows:
        return {}
    for item in rows:
        if item.get('enabled') is not False:
            return item
    return rows[0]


WHATSAPP_APPROVAL_AREA_OPTIONS: list[dict[str, str]] = _normalize_area_options(
    [item['value'] for item in WHATSAPP_APPROVAL_DEFAULT_AREA_OPTIONS]
)
WHATSAPP_APPROVAL_AREA_VALUES: set[str] = {item['value'] for item in WHATSAPP_APPROVAL_AREA_OPTIONS}


class ApprovalBatchEvaluateRequest(BaseModel):
    approval_type: str
    registration_group: str
    pending_count: int
    oldest_pending_at: Optional[str] = None
    now: str
    batch_size: Optional[int] = None
    timeout_minutes: Optional[int] = None
    cycle_anchor_at: Optional[str] = None


class OpsAuthBootstrapRequest(BaseModel):
    username: str
    password: str
    display_name: Optional[str] = None


class OpsAuthLoginRequest(BaseModel):
    username: str
    password: str


class OpsAccountCreateRequest(BaseModel):
    username: str
    password: str
    role: str = 'operator'
    display_name: Optional[str] = None
    enabled: bool = True


class OpsAccountUpdateRequest(BaseModel):
    role: Optional[str] = None
    display_name: Optional[str] = None
    enabled: Optional[bool] = None
    password: Optional[str] = None


class OpsPasswordChangeRequest(BaseModel):
    current_password: str
    new_password: str


class GroupAtmosphereTemplate(BaseModel):
    template_id: Optional[str] = None
    category: Optional[str] = None
    text: str
    candidate_id: Optional[str] = None
    media_id: Optional[str] = None
    media_path: Optional[str] = None
    media_mime_type: Optional[str] = None
    media_filename: Optional[str] = None
    asset_type: Optional[str] = None
    moved_from_role_key: Optional[str] = None
    moved_from_template_id: Optional[str] = None
    moved_to_role_key: Optional[str] = None
    moved_at: Optional[str] = None
    role_positioning: Optional[str] = None
    phrase_type: Optional[str] = None
    region: Optional[str] = None
    language: Optional[str] = None
    role_selected: bool = False
    role_send_enabled: bool = False
    role_selection_order: Optional[int] = None
    source_role: Optional[str] = None
    source_type: Optional[str] = None
    text_zh: Optional[str] = None
    text_zh_source: Optional[str] = None
    text_zh_status: Optional[str] = None
    text_zh_updated_at: Optional[str] = None
    text_zh_failure_reason: Optional[str] = None
    text_zh_retry_count: int = 0
    customized: bool = False
    customized_at: Optional[str] = None
    score: Optional[int] = None
    frequency: int = 1
    quality_decision: Optional[str] = None
    quality_status: Optional[str] = None
    quality_score: Optional[int] = None
    quality_reasons: list[str] = []
    normalized_key: Optional[str] = None
    semantic_key: Optional[str] = None
    safe_to_send: bool = True
    enabled: bool = True


class GroupAtmosphereFaqRule(BaseModel):
    keyword: str
    reply: str


class GroupAtmosphereAccountGroupRequest(BaseModel):
    target_group: str
    group_name: Optional[str] = None
    enabled: bool = True
    daily_max_messages: Optional[int] = None
    min_interval_seconds: Optional[int] = None
    max_interval_seconds: Optional[int] = None
    min_interval_minutes: Optional[int] = None
    max_interval_minutes: Optional[int] = None
    allowed_windows: list[dict[str, Any]] = []
    language: Optional[str] = None
    speech_plan_config_name: Optional[str] = None


class GroupAtmosphereWhatsAppAccountRequest(BaseModel):
    account_key: Optional[str] = None
    account_name: Optional[str] = None
    baileys_account_id: Optional[str] = None
    provider_account_id: Optional[str] = None
    account_id: Optional[str] = None
    region: Optional[str] = None
    language: Optional[str] = None
    role_positioning: Optional[str] = None
    speaking_style: Optional[str] = None
    randomness_level: Optional[str] = None
    daily_max_messages: Optional[int] = None
    min_interval_seconds: Optional[int] = None
    max_interval_seconds: Optional[int] = None
    min_interval_minutes: Optional[int] = None
    max_interval_minutes: Optional[int] = None
    allowed_windows: list[dict[str, Any]] = []
    target_group: Optional[str] = None
    group_name: Optional[str] = None
    groups: list[GroupAtmosphereAccountGroupRequest] = []
    enabled: bool = True


class GroupAtmosphereConfigRequest(BaseModel):
    config_name: str
    enabled: bool = True
    account_key: str
    target_group: str
    group_name: Optional[str] = None
    language: str = 'en'
    timezone: str = 'UTC'
    worker_base_url: Optional[str] = None
    daily_max_messages: int = Field(default=4, ge=0, le=10000)
    min_interval_seconds: Optional[int] = Field(default=None, ge=0, le=86400)
    max_interval_seconds: Optional[int] = Field(default=None, ge=0, le=86400)
    min_interval_minutes: Optional[int] = Field(default=None, ge=0, le=86400)
    max_interval_minutes: Optional[int] = Field(default=None, ge=0, le=86400)
    allowed_windows: List[Dict[str, Any]] = Field(default_factory=list)
    template_pool: List[GroupAtmosphereTemplate] = Field(default_factory=list)
    mention_reply_enabled: bool = True
    faq_rules: List[GroupAtmosphereFaqRule] = Field(default_factory=list)
    status: Optional[str] = None


class GroupAtmosphereDispatchRequest(BaseModel):
    config_name: str
    message_text: Optional[str] = None
    trigger_type: str = 'manual'
    client_send_key: Optional[str] = None
    scheduled_at: Optional[str] = None


class GroupAtmosphereManualSendRequest(BaseModel):
    message_text: Optional[str] = ''
    trigger_type: str = 'manual_ops_page'
    media_id: Optional[str] = None
    media_path: Optional[str] = None
    media_mime_type: Optional[str] = None
    media_filename: Optional[str] = None
    client_send_key: Optional[str] = None
    scheduled_at: Optional[str] = None


class GroupAtmosphereInboundMessageRequest(BaseModel):
    account_key: str
    target_group: str
    sender_id: Optional[str] = None
    text: str = ''
    mentioned: bool = False
    quoted_own_message: bool = False


class GroupAtmosphereTriggerEventRequest(BaseModel):
    account_key: str
    target_group: str
    trigger_type: str
    sender_id: Optional[str] = None
    event_payload: Dict[str, Any] = Field(default_factory=dict)


class GroupAtmosphereChatRecord(BaseModel):
    sender: Optional[str] = None
    text: str
    created_at: Optional[str] = None
    message_id: Optional[str] = None


class GroupAtmosphereImportChatRecordsRequest(BaseModel):
    config_name: str
    records: List[GroupAtmosphereChatRecord] = Field(default_factory=list)


class GroupAtmosphereAiCandidateRequest(BaseModel):
    config_name: str
    topic: str = 'general'
    count: int = Field(default=3, ge=1, le=100)


class GroupAtmosphereCandidateEnableRequest(BaseModel):
    config_name: str
    candidate_ids: List[str] = Field(default_factory=list)
    account_key: Optional[str] = None
    target_group: Optional[str] = None
    group_name: Optional[str] = None
    worker_base_url: Optional[str] = None
    daily_max_messages: Optional[int] = Field(default=None, ge=0, le=10000)
    min_interval_seconds: Optional[int] = Field(default=None, ge=0, le=86400)
    max_interval_seconds: Optional[int] = Field(default=None, ge=0, le=86400)
    min_interval_minutes: Optional[int] = Field(default=None, ge=0, le=86400)
    max_interval_minutes: Optional[int] = Field(default=None, ge=0, le=86400)


class GroupAtmosphereCandidateCustomRequest(BaseModel):
    config_name: str
    text: str = Field(..., min_length=1, max_length=1500)
    candidate_id: Optional[str] = None
    role_positioning: Optional[str] = None
    region: Optional[str] = None
    language: Optional[str] = None
    media_id: Optional[str] = None
    remove_media: bool = False


class GroupAtmosphereCandidateTranslateRequest(BaseModel):
    config_name: str
    candidate_id: str
    text_zh: Optional[str] = Field(default=None, max_length=1500)
    force: bool = False


class GroupAtmosphereSchedulerRunRequest(BaseModel):
    limit: int = Field(default=20, ge=1, le=100)


class GroupAtmosphereSimulatedInboundMessage(BaseModel):
    sender_id: Optional[str] = None
    text: str = ''
    mentioned: bool = False
    quoted_own_message: bool = False


class GroupAtmosphereSimulationRequest(BaseModel):
    config_name: str
    scenario: str = 'full_stage_4'
    inbound_messages: List[GroupAtmosphereSimulatedInboundMessage] = Field(default_factory=list)


class Database:
    def __init__(self, db_path: str) -> None:
        self.db_path = db_path
        self._memory_conn: Optional[sqlite3.Connection] = None
        self._ensure_parent()
        if self.db_path != ":memory:":
            ensure_sqlite_ready(self.db_path, profile="online")
        self._init_schema()

    def _ensure_parent(self) -> None:
        if self.db_path != ":memory:":
            Path(self.db_path).parent.mkdir(parents=True, exist_ok=True)

    def _configure_connection(self, conn: sqlite3.Connection) -> None:
        conn.row_factory = sqlite3.Row
        conn.execute(f'PRAGMA busy_timeout = {sqlite_busy_timeout_ms("online")}')
        if self.db_path != ":memory:":
            conn.execute('PRAGMA synchronous=NORMAL')

    def connect(self) -> sqlite3.Connection:
        if self.db_path == ":memory:":
            if self._memory_conn is None:
                self._memory_conn = connect_observed_sqlite(
                    ":memory:",
                    source='app.main:memory',
                    check_same_thread=False,
                    timeout=30.0,
                )
                self._configure_connection(self._memory_conn)
            return self._memory_conn
        conn = connect_observed_sqlite(self.db_path, source='app.main', timeout=30.0)
        self._configure_connection(conn)
        return conn

    def _init_schema(self) -> None:
        conn = self.connect()
        conn.executescript(
                """
                CREATE TABLE IF NOT EXISTS leads (
                    lead_id TEXT PRIMARY KEY,
                    trace_id TEXT NOT NULL,
                    source_platform TEXT NOT NULL,
                    source_campaign TEXT,
                    source_page_id TEXT NOT NULL,
                    country TEXT NOT NULL,
                    assigned_guild_country TEXT NOT NULL DEFAULT '',
                    cross_country_fallback INTEGER NOT NULL DEFAULT 0,
                    cross_country_fallback_reason TEXT NOT NULL DEFAULT '',
                    area_code INTEGER NOT NULL,
                    mobile TEXT NOT NULL,
                    yw_id TEXT,
                    app_name TEXT,
                    dept_name TEXT,
                    pendaftaran_group TEXT,
                    inviter_id TEXT,
                    parser_confidence REAL,
                    parser_missing_fields TEXT NOT NULL DEFAULT '[]',
                    parser_conflicts TEXT NOT NULL DEFAULT '[]',
                    parser_raw_text TEXT,
                    parser_raw_ocr_text TEXT,
                    parser_version TEXT NOT NULL DEFAULT 'manual_cs_parser_v2',
                    parser_status TEXT NOT NULL DEFAULT 'unknown',
                    review_reason_codes TEXT NOT NULL DEFAULT '[]',
                    routing_decision TEXT,
                    recommended_next_action TEXT,
                    review_status TEXT NOT NULL DEFAULT 'not_needed',
                    review_notes TEXT,
                    reviewed_by TEXT,
                    reviewed_at TEXT,
                    correction_count INTEGER NOT NULL DEFAULT 0,
                    current_status TEXT NOT NULL,
                    matched_customer_id TEXT NOT NULL,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL,
                    UNIQUE(area_code, mobile)
                );

                CREATE TABLE IF NOT EXISTS customer_projection (
                    customer_id TEXT PRIMARY KEY,
                    lead_id TEXT NOT NULL,
                    mobile TEXT NOT NULL,
                    area_code INTEGER NOT NULL,
                    yw_id TEXT,
                    pendaftaran_group TEXT,
                    payment_status TEXT,
                    user_quality TEXT,
                    remark TEXT,
                    join_group TEXT,
                    file_url TEXT,
                    pz_status INTEGER,
                    updated_at TEXT NOT NULL,
                    UNIQUE(area_code, mobile)
                );

                CREATE TABLE IF NOT EXISTS lead_events (
                    event_id TEXT PRIMARY KEY,
                    lead_id TEXT,
                    trace_id TEXT NOT NULL,
                    event_type TEXT NOT NULL,
                    event_source TEXT NOT NULL,
                    event_value TEXT,
                    page_id TEXT,
                    session_id TEXT,
                    operator_id TEXT,
                    operator_name TEXT,
                    raw_payload TEXT NOT NULL,
                    happened_at TEXT NOT NULL,
                    created_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS automation_tasks (
                    task_id TEXT PRIMARY KEY,
                    lead_id TEXT NOT NULL,
                    task_type TEXT NOT NULL,
                    priority TEXT NOT NULL,
                    payload TEXT NOT NULL,
                    dedupe_key TEXT NOT NULL UNIQUE,
                    created_by TEXT NOT NULL,
                    created_at TEXT NOT NULL,
                    status TEXT NOT NULL,
                    result_code TEXT,
                    result_reason TEXT,
                    toast_text TEXT,
                    evidence_url TEXT,
                    retry_count INTEGER NOT NULL DEFAULT 0,
                    executor_type TEXT,
                    executor_id TEXT,
                    started_at TEXT,
                    finished_at TEXT,
                    raw_result TEXT NOT NULL DEFAULT '{}',
                    requires_human_action INTEGER NOT NULL DEFAULT 0,
                    human_action_type TEXT,
                    human_action_payload TEXT,
                    worker_id TEXT NOT NULL DEFAULT '',
                    lease_until TEXT NOT NULL DEFAULT '',
                    heartbeat_at TEXT NOT NULL DEFAULT ''
                );

                CREATE TABLE IF NOT EXISTS sync_logs (
                    sync_log_id TEXT PRIMARY KEY,
                    lead_id TEXT,
                    task_id TEXT,
                    sync_type TEXT NOT NULL,
                    target_system TEXT NOT NULL,
                    status TEXT NOT NULL,
                    request_snapshot TEXT NOT NULL,
                    response_snapshot TEXT NOT NULL,
                    created_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS account_submissions (
                    submission_id TEXT PRIMARY KEY,
                    lead_id TEXT NOT NULL,
                    task_id TEXT,
                    submission_type TEXT NOT NULL,
                    account_id TEXT,
                    account_id_type TEXT,
                    file_url TEXT,
                    file_type TEXT,
                    source_channel TEXT,
                    submitted_by TEXT,
                    recognition_status TEXT NOT NULL DEFAULT 'not_needed',
                    recognized_account_id TEXT,
                    recognition_raw TEXT NOT NULL DEFAULT '{}',
                    submitted_at TEXT NOT NULL,
                    remark TEXT,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS bind_check_jobs (
                    job_id TEXT PRIMARY KEY,
                    lead_id TEXT NOT NULL,
                    submission_id TEXT,
                    account_id TEXT NOT NULL,
                    guild_code TEXT,
                    check_source TEXT NOT NULL,
                    status TEXT NOT NULL,
                    result_code TEXT,
                    result_reason TEXT,
                    raw_result TEXT NOT NULL DEFAULT '{}',
                    retry_count INTEGER NOT NULL DEFAULT 0,
                    scheduled_at TEXT NOT NULL,
                    finished_at TEXT,
                    created_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS group_join_jobs (
                    job_id TEXT PRIMARY KEY,
                    lead_id TEXT NOT NULL,
                    submission_id TEXT,
                    account_id TEXT,
                    target_group TEXT,
                    join_type TEXT NOT NULL,
                    status TEXT NOT NULL,
                    result_code TEXT,
                    result_reason TEXT,
                    raw_result TEXT NOT NULL DEFAULT '{}',
                    retry_count INTEGER NOT NULL DEFAULT 0,
                    scheduled_at TEXT NOT NULL,
                    finished_at TEXT,
                    created_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS guild_executors (
                    guild_name TEXT PRIMARY KEY,
                    app_name TEXT NOT NULL DEFAULT 'linky',
                    backend_url TEXT NOT NULL,
                    login_username TEXT NOT NULL,
                    password_secret_ref TEXT,
                    guild_backend_token TEXT,
                    oauth_token TEXT,
                    oauth_token_secret TEXT,
                    platform_backend_url TEXT,
                    platform_authorization TEXT,
                    cms_guild_id TEXT,
                    cms_guild_sid TEXT,
                    country TEXT,
                    guild_country TEXT,
                    eligible_user_countries TEXT NOT NULL DEFAULT '[]',
                    routing_region TEXT NOT NULL DEFAULT '',
                    proxy_url TEXT,
                    proxy_region TEXT,
                    proxy_type TEXT,
                    enabled INTEGER NOT NULL DEFAULT 1,
                    browser_profile_key TEXT,
                    bind_concurrency INTEGER NOT NULL DEFAULT 1,
                    request_timeout_seconds INTEGER NOT NULL DEFAULT 30,
                    notes TEXT,
                    updated_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS cms_executor_tokens (
                    guild_name TEXT PRIMARY KEY,
                    refresh_token TEXT,
                    refresh_token_deadtime INTEGER,
                    access_token_exp INTEGER,
                    last_refresh_at INTEGER,
                    last_refresh_error TEXT,
                    updated_at TEXT DEFAULT CURRENT_TIMESTAMP
                );

                CREATE TABLE IF NOT EXISTS guild_anchor_daily_stats (
                    guild_executor_key TEXT NOT NULL DEFAULT '',
                    guild_name TEXT NOT NULL,
                    guild_display_name TEXT NOT NULL DEFAULT '',
                    stat_date TEXT NOT NULL,
                    joined_count INTEGER NOT NULL DEFAULT 0,
                    real_person_count INTEGER NOT NULL DEFAULT 0,
                    total_anchors INTEGER,
                    scanned_count INTEGER NOT NULL DEFAULT 0,
                    page_count INTEGER NOT NULL DEFAULT 0,
                    sort_direction TEXT NOT NULL DEFAULT '',
                    sort_confidence TEXT NOT NULL DEFAULT '',
                    stale_after TEXT NOT NULL DEFAULT '',
                    status TEXT NOT NULL DEFAULT 'pending',
                    error TEXT NOT NULL DEFAULT '',
                    refreshed_at TEXT NOT NULL,
                    PRIMARY KEY (guild_name, stat_date)
                );

                CREATE TABLE IF NOT EXISTS guild_anchor_daily_stat_jobs (
                    job_id TEXT PRIMARY KEY,
                    guild_executor_key TEXT NOT NULL,
                    guild_name TEXT NOT NULL,
                    stat_date TEXT NOT NULL,
                    status TEXT NOT NULL DEFAULT 'pending',
                    attempt_count INTEGER NOT NULL DEFAULT 0,
                    max_attempts INTEGER NOT NULL DEFAULT 5,
                    next_retry_at TEXT NOT NULL DEFAULT '',
                    lease_owner TEXT NOT NULL DEFAULT '',
                    lease_until TEXT NOT NULL DEFAULT '',
                    started_at TEXT NOT NULL DEFAULT '',
                    finished_at TEXT NOT NULL DEFAULT '',
                    scanned_count INTEGER NOT NULL DEFAULT 0,
                    page_count INTEGER NOT NULL DEFAULT 0,
                    total_anchors INTEGER,
                    joined_count INTEGER,
                    real_person_count INTEGER,
                    sort_direction TEXT NOT NULL DEFAULT '',
                    sort_confidence TEXT NOT NULL DEFAULT '',
                    error TEXT NOT NULL DEFAULT '',
                    last_error TEXT NOT NULL DEFAULT '',
                    recovery_count INTEGER NOT NULL DEFAULT 0,
                    last_recovered_at TEXT NOT NULL DEFAULT '',
                    source TEXT NOT NULL DEFAULT 'schedule',
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL,
                    UNIQUE (guild_executor_key, stat_date)
                );
                CREATE INDEX IF NOT EXISTS idx_guild_anchor_daily_stat_jobs_status_retry
                    ON guild_anchor_daily_stat_jobs(status, next_retry_at, lease_until);

                CREATE TABLE IF NOT EXISTS guild_anchor_seen (
                    guild_executor_key TEXT NOT NULL,
                    guild_name TEXT NOT NULL,
                    anchor_id TEXT NOT NULL,
                    streamer_sid TEXT NOT NULL DEFAULT '',
                    streamer_sid_source_contract TEXT NOT NULL DEFAULT '',
                    anchor_name TEXT NOT NULL DEFAULT '',
                    created_at INTEGER NOT NULL,
                    created_date_utc TEXT NOT NULL DEFAULT '',
                    created_date_bj TEXT NOT NULL,
                    is_real_person INTEGER NOT NULL DEFAULT 0,
                    first_seen_at TEXT NOT NULL,
                    last_seen_at TEXT NOT NULL,
                    source_total_anchors INTEGER,
                    source_page INTEGER,
                    source_page_size INTEGER,
                    raw_hash TEXT NOT NULL DEFAULT '',
                    PRIMARY KEY (guild_executor_key, anchor_id)
                );
                CREATE INDEX IF NOT EXISTS idx_guild_anchor_seen_date
                    ON guild_anchor_seen(guild_executor_key, created_date_utc);
                CREATE INDEX IF NOT EXISTS idx_guild_anchor_seen_bj_date_stats
                    ON guild_anchor_seen(created_date_bj, guild_name, is_real_person, last_seen_at);

                CREATE TABLE IF NOT EXISTS guild_anchor_newcomer_identity_snapshots (
                    guild_executor_key TEXT NOT NULL,
                    guild_name TEXT NOT NULL,
                    stat_date TEXT NOT NULL,
                    anchor_id TEXT NOT NULL,
                    streamer_sid TEXT NOT NULL,
                    anchor_name TEXT NOT NULL DEFAULT '',
                    source_created_at INTEGER NOT NULL DEFAULT 0,
                    source_first_seen_at TEXT NOT NULL DEFAULT '',
                    is_real_person INTEGER NOT NULL DEFAULT 0,
                    snapshot_refreshed_at TEXT NOT NULL,
                    recorded_at TEXT NOT NULL,
                    PRIMARY KEY (guild_executor_key, stat_date, anchor_id)
                );
                CREATE INDEX IF NOT EXISTS idx_guild_anchor_newcomer_snapshot_date
                    ON guild_anchor_newcomer_identity_snapshots(stat_date, guild_name);

                CREATE TABLE IF NOT EXISTS guild_anchor_newcomer_snapshot_runs (
                    guild_executor_key TEXT NOT NULL,
                    guild_name TEXT NOT NULL,
                    stat_date TEXT NOT NULL,
                    member_count INTEGER NOT NULL,
                    source_contract TEXT NOT NULL,
                    snapshot_refreshed_at TEXT NOT NULL,
                    recorded_at TEXT NOT NULL,
                    PRIMARY KEY (guild_executor_key, stat_date)
                );
                CREATE INDEX IF NOT EXISTS idx_guild_anchor_newcomer_snapshot_run_date
                    ON guild_anchor_newcomer_snapshot_runs(stat_date, guild_name);

                CREATE TABLE IF NOT EXISTS timo_anchor_exported_ids (
                    guild_executor_key TEXT NOT NULL,
                    guild_name TEXT NOT NULL,
                    export_kind TEXT NOT NULL,
                    timo_id TEXT NOT NULL,
                    first_exported_at TEXT NOT NULL,
                    last_seen_at TEXT NOT NULL DEFAULT '',
                    source_value REAL,
                    source_payload TEXT NOT NULL DEFAULT '{}',
                    PRIMARY KEY (guild_executor_key, export_kind, timo_id)
                );
                CREATE INDEX IF NOT EXISTS idx_timo_anchor_exported_ids_kind
                    ON timo_anchor_exported_ids(export_kind, first_exported_at);

                CREATE TABLE IF NOT EXISTS timo_anchor_export_cache (
                    guild_executor_key TEXT NOT NULL,
                    guild_name TEXT NOT NULL,
                    export_kind TEXT NOT NULL,
                    data_date_bj TEXT NOT NULL,
                    period_type TEXT NOT NULL DEFAULT 'day',
                    timo_id TEXT NOT NULL,
                    anchor_name TEXT NOT NULL DEFAULT '',
                    diamond_amount REAL NOT NULL DEFAULT 0,
                    source_payload TEXT NOT NULL DEFAULT '{}',
                    first_cached_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL,
                    PRIMARY KEY (guild_executor_key, export_kind, timo_id)
                );
                CREATE INDEX IF NOT EXISTS idx_timo_anchor_export_cache_date
                    ON timo_anchor_export_cache(guild_executor_key, export_kind, data_date_bj);

                CREATE TABLE IF NOT EXISTS timo_revenue_export_cache (
                    guild_executor_key TEXT NOT NULL,
                    guild_name TEXT NOT NULL,
                    export_type TEXT NOT NULL,
                    date_from_bj TEXT NOT NULL,
                    date_to_bj TEXT NOT NULL,
                    filename TEXT NOT NULL,
                    file_path TEXT NOT NULL,
                    file_size INTEGER NOT NULL DEFAULT 0,
                    status TEXT NOT NULL DEFAULT 'success',
                    error TEXT NOT NULL DEFAULT '',
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL,
                    PRIMARY KEY (guild_executor_key, export_type, date_from_bj, date_to_bj)
                );

                CREATE TABLE IF NOT EXISTS guild_anchor_scan_markers (
                    guild_executor_key TEXT PRIMARY KEY,
                    guild_name TEXT NOT NULL,
                    last_total_anchors INTEGER NOT NULL DEFAULT 0,
                    last_full_scan_at TEXT NOT NULL DEFAULT '',
                    last_incremental_scan_at TEXT NOT NULL DEFAULT '',
                    marker_anchor_ids_json TEXT NOT NULL DEFAULT '[]',
                    marker_page INTEGER NOT NULL DEFAULT 0,
                    marker_page_size INTEGER NOT NULL DEFAULT 0,
                    sort_confidence TEXT NOT NULL DEFAULT '',
                    updated_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS production_ops_daemon_configs (
                    config_name TEXT PRIMARY KEY,
                    enabled INTEGER NOT NULL DEFAULT 0,
                    registration_group TEXT NOT NULL,
                    api_base_url TEXT NOT NULL,
                    worker_base_url TEXT NOT NULL,
                    interval_seconds REAL NOT NULL DEFAULT 20,
                    notify_chat_id TEXT,
                    area TEXT,
                    remark TEXT,
                    approved_count INTEGER NOT NULL DEFAULT 1,
                    auto_recover_worker INTEGER NOT NULL DEFAULT 1,
                    updated_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS whatsapp_approval_accounts (
                    account_key TEXT PRIMARY KEY,
                    account_name TEXT NOT NULL,
                    responsible_type TEXT NOT NULL,
                    group_links TEXT NOT NULL DEFAULT '[]',
                    area TEXT,
                    notify_profile_name TEXT,
                    approval_rule TEXT NOT NULL DEFAULT 'count_30',
                    approval_count_threshold INTEGER NOT NULL DEFAULT 30,
                    approval_timeout_minutes INTEGER NOT NULL DEFAULT 30,
                    auto_recover_worker INTEGER NOT NULL DEFAULT 1,
                    schedule_windows TEXT NOT NULL DEFAULT '[]',
                    enabled INTEGER NOT NULL DEFAULT 1,
                    verification_status TEXT NOT NULL DEFAULT 'pending_verification',
                    assigned_customer_service_user_id TEXT,
                    assigned_customer_service_username TEXT,
                    assigned_customer_service_display_name TEXT,
                    notes TEXT,
                    created_at TEXT,
                    updated_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS whatsapp_approval_area_options (
                    option_key TEXT PRIMARY KEY,
                    options_json TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS mcn_region_options (
                    code TEXT PRIMARY KEY,
                    enabled INTEGER NOT NULL DEFAULT 1,
                    sort_order INTEGER,
                    updated_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS lead_status_history (
                    history_id TEXT PRIMARY KEY,
                    lead_id TEXT NOT NULL,
                    from_status TEXT,
                    to_status TEXT NOT NULL,
                    trigger_type TEXT NOT NULL,
                    trigger_source TEXT NOT NULL,
                    trigger_event_id TEXT,
                    trigger_task_id TEXT,
                    operator_id TEXT,
                    operator_name TEXT,
                    remark TEXT,
                    created_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS operator_notifications (
                    notification_id TEXT PRIMARY KEY,
                    lead_id TEXT NOT NULL,
                    notification_type TEXT NOT NULL,
                    mobile TEXT NOT NULL,
                    yw_id TEXT,
                    write_result TEXT NOT NULL,
                    reason TEXT,
                    is_read INTEGER NOT NULL DEFAULT 0,
                    read_at TEXT,
                    read_by TEXT,
                    created_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS manual_review_history (
                    review_id TEXT PRIMARY KEY,
                    lead_id TEXT NOT NULL,
                    decision TEXT NOT NULL,
                    reviewed_by TEXT NOT NULL,
                    review_note TEXT,
                    snapshot_before TEXT NOT NULL DEFAULT '{}',
                    snapshot_after TEXT NOT NULL DEFAULT '{}',
                    created_task_id TEXT,
                    submitted_at TEXT NOT NULL,
                    created_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS lead_corrections (
                    correction_id TEXT PRIMARY KEY,
                    lead_id TEXT NOT NULL,
                    field_name TEXT NOT NULL,
                    old_value TEXT,
                    new_value TEXT,
                    corrected_by TEXT NOT NULL,
                    review_id TEXT,
                    created_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS intake_bot_presets (
                    profile_name TEXT PRIMARY KEY,
                    app_id TEXT,
                    robot_name TEXT,
                    default_app TEXT,
                    default_guild TEXT,
                    enabled INTEGER NOT NULL DEFAULT 1,
                    updated_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS crm_option_cache (
                    option_type TEXT NOT NULL,
                    display_name TEXT NOT NULL,
                    row_json TEXT NOT NULL,
                    updated_at TEXT NOT NULL,
                    PRIMARY KEY (option_type, display_name)
                );

                CREATE TABLE IF NOT EXISTS intake_guild_assignees (
                    guild_name TEXT NOT NULL,
                    user_id TEXT NOT NULL,
                    assigned_by TEXT,
                    updated_at TEXT NOT NULL,
                    PRIMARY KEY (guild_name, user_id)
                );

                CREATE TABLE IF NOT EXISTS ops_intake_items (
                    item_id TEXT PRIMARY KEY,
                    guild_name TEXT NOT NULL,
                    submitted_by_user_id TEXT,
                    submitted_by_username TEXT,
                    raw_text TEXT NOT NULL,
                    parsed_phone TEXT,
                    parsed_account_id TEXT,
                    parsed_group TEXT,
                    parsed_code TEXT,
                    parsed_app TEXT,
                    parsed_agency TEXT,
                    system_status TEXT NOT NULL,
                    feedback_status TEXT NOT NULL,
                    reply_text TEXT,
                    result_code TEXT,
                    result_reason TEXT,
                    result_snapshot TEXT NOT NULL DEFAULT '{}',
                    created_at TEXT NOT NULL,
                    processed_at TEXT,
                    feedback_done_at TEXT,
                    feedback_done_by TEXT,
                    template_copied_at TEXT,
                    template_copied_by TEXT,
                    force_feedback_reason TEXT,
                    idempotency_key TEXT,
                    dedupe_route TEXT,
                    route_snapshot TEXT NOT NULL DEFAULT '{}',
                    config_drift TEXT NOT NULL DEFAULT '{}',
                    group_auto_filled INTEGER NOT NULL DEFAULT 0,
                    group_auto_fill_source TEXT,
                    group_auto_fill_confidence TEXT,
                    group_auto_fill_confirmed INTEGER NOT NULL DEFAULT 0,
                    group_auto_fill_confirmed_by TEXT,
                    group_auto_fill_confirmed_at TEXT,
                    source TEXT NOT NULL DEFAULT 'ops_intake_workbench',
                    external_user_id TEXT,
                    external_session_id TEXT,
                    external_message_id TEXT,
                    external_customer_service_id TEXT,
                    external_customer_service_name TEXT,
                    external_payload TEXT NOT NULL DEFAULT '{}'
                );

                CREATE TABLE IF NOT EXISTS external_app_phone_backfill_requests (
                    request_id TEXT PRIMARY KEY,
                    source TEXT NOT NULL,
                    app TEXT NOT NULL DEFAULT '',
                    guild_name TEXT NOT NULL DEFAULT '',
                    linky_account_id TEXT NOT NULL DEFAULT '',
                    phone TEXT NOT NULL DEFAULT '',
                    external_user_id TEXT,
                    external_session_id TEXT,
                    external_message_id TEXT,
                    external_customer_service_id TEXT,
                    external_customer_service_name TEXT,
                    request_payload TEXT NOT NULL DEFAULT '{}',
                    status TEXT NOT NULL DEFAULT 'received',
                    result_code TEXT NOT NULL DEFAULT '',
                    result_reason TEXT NOT NULL DEFAULT '',
                    submission_id TEXT NOT NULL DEFAULT '',
                    result_snapshot TEXT NOT NULL DEFAULT '{}',
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL,
                    processed_at TEXT
                );

                CREATE TABLE IF NOT EXISTS ops_timo_intake_items (
                    item_id TEXT PRIMARY KEY,
                    guild_name TEXT NOT NULL DEFAULT '',
                    mobile TEXT NOT NULL DEFAULT '',
                    timo_id TEXT NOT NULL,
                    group_name TEXT NOT NULL DEFAULT '',
                    app_name TEXT NOT NULL DEFAULT 'Timo',
                    source_text TEXT NOT NULL DEFAULT '',
                    source_channel TEXT NOT NULL DEFAULT 'ops_timo_intake',
                    source TEXT NOT NULL DEFAULT 'ops_timo_intake',
                    external_user_id TEXT,
                    external_session_id TEXT,
                    external_message_id TEXT,
                    external_customer_service_id TEXT,
                    external_customer_service_name TEXT,
                    external_payload TEXT NOT NULL DEFAULT '{}',
                    profile_name TEXT NOT NULL DEFAULT '',
                    submitted_by_user_id TEXT NOT NULL DEFAULT '',
                    submitted_by_username TEXT NOT NULL DEFAULT '',
                    system_status TEXT NOT NULL DEFAULT 'pending_verification',
                    feedback_status TEXT NOT NULL DEFAULT 'not_feedbackable',
                    feedback_done_at TEXT,
                    feedback_done_by TEXT,
                    template_copied_at TEXT,
                    template_copied_by TEXT,
                    timo_verify_status TEXT NOT NULL DEFAULT 'not_checked',
                    timo_result_code TEXT NOT NULL DEFAULT '',
                    timo_result_reason TEXT NOT NULL DEFAULT '',
                    timo_result_snapshot TEXT NOT NULL DEFAULT '{}',
                    timo_verified_at TEXT,
                    crm_sync_status TEXT NOT NULL DEFAULT 'not_started',
                    crm_result_code TEXT NOT NULL DEFAULT '',
                    crm_result_reason TEXT NOT NULL DEFAULT '',
                    crm_payload TEXT NOT NULL DEFAULT '{}',
                    crm_response TEXT NOT NULL DEFAULT '{}',
                    crm_synced_at TEXT,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS ops_intake_binding_history_attempts (
                    attempt_id TEXT PRIMARY KEY,
                    source_type TEXT NOT NULL,
                    source_id TEXT NOT NULL,
                    item_id TEXT NOT NULL DEFAULT '',
                    lead_id TEXT NOT NULL DEFAULT '',
                    task_id TEXT NOT NULL DEFAULT '',
                    dedupe_key TEXT NOT NULL,
                    normalized_phone TEXT NOT NULL DEFAULT '',
                    normalized_phone_digits TEXT NOT NULL DEFAULT '',
                    created_date_bj TEXT NOT NULL DEFAULT '',
                    guild_name TEXT NOT NULL DEFAULT '',
                    submitted_by_user_id TEXT NOT NULL DEFAULT '',
                    submitted_by_username TEXT NOT NULL DEFAULT '',
                    external_customer_service_id TEXT NOT NULL DEFAULT '',
                    external_customer_service_name TEXT NOT NULL DEFAULT '',
                    display_initiator TEXT NOT NULL DEFAULT '',
                    parsed_phone TEXT NOT NULL DEFAULT '',
                    parsed_account_id TEXT NOT NULL DEFAULT '',
                    parsed_group TEXT NOT NULL DEFAULT '',
                    parsed_code TEXT NOT NULL DEFAULT '',
                    parsed_app TEXT NOT NULL DEFAULT '',
                    parsed_agency TEXT NOT NULL DEFAULT '',
                    system_status TEXT NOT NULL DEFAULT '',
                    feedback_status TEXT NOT NULL DEFAULT '',
                    result_code TEXT NOT NULL DEFAULT '',
                    result_reason TEXT NOT NULL DEFAULT '',
                    created_at TEXT NOT NULL,
                    processed_at TEXT NOT NULL DEFAULT '',
                    closure_status TEXT NOT NULL DEFAULT '',
                    closure_reason TEXT NOT NULL DEFAULT '',
                    closure_note TEXT NOT NULL DEFAULT '',
                    current_exception INTEGER NOT NULL DEFAULT 0,
                    is_failure INTEGER NOT NULL DEFAULT 0,
                    is_duplicate INTEGER NOT NULL DEFAULT 0,
                    is_success INTEGER NOT NULL DEFAULT 0,
                    is_closed INTEGER NOT NULL DEFAULT 0,
                    current_truth_json TEXT NOT NULL DEFAULT '{}'
                );

                CREATE TABLE IF NOT EXISTS ops_intake_binding_history_projection_meta (
                    projection_key TEXT PRIMARY KEY,
                    signature TEXT NOT NULL DEFAULT '',
                    refreshed_at TEXT NOT NULL DEFAULT ''
                );

                CREATE TABLE IF NOT EXISTS ingress_events (
                    event_id TEXT PRIMARY KEY,
                    ingress_type TEXT NOT NULL,
                    source_key TEXT,
                    idempotency_key TEXT NOT NULL UNIQUE,
                    payload TEXT NOT NULL,
                    status TEXT NOT NULL,
                    result_snapshot TEXT,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL,
                    processed_at TEXT
                );

                CREATE TABLE IF NOT EXISTS ingress_jobs (
                    job_id TEXT PRIMARY KEY,
                    event_id TEXT NOT NULL,
                    status TEXT NOT NULL,
                    attempt_count INTEGER NOT NULL DEFAULT 0,
                    available_at TEXT NOT NULL,
                    worker_id TEXT NOT NULL DEFAULT '',
                    lease_until TEXT NOT NULL DEFAULT '',
                    heartbeat_at TEXT NOT NULL DEFAULT '',
                    last_error TEXT,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS registration_group_approval_batch_runs (
                    approval_run_id TEXT PRIMARY KEY,
                    sync_log_id TEXT,
                    status TEXT NOT NULL,
                    request_snapshot TEXT NOT NULL,
                    response_snapshot TEXT NOT NULL,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS registration_group_approval_batch_members (
                    member_id TEXT PRIMARY KEY,
                    approval_run_id TEXT NOT NULL,
                    group_type TEXT NOT NULL DEFAULT 'registration_group',
                    registration_group TEXT NOT NULL,
                    registration_group_name TEXT NOT NULL DEFAULT '',
                    lead_id TEXT NOT NULL DEFAULT '',
                    matched_customer_id TEXT NOT NULL DEFAULT '',
                    registration_status_snapshot TEXT NOT NULL DEFAULT '',
                    registration_status_label_snapshot TEXT NOT NULL DEFAULT '',
                    eligibility_source TEXT NOT NULL DEFAULT '',
                    eligibility_snapshot TEXT NOT NULL DEFAULT '',
                    requester_id TEXT NOT NULL DEFAULT '',
                    display_name TEXT NOT NULL DEFAULT '',
                    display_name_source TEXT NOT NULL DEFAULT '',
                    display_name_enhanced_at TEXT,
                    wa_phone_raw TEXT NOT NULL DEFAULT '',
                    wa_phone_normalized TEXT NOT NULL DEFAULT '',
                    requested_at TEXT,
                    approved_at TEXT NOT NULL,
                    batch_index INTEGER NOT NULL DEFAULT 0,
                    repair_last_attempt_at TEXT,
                    repair_last_result TEXT NOT NULL DEFAULT '',
                    repair_next_attempt_at TEXT,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS operator_audit_log (
                    audit_id TEXT PRIMARY KEY,
                    lead_id TEXT,
                    ingress_event_id TEXT,
                    event_type TEXT NOT NULL,
                    event_source TEXT NOT NULL,
                    payload TEXT NOT NULL,
                    created_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS mcn_event_ledger (
                    event_id TEXT PRIMARY KEY,
                    event_type TEXT NOT NULL,
                    object_type TEXT NOT NULL,
                    object_key TEXT NOT NULL,
                    actor_type TEXT NOT NULL DEFAULT 'system',
                    actor_id TEXT NOT NULL DEFAULT '',
                    status TEXT NOT NULL,
                    evidence_level TEXT NOT NULL DEFAULT '',
                    external_id TEXT NOT NULL DEFAULT '',
                    payload_json TEXT NOT NULL DEFAULT '{}',
                    created_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS mcn_truth_snapshots (
                    snapshot_id TEXT PRIMARY KEY,
                    object_type TEXT NOT NULL,
                    object_key TEXT NOT NULL,
                    snapshot_type TEXT NOT NULL,
                    truth_status TEXT NOT NULL,
                    confidence TEXT NOT NULL,
                    confidence_reason TEXT NOT NULL DEFAULT '',
                    facts_json TEXT NOT NULL DEFAULT '{}',
                    source_json TEXT NOT NULL DEFAULT '{}',
                    checked_at TEXT NOT NULL,
                    expires_at TEXT,
                    recommended_action TEXT NOT NULL DEFAULT '',
                    updated_at TEXT NOT NULL,
                    UNIQUE(object_type, object_key, snapshot_type)
                );

                CREATE TABLE IF NOT EXISTS wa_accounts (
                    account_key TEXT PRIMARY KEY,
                    responsible_type TEXT NOT NULL DEFAULT '',
                    provider_name TEXT NOT NULL DEFAULT 'legacy_playwright',
                    provider_mode TEXT NOT NULL DEFAULT 'legacy_only',
                    health_status TEXT NOT NULL DEFAULT 'unknown',
                    runtime_generation INTEGER NOT NULL DEFAULT 0,
                    metadata_json TEXT NOT NULL DEFAULT '{}',
                    updated_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS wa_group_bindings (
                    binding_id TEXT PRIMARY KEY,
                    account_key TEXT NOT NULL DEFAULT '',
                    responsible_type TEXT NOT NULL DEFAULT '',
                    link TEXT NOT NULL DEFAULT '',
                    group_id TEXT NOT NULL DEFAULT '',
                    registration_group TEXT NOT NULL DEFAULT '',
                    group_name TEXT NOT NULL DEFAULT '',
                    identity_status TEXT NOT NULL DEFAULT '',
                    config_fingerprint TEXT NOT NULL DEFAULT '',
                    provider_mode TEXT NOT NULL DEFAULT 'legacy_only',
                    provider_capabilities_json TEXT NOT NULL DEFAULT '{}',
                    binding_json TEXT NOT NULL DEFAULT '{}',
                    updated_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS wa_truth_snapshots (
                    snapshot_id TEXT PRIMARY KEY,
                    binding_id TEXT NOT NULL DEFAULT '',
                    account_key TEXT NOT NULL DEFAULT '',
                    snapshot_type TEXT NOT NULL,
                    truth_status TEXT NOT NULL,
                    trusted_pending_count INTEGER,
                    requester_ids_json TEXT NOT NULL DEFAULT '[]',
                    facts_json TEXT NOT NULL DEFAULT '{}',
                    source_json TEXT NOT NULL DEFAULT '{}',
                    checked_at TEXT NOT NULL,
                    expires_at TEXT,
                    updated_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS wa_runtime_actions (
                    action_id TEXT PRIMARY KEY,
                    account_key TEXT NOT NULL DEFAULT '',
                    binding_id TEXT NOT NULL DEFAULT '',
                    action_type TEXT NOT NULL,
                    provider_name TEXT NOT NULL DEFAULT '',
                    provider_mode TEXT NOT NULL DEFAULT '',
                    status TEXT NOT NULL,
                    request_json TEXT NOT NULL DEFAULT '{}',
                    result_json TEXT NOT NULL DEFAULT '{}',
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS wa_identity_map (
                    identity_key TEXT PRIMARY KEY,
                    provider_name TEXT NOT NULL DEFAULT '',
                    provider_requester_id TEXT NOT NULL DEFAULT '',
                    normalized_requester_id TEXT NOT NULL DEFAULT '',
                    wa_phone_normalized TEXT NOT NULL DEFAULT '',
                    lid TEXT NOT NULL DEFAULT '',
                    metadata_json TEXT NOT NULL DEFAULT '{}',
                    updated_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS truth_acquisition_logs (
                    acquisition_id TEXT PRIMARY KEY,
                    account_key TEXT NOT NULL DEFAULT '',
                    binding_id TEXT NOT NULL DEFAULT '',
                    trigger TEXT NOT NULL DEFAULT '',
                    final_state TEXT NOT NULL DEFAULT '',
                    trust_status TEXT NOT NULL DEFAULT '',
                    current_truth_written INTEGER NOT NULL DEFAULT 0,
                    latest_probe_written INTEGER NOT NULL DEFAULT 0,
                    result_json TEXT NOT NULL DEFAULT '{}',
                    stages_json TEXT NOT NULL DEFAULT '[]',
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS idx_truth_acquisition_logs_account_binding_created ON truth_acquisition_logs(account_key, binding_id, created_at DESC);

                CREATE TABLE IF NOT EXISTS mcn_operation_tasks (
                    task_id TEXT PRIMARY KEY,
                    task_type TEXT NOT NULL,
                    object_type TEXT NOT NULL,
                    object_key TEXT NOT NULL,
                    idempotency_key TEXT NOT NULL,
                    status TEXT NOT NULL,
                    stage TEXT NOT NULL DEFAULT '',
                    priority INTEGER NOT NULL DEFAULT 100,
                    retry_count INTEGER NOT NULL DEFAULT 0,
                    max_retries INTEGER NOT NULL DEFAULT 3,
                    input_json TEXT NOT NULL DEFAULT '{}',
                    result_json TEXT NOT NULL DEFAULT '{}',
                    error_code TEXT NOT NULL DEFAULT '',
                    error_message TEXT NOT NULL DEFAULT '',
                    created_by TEXT NOT NULL DEFAULT '',
                    created_at TEXT NOT NULL,
                    available_at TEXT NOT NULL DEFAULT '',
                    lease_owner TEXT NOT NULL DEFAULT '',
                    lease_until TEXT NOT NULL DEFAULT '',
                    timeout_seconds INTEGER NOT NULL DEFAULT 60,
                    started_at TEXT,
                    finished_at TEXT,
                    UNIQUE(task_type, idempotency_key)
                );

                CREATE INDEX IF NOT EXISTS idx_mcn_event_ledger_object
                ON mcn_event_ledger(object_type, object_key, event_type, created_at);

                CREATE INDEX IF NOT EXISTS idx_mcn_event_ledger_external
                ON mcn_event_ledger(event_type, external_id);

                CREATE INDEX IF NOT EXISTS idx_mcn_event_ledger_type_created
                ON mcn_event_ledger(event_type, created_at);

                CREATE INDEX IF NOT EXISTS idx_mcn_truth_snapshots_lookup
                ON mcn_truth_snapshots(object_type, object_key, snapshot_type);

                CREATE INDEX IF NOT EXISTS idx_mcn_truth_snapshots_expiry
                ON mcn_truth_snapshots(expires_at);

                CREATE INDEX IF NOT EXISTS idx_mcn_operation_tasks_status
                ON mcn_operation_tasks(status, priority, created_at);

                CREATE TABLE IF NOT EXISTS whatsapp_group_atmosphere_configs (
                    config_name TEXT PRIMARY KEY,
                    enabled INTEGER NOT NULL DEFAULT 1,
                    account_key TEXT NOT NULL,
                    target_group TEXT NOT NULL,
                    group_name TEXT,
                    language TEXT NOT NULL DEFAULT 'en',
                    timezone TEXT NOT NULL DEFAULT 'UTC',
                    worker_base_url TEXT,
                    daily_max_messages INTEGER NOT NULL DEFAULT 4,
                    min_interval_minutes INTEGER NOT NULL DEFAULT 60,
                    max_interval_minutes INTEGER NOT NULL DEFAULT 240,
                    next_due_at TEXT,
                    allowed_windows TEXT NOT NULL DEFAULT '[]',
                    template_pool TEXT NOT NULL DEFAULT '[]',
                    mention_reply_enabled INTEGER NOT NULL DEFAULT 1,
                    faq_rules TEXT NOT NULL DEFAULT '[]',
                    status TEXT NOT NULL DEFAULT 'enabled',
                    last_sent_at TEXT,
                    sent_count_today INTEGER NOT NULL DEFAULT 0,
                    sent_count_date TEXT,
                    scheduler_lease_owner TEXT NOT NULL DEFAULT '',
                    scheduler_lease_until TEXT NOT NULL DEFAULT '',
                    updated_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS whatsapp_group_atmosphere_logs (
                    log_id TEXT PRIMARY KEY,
                    config_name TEXT,
                    account_key TEXT NOT NULL,
                    target_group TEXT NOT NULL,
                    direction TEXT NOT NULL,
                    trigger_type TEXT NOT NULL,
                    message_text TEXT NOT NULL,
                    status TEXT NOT NULL,
                    result_code TEXT,
                    result_reason TEXT,
                    raw_result TEXT NOT NULL DEFAULT '{}',
                    created_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS whatsapp_group_atmosphere_language_profiles (
                    config_name TEXT PRIMARY KEY,
                    language TEXT NOT NULL DEFAULT 'en',
                    sample_count INTEGER NOT NULL DEFAULT 0,
                    frequent_terms TEXT NOT NULL DEFAULT '[]',
                    phrase_samples TEXT NOT NULL DEFAULT '[]',
                    tone_markers TEXT NOT NULL DEFAULT '{}',
                    updated_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS whatsapp_group_atmosphere_chat_records (
                    record_id TEXT PRIMARY KEY,
                    config_name TEXT NOT NULL,
                    sender TEXT,
                    message_text TEXT NOT NULL,
                    created_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS whatsapp_group_atmosphere_candidates (
                    candidate_row_id TEXT PRIMARY KEY,
                    config_name TEXT NOT NULL,
                    candidate_id TEXT NOT NULL,
                    template_id TEXT NOT NULL DEFAULT '',
                    language TEXT NOT NULL DEFAULT '',
                    role_positioning TEXT NOT NULL DEFAULT '',
                    source_type TEXT NOT NULL DEFAULT '',
                    normalized_key TEXT NOT NULL DEFAULT '',
                    semantic_key TEXT NOT NULL DEFAULT '',
                    text TEXT NOT NULL DEFAULT '',
                    safe_to_send INTEGER NOT NULL DEFAULT 0,
                    enabled INTEGER NOT NULL DEFAULT 0,
                    sort_order INTEGER,
                    payload_json TEXT NOT NULL DEFAULT '{}',
                    updated_at TEXT NOT NULL,
                    UNIQUE(config_name, candidate_id)
                );

                CREATE TABLE IF NOT EXISTS whatsapp_group_atmosphere_translation_cache (
                    cache_key TEXT PRIMARY KEY,
                    text TEXT NOT NULL,
                    language TEXT NOT NULL DEFAULT '',
                    region TEXT NOT NULL DEFAULT '',
                    text_zh TEXT NOT NULL DEFAULT '',
                    text_zh_source TEXT NOT NULL DEFAULT '',
                    text_zh_status TEXT NOT NULL DEFAULT '',
                    failure_reason TEXT NOT NULL DEFAULT '',
                    retry_count INTEGER NOT NULL DEFAULT 0,
                    next_retry_at TEXT NOT NULL DEFAULT '',
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS whatsapp_group_atmosphere_phrase_types (
                    type_key TEXT PRIMARY KEY,
                    type_name TEXT NOT NULL,
                    description TEXT,
                    enabled INTEGER NOT NULL DEFAULT 1,
                    is_system INTEGER NOT NULL DEFAULT 0,
                    sort_order INTEGER NOT NULL DEFAULT 100,
                    region_scope TEXT NOT NULL DEFAULT '[]',
                    created_by TEXT,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS whatsapp_group_atmosphere_media_assets (
                    media_id TEXT PRIMARY KEY,
                    filename TEXT NOT NULL,
                    media_path TEXT NOT NULL,
                    mime_type TEXT NOT NULL,
                    file_size INTEGER NOT NULL DEFAULT 0,
                    sha256 TEXT,
                    created_by TEXT,
                    created_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS whatsapp_group_atmosphere_role_bindings (
                    binding_id TEXT PRIMARY KEY,
                    role_key TEXT NOT NULL,
                    account_key TEXT NOT NULL,
                    group_index INTEGER NOT NULL DEFAULT 0,
                    target_group TEXT NOT NULL,
                    group_name TEXT,
                    enabled INTEGER NOT NULL DEFAULT 1,
                    auto_speaking_enabled INTEGER NOT NULL DEFAULT 1,
                    trigger_speaking_enabled INTEGER NOT NULL DEFAULT 1,
                    group_send_permission_enabled INTEGER NOT NULL DEFAULT 1,
                    worker_base_url TEXT,
                    daily_max_messages INTEGER NOT NULL DEFAULT 4,
                    min_interval_minutes INTEGER NOT NULL DEFAULT 60,
                    max_interval_minutes INTEGER NOT NULL DEFAULT 240,
                    randomness_level TEXT NOT NULL DEFAULT 'medium',
                    phrase_send_order TEXT NOT NULL DEFAULT 'random',
                    allowed_windows TEXT NOT NULL DEFAULT '[]',
                    schedule_strategies TEXT NOT NULL DEFAULT '[]',
                    last_sent_at TEXT,
                    sent_count_today INTEGER NOT NULL DEFAULT 0,
                    sent_count_date TEXT,
                    next_due_at TEXT,
                    scheduler_lease_owner TEXT NOT NULL DEFAULT '',
                    scheduler_lease_until TEXT NOT NULL DEFAULT '',
                    status TEXT NOT NULL DEFAULT 'enabled',
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL,
                    UNIQUE(role_key, account_key, group_index)
                );

                CREATE TABLE IF NOT EXISTS whatsapp_group_atmosphere_trigger_rules (
                    rule_id TEXT PRIMARY KEY,
                    relationship_key TEXT NOT NULL,
                    rule_name TEXT NOT NULL,
                    trigger_type TEXT NOT NULL,
                    enabled INTEGER NOT NULL DEFAULT 1,
                    priority INTEGER NOT NULL DEFAULT 0,
                    conditions_json TEXT NOT NULL DEFAULT '{}',
                    message_sequence_json TEXT NOT NULL DEFAULT '[]',
                    send_mode TEXT NOT NULL DEFAULT 'sequence',
                    delay_min_seconds INTEGER NOT NULL DEFAULT 2,
                    delay_max_seconds INTEGER NOT NULL DEFAULT 5,
                    cooldown_seconds INTEGER NOT NULL DEFAULT 60,
                    per_user_cooldown_seconds INTEGER NOT NULL DEFAULT 10,
                    daily_max_triggers INTEGER NOT NULL DEFAULT 0,
                    last_triggered_at TEXT,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS whatsapp_group_atmosphere_trigger_events (
                    event_id TEXT PRIMARY KEY,
                    rule_id TEXT,
                    relationship_key TEXT,
                    binding_id TEXT,
                    account_key TEXT,
                    target_group TEXT,
                    sender_id TEXT,
                    trigger_type TEXT NOT NULL,
                    status TEXT NOT NULL,
                    result_code TEXT,
                    result_reason TEXT,
                    trigger_payload_json TEXT NOT NULL DEFAULT '{}',
                    message_sequence_snapshot TEXT NOT NULL DEFAULT '[]',
                    created_at TEXT NOT NULL,
                    sent_at TEXT
                );

                CREATE TABLE IF NOT EXISTS whatsapp_group_atmosphere_learning_accounts (
                    learning_account_key TEXT PRIMARY KEY,
                    account_name TEXT NOT NULL,
                    region TEXT NOT NULL,
                    language TEXT NOT NULL,
                    enabled INTEGER NOT NULL DEFAULT 1,
                    group_links TEXT NOT NULL DEFAULT '[]',
                    target_role_keys TEXT NOT NULL DEFAULT '[]',
                    daily_learning_time TEXT NOT NULL DEFAULT '03:00',
                    read_recent_hours INTEGER NOT NULL DEFAULT 24,
                    max_messages_per_run INTEGER NOT NULL DEFAULT 300,
                    worker_base_url TEXT,
                    last_learned_at TEXT,
                    last_result_summary TEXT NOT NULL DEFAULT '{}',
                    status TEXT NOT NULL DEFAULT 'pending_login',
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS ops_users (
                    user_id TEXT PRIMARY KEY,
                    username TEXT NOT NULL UNIQUE,
                    password_hash TEXT NOT NULL,
                    role TEXT NOT NULL,
                    display_name TEXT,
                    enabled INTEGER NOT NULL DEFAULT 1,
                    last_login_at TEXT,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS ops_sessions (
                    session_id TEXT PRIMARY KEY,
                    user_id TEXT NOT NULL,
                    session_token_hash TEXT NOT NULL UNIQUE,
                    created_at TEXT NOT NULL,
                    expires_at TEXT NOT NULL,
                    last_seen_at TEXT NOT NULL,
                    ip_address TEXT,
                    user_agent TEXT
                );
                """
            )
        for statement in [
            "ALTER TABLE leads ADD COLUMN occurred_at TEXT",
            "ALTER TABLE leads ADD COLUMN parser_confidence REAL",
            "ALTER TABLE leads ADD COLUMN parser_missing_fields TEXT NOT NULL DEFAULT '[]'",
            "ALTER TABLE leads ADD COLUMN parser_conflicts TEXT NOT NULL DEFAULT '[]'",
            "ALTER TABLE leads ADD COLUMN parser_raw_text TEXT",
            "ALTER TABLE leads ADD COLUMN parser_raw_ocr_text TEXT",
            "ALTER TABLE leads ADD COLUMN parser_version TEXT NOT NULL DEFAULT 'manual_cs_parser_v2'",
            "ALTER TABLE leads ADD COLUMN parser_status TEXT NOT NULL DEFAULT 'unknown'",
            "ALTER TABLE leads ADD COLUMN review_reason_codes TEXT NOT NULL DEFAULT '[]'",
            "ALTER TABLE leads ADD COLUMN routing_decision TEXT",
            "ALTER TABLE leads ADD COLUMN recommended_next_action TEXT",
            "ALTER TABLE leads ADD COLUMN review_status TEXT NOT NULL DEFAULT 'not_needed'",
            "ALTER TABLE leads ADD COLUMN review_notes TEXT",
            "ALTER TABLE leads ADD COLUMN reviewed_by TEXT",
            "ALTER TABLE leads ADD COLUMN reviewed_at TEXT",
            "ALTER TABLE leads ADD COLUMN correction_count INTEGER NOT NULL DEFAULT 0",
            "ALTER TABLE leads ADD COLUMN crm_verified_payload TEXT",
            "ALTER TABLE leads ADD COLUMN crm_verified_app_name TEXT",
            "ALTER TABLE leads ADD COLUMN crm_verified_dept_name TEXT",
            "ALTER TABLE leads ADD COLUMN crm_verified_registration_group TEXT",
            "ALTER TABLE leads ADD COLUMN crm_verified_official_group TEXT",
            "ALTER TABLE leads ADD COLUMN crm_verified_at TEXT",
            "ALTER TABLE operator_notifications ADD COLUMN is_read INTEGER NOT NULL DEFAULT 0",
            "ALTER TABLE operator_notifications ADD COLUMN read_at TEXT",
            "ALTER TABLE operator_notifications ADD COLUMN read_by TEXT",
            "ALTER TABLE intake_bot_presets ADD COLUMN robot_name TEXT",
            "ALTER TABLE automation_tasks ADD COLUMN started_at TEXT",
            "ALTER TABLE automation_tasks ADD COLUMN requires_human_action INTEGER NOT NULL DEFAULT 0",
            "ALTER TABLE automation_tasks ADD COLUMN human_action_type TEXT",
            "ALTER TABLE automation_tasks ADD COLUMN human_action_payload TEXT",
            "ALTER TABLE automation_tasks ADD COLUMN worker_id TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE automation_tasks ADD COLUMN lease_until TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE automation_tasks ADD COLUMN heartbeat_at TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE guild_executors ADD COLUMN guild_backend_token TEXT",
            "ALTER TABLE guild_executors ADD COLUMN app_name TEXT NOT NULL DEFAULT 'linky'",
            "ALTER TABLE guild_executors ADD COLUMN oauth_token TEXT",
            "ALTER TABLE guild_executors ADD COLUMN oauth_token_secret TEXT",
            "ALTER TABLE guild_executors ADD COLUMN platform_backend_url TEXT",
            "ALTER TABLE guild_executors ADD COLUMN platform_authorization TEXT",
            "ALTER TABLE guild_executors ADD COLUMN cms_guild_id TEXT",
            "ALTER TABLE guild_executors ADD COLUMN cms_guild_sid TEXT",
            "ALTER TABLE guild_executors ADD COLUMN country TEXT",
            "ALTER TABLE guild_executors ADD COLUMN guild_country TEXT",
            "ALTER TABLE guild_executors ADD COLUMN eligible_user_countries TEXT NOT NULL DEFAULT '[]'",
            "ALTER TABLE guild_executors ADD COLUMN routing_region TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE leads ADD COLUMN assigned_guild_country TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE leads ADD COLUMN cross_country_fallback INTEGER NOT NULL DEFAULT 0",
            "ALTER TABLE leads ADD COLUMN cross_country_fallback_reason TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE guild_anchor_daily_stats ADD COLUMN guild_executor_key TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE guild_anchor_daily_stats ADD COLUMN guild_display_name TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE guild_anchor_daily_stats ADD COLUMN sort_direction TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE guild_anchor_daily_stats ADD COLUMN sort_confidence TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE guild_anchor_daily_stats ADD COLUMN stale_after TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE guild_anchor_daily_stats ADD COLUMN real_person_count INTEGER NOT NULL DEFAULT 0",
            "ALTER TABLE guild_anchor_daily_stat_jobs ADD COLUMN real_person_count INTEGER",
            "ALTER TABLE guild_anchor_daily_stat_jobs ADD COLUMN last_error TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE guild_anchor_daily_stat_jobs ADD COLUMN recovery_count INTEGER NOT NULL DEFAULT 0",
            "ALTER TABLE guild_anchor_daily_stat_jobs ADD COLUMN last_recovered_at TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE guild_anchor_seen ADD COLUMN created_date_utc TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE guild_anchor_seen ADD COLUMN is_real_person INTEGER NOT NULL DEFAULT 0",
            "ALTER TABLE guild_anchor_seen ADD COLUMN anchor_name TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE guild_anchor_seen ADD COLUMN streamer_sid TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE guild_anchor_seen ADD COLUMN streamer_sid_source_contract TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE guild_anchor_newcomer_identity_snapshots ADD COLUMN streamer_sid TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE registration_group_approval_batch_members ADD COLUMN repair_last_attempt_at TEXT",
            "ALTER TABLE registration_group_approval_batch_members ADD COLUMN repair_last_result TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE registration_group_approval_batch_members ADD COLUMN repair_next_attempt_at TEXT",
            "ALTER TABLE registration_group_approval_batch_members ADD COLUMN display_name_source TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE registration_group_approval_batch_members ADD COLUMN display_name_enhanced_at TEXT",
            "ALTER TABLE registration_group_approval_batch_members ADD COLUMN group_type TEXT NOT NULL DEFAULT 'registration_group'",
            "ALTER TABLE registration_group_approval_batch_members ADD COLUMN lead_id TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE registration_group_approval_batch_members ADD COLUMN matched_customer_id TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE registration_group_approval_batch_members ADD COLUMN registration_status_snapshot TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE registration_group_approval_batch_members ADD COLUMN registration_status_label_snapshot TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE registration_group_approval_batch_members ADD COLUMN eligibility_source TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE registration_group_approval_batch_members ADD COLUMN eligibility_snapshot TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE ops_users ADD COLUMN display_name TEXT",
            "ALTER TABLE ops_users ADD COLUMN enabled INTEGER NOT NULL DEFAULT 1",
            "ALTER TABLE ops_users ADD COLUMN last_login_at TEXT",
            "ALTER TABLE whatsapp_group_atmosphere_configs ADD COLUMN max_interval_minutes INTEGER NOT NULL DEFAULT 240",
            "ALTER TABLE whatsapp_group_atmosphere_configs ADD COLUMN next_due_at TEXT",
            "ALTER TABLE whatsapp_group_atmosphere_logs ADD COLUMN delivery_state TEXT NOT NULL DEFAULT 'unknown'",
            "ALTER TABLE whatsapp_group_atmosphere_logs ADD COLUMN evidence_level TEXT NOT NULL DEFAULT 'none'",
            "ALTER TABLE whatsapp_group_atmosphere_logs ADD COLUMN frontend_verified INTEGER NOT NULL DEFAULT 0",
            "ALTER TABLE whatsapp_group_atmosphere_logs ADD COLUMN client_send_key TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE whatsapp_group_atmosphere_logs ADD COLUMN legacy_status TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE whatsapp_group_atmosphere_logs ADD COLUMN legacy_result_code TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE whatsapp_group_atmosphere_logs ADD COLUMN legacy_message_id TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE whatsapp_group_atmosphere_logs ADD COLUMN migration_note TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE whatsapp_group_atmosphere_logs ADD COLUMN preflight_status TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE whatsapp_group_atmosphere_logs ADD COLUMN preflight_reason TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE whatsapp_group_atmosphere_logs ADD COLUMN preflight_details TEXT NOT NULL DEFAULT '{}'",
            "ALTER TABLE whatsapp_group_atmosphere_logs ADD COLUMN readback_matched INTEGER NOT NULL DEFAULT 0",
            "ALTER TABLE whatsapp_group_atmosphere_logs ADD COLUMN readback_match_reason TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE whatsapp_group_atmosphere_logs ADD COLUMN readback_message_id TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE whatsapp_group_atmosphere_logs ADD COLUMN readback_text TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE whatsapp_group_atmosphere_logs ADD COLUMN readback_timestamp TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE whatsapp_group_atmosphere_logs ADD COLUMN readback_attempt_count INTEGER NOT NULL DEFAULT 0",
            "ALTER TABLE whatsapp_group_atmosphere_role_bindings ADD COLUMN randomness_level TEXT NOT NULL DEFAULT 'medium'",
            "ALTER TABLE whatsapp_group_atmosphere_role_bindings ADD COLUMN phrase_send_order TEXT NOT NULL DEFAULT 'random'",
            "ALTER TABLE whatsapp_group_atmosphere_role_bindings ADD COLUMN trigger_speaking_enabled INTEGER NOT NULL DEFAULT 1",
            "ALTER TABLE whatsapp_group_atmosphere_role_bindings ADD COLUMN schedule_strategies TEXT NOT NULL DEFAULT '[]'",
            "ALTER TABLE whatsapp_group_atmosphere_trigger_rules ADD COLUMN send_mode TEXT NOT NULL DEFAULT 'sequence'",
            "ALTER TABLE whatsapp_group_atmosphere_learning_accounts ADD COLUMN worker_base_url TEXT",
            "ALTER TABLE ops_intake_items ADD COLUMN template_copied_at TEXT",
            "ALTER TABLE ops_intake_items ADD COLUMN template_copied_by TEXT",
            "ALTER TABLE ops_intake_items ADD COLUMN force_feedback_reason TEXT",
            "ALTER TABLE ops_intake_items ADD COLUMN idempotency_key TEXT",
            "ALTER TABLE ops_intake_items ADD COLUMN dedupe_route TEXT",
            "ALTER TABLE ops_intake_items ADD COLUMN route_snapshot TEXT NOT NULL DEFAULT '{}'",
            "ALTER TABLE ops_intake_items ADD COLUMN config_drift TEXT NOT NULL DEFAULT '{}'",
            "ALTER TABLE ops_intake_items ADD COLUMN group_auto_filled INTEGER NOT NULL DEFAULT 0",
            "ALTER TABLE ops_intake_items ADD COLUMN group_auto_fill_source TEXT",
            "ALTER TABLE ops_intake_items ADD COLUMN group_auto_fill_confidence TEXT",
            "ALTER TABLE ops_intake_items ADD COLUMN group_auto_fill_confirmed INTEGER NOT NULL DEFAULT 0",
            "ALTER TABLE ops_intake_items ADD COLUMN group_auto_fill_confirmed_by TEXT",
            "ALTER TABLE ops_intake_items ADD COLUMN group_auto_fill_confirmed_at TEXT",
            "ALTER TABLE ops_intake_items ADD COLUMN source TEXT NOT NULL DEFAULT 'ops_intake_workbench'",
            "ALTER TABLE ops_intake_items ADD COLUMN external_user_id TEXT",
            "ALTER TABLE ops_intake_items ADD COLUMN external_session_id TEXT",
            "ALTER TABLE ops_intake_items ADD COLUMN external_message_id TEXT",
            "ALTER TABLE ops_intake_items ADD COLUMN external_customer_service_id TEXT",
            "ALTER TABLE ops_intake_items ADD COLUMN external_customer_service_name TEXT",
            "ALTER TABLE ops_intake_items ADD COLUMN external_payload TEXT NOT NULL DEFAULT '{}'",
            "ALTER TABLE ops_timo_intake_items ADD COLUMN app_name TEXT NOT NULL DEFAULT 'Timo'",
            "ALTER TABLE ops_timo_intake_items ADD COLUMN guild_name TEXT NOT NULL DEFAULT ''",
            "ALTER TABLE ops_timo_intake_items ADD COLUMN source TEXT NOT NULL DEFAULT 'ops_timo_intake'",
            "ALTER TABLE ops_timo_intake_items ADD COLUMN external_user_id TEXT",
            "ALTER TABLE ops_timo_intake_items ADD COLUMN external_session_id TEXT",
            "ALTER TABLE ops_timo_intake_items ADD COLUMN external_message_id TEXT",
            "ALTER TABLE ops_timo_intake_items ADD COLUMN external_customer_service_id TEXT",
            "ALTER TABLE ops_timo_intake_items ADD COLUMN external_customer_service_name TEXT",
            "ALTER TABLE ops_timo_intake_items ADD COLUMN external_payload TEXT NOT NULL DEFAULT '{}'",
            "ALTER TABLE ops_timo_intake_items ADD COLUMN feedback_status TEXT NOT NULL DEFAULT 'not_feedbackable'",
            "ALTER TABLE ops_timo_intake_items ADD COLUMN feedback_done_at TEXT",
            "ALTER TABLE ops_timo_intake_items ADD COLUMN feedback_done_by TEXT",
            "ALTER TABLE ops_timo_intake_items ADD COLUMN template_copied_at TEXT",
            "ALTER TABLE ops_timo_intake_items ADD COLUMN template_copied_by TEXT",
        ]:
            try:
                conn.execute(statement)
            except sqlite3.OperationalError:
                pass
        apply_schema_migration_registry(conn)
        ensure_timo_incremental_schema(conn)
        ensure_streamer_analytics_views(conn)
        ensure_streamer_roi_tables(conn)
        ensure_im_diagnostics_tables(conn)
        try:
            conn.execute(
                """
                UPDATE whatsapp_group_atmosphere_logs
                SET
                    legacy_status = CASE WHEN COALESCE(legacy_status, '') = '' THEN COALESCE(status, '') ELSE legacy_status END,
                    legacy_result_code = CASE WHEN COALESCE(legacy_result_code, '') = '' THEN COALESCE(result_code, '') ELSE legacy_result_code END,
                    legacy_message_id = CASE
                        WHEN COALESCE(legacy_message_id, '') = ''
                        THEN COALESCE(json_extract(raw_result, '$.message_id'), json_extract(raw_result, '$.raw_result.message_id'), '')
                        ELSE legacy_message_id
                    END,
                    delivery_state = CASE
                        WHEN COALESCE(delivery_state, '') <> '' AND delivery_state <> 'unknown' THEN delivery_state
                        WHEN LOWER(COALESCE(status, '')) IN ('success', 'sent') THEN 'api_accepted'
                        WHEN LOWER(COALESCE(status, '')) IN ('failed', 'error') THEN 'send_failed'
                        ELSE 'unknown'
                    END,
                    evidence_level = CASE
                        WHEN COALESCE(evidence_level, '') <> '' AND evidence_level <> 'none' THEN evidence_level
                        WHEN LOWER(COALESCE(status, '')) IN ('success', 'sent') THEN 'accepted_by_runtime_api'
                        ELSE 'none'
                    END,
                    migration_note = CASE
                        WHEN COALESCE(migration_note, '') <> '' THEN migration_note
                        WHEN LOWER(COALESCE(status, '')) IN ('success', 'sent') THEN 'legacy success means runtime api accepted, not frontend verified'
                        ELSE migration_note
                    END,
                    frontend_verified = CASE
                        WHEN COALESCE(frontend_verified, 0) NOT IN (0, 1) THEN 0
                        ELSE COALESCE(frontend_verified, 0)
                    END
                WHERE
                    COALESCE(legacy_status, '') = ''
                    OR COALESCE(legacy_result_code, '') = ''
                    OR COALESCE(delivery_state, '') IN ('', 'unknown')
                    OR COALESCE(evidence_level, '') IN ('', 'none')
                    OR COALESCE(migration_note, '') = ''
                """
            )
        except sqlite3.OperationalError:
            pass
        for statement in [
            "CREATE INDEX IF NOT EXISTS idx_sync_logs_lead_created_at ON sync_logs (lead_id, created_at)",
            "CREATE INDEX IF NOT EXISTS idx_account_submissions_lead_created_at ON account_submissions (lead_id, created_at)",
            "CREATE INDEX IF NOT EXISTS idx_operator_notifications_lead_created_at ON operator_notifications (lead_id, created_at)",
            "CREATE INDEX IF NOT EXISTS idx_operator_notifications_unread ON operator_notifications (is_read, created_at)",
            "CREATE INDEX IF NOT EXISTS idx_automation_tasks_lead_type_status ON automation_tasks (lead_id, task_type, status, created_at)",
            "CREATE UNIQUE INDEX IF NOT EXISTS idx_guild_anchor_newcomer_snapshot_sid ON guild_anchor_newcomer_identity_snapshots (guild_executor_key, stat_date, streamer_sid) WHERE streamer_sid <> ''",
            "CREATE INDEX IF NOT EXISTS idx_customer_projection_mobile_area_updated ON customer_projection (mobile, area_code, updated_at DESC)",
            "CREATE INDEX IF NOT EXISTS idx_customer_projection_mobile_updated ON customer_projection (mobile, updated_at DESC)",
            "CREATE INDEX IF NOT EXISTS idx_leads_mobile_area_updated ON leads (mobile, area_code, updated_at DESC)",
            "CREATE INDEX IF NOT EXISTS idx_leads_mobile_updated ON leads (mobile, updated_at DESC)",
            "CREATE INDEX IF NOT EXISTS idx_ingress_events_status_created_at ON ingress_events (status, created_at)",
            "CREATE INDEX IF NOT EXISTS idx_ingress_jobs_status_available_at ON ingress_jobs (status, available_at)",
            "CREATE INDEX IF NOT EXISTS idx_registration_group_approval_batch_runs_updated_at ON registration_group_approval_batch_runs (updated_at)",
            "CREATE INDEX IF NOT EXISTS idx_group_atmosphere_phrase_types_enabled_sort ON whatsapp_group_atmosphere_phrase_types (enabled, sort_order)",
            "CREATE INDEX IF NOT EXISTS idx_group_atmosphere_media_sha256 ON whatsapp_group_atmosphere_media_assets (sha256)",
            "CREATE INDEX IF NOT EXISTS idx_registration_group_approval_batch_members_run_idx ON registration_group_approval_batch_members (approval_run_id, batch_index)",
            "CREATE INDEX IF NOT EXISTS idx_registration_group_approval_batch_members_phone_idx ON registration_group_approval_batch_members (wa_phone_normalized, created_at)",
            "CREATE INDEX IF NOT EXISTS idx_rgm_group_type_group_approved ON registration_group_approval_batch_members (group_type, registration_group, approved_at DESC)",
            "CREATE INDEX IF NOT EXISTS idx_rgm_group_approved ON registration_group_approval_batch_members (registration_group, approved_at DESC)",
            "CREATE INDEX IF NOT EXISTS idx_rgm_approved_at ON registration_group_approval_batch_members (approved_at DESC)",
            "CREATE INDEX IF NOT EXISTS idx_rgm_created_at ON registration_group_approval_batch_members (created_at DESC)",
            "CREATE INDEX IF NOT EXISTS idx_rgm_requester ON registration_group_approval_batch_members (requester_id)",
            "CREATE INDEX IF NOT EXISTS idx_rgm_lead ON registration_group_approval_batch_members (lead_id, approved_at DESC)",
            "CREATE INDEX IF NOT EXISTS idx_rgm_customer ON registration_group_approval_batch_members (matched_customer_id, approved_at DESC)",
            "CREATE INDEX IF NOT EXISTS idx_rgm_run ON registration_group_approval_batch_members (approval_run_id)",
            "CREATE INDEX IF NOT EXISTS idx_operator_audit_log_lead_created_at ON operator_audit_log (lead_id, created_at)",
            "CREATE INDEX IF NOT EXISTS idx_group_atmosphere_config_account_group ON whatsapp_group_atmosphere_configs (account_key, target_group)",
            "CREATE INDEX IF NOT EXISTS idx_group_atmosphere_logs_config_created ON whatsapp_group_atmosphere_logs (config_name, created_at)",
            "CREATE INDEX IF NOT EXISTS idx_group_atmosphere_chat_config_created ON whatsapp_group_atmosphere_chat_records (config_name, created_at)",
            "CREATE INDEX IF NOT EXISTS idx_group_atmosphere_translation_cache_status_retry ON whatsapp_group_atmosphere_translation_cache (text_zh_status, retry_count, next_retry_at)",
            "CREATE INDEX IF NOT EXISTS idx_group_atmosphere_bindings_role ON whatsapp_group_atmosphere_role_bindings (role_key, enabled)",
            "CREATE INDEX IF NOT EXISTS idx_group_atmosphere_bindings_account_group ON whatsapp_group_atmosphere_role_bindings (account_key, group_index)",
            "CREATE INDEX IF NOT EXISTS idx_group_atmosphere_learning_enabled ON whatsapp_group_atmosphere_learning_accounts (enabled, region)",
            "CREATE INDEX IF NOT EXISTS idx_ops_users_username ON ops_users (username)",
            "CREATE INDEX IF NOT EXISTS idx_intake_guild_assignees_user ON intake_guild_assignees (user_id, guild_name)",
            "CREATE INDEX IF NOT EXISTS idx_ops_intake_items_guild_feedback ON ops_intake_items (guild_name, feedback_status, created_at)",
            "CREATE INDEX IF NOT EXISTS idx_ops_intake_items_crm_compensation_patrol ON ops_intake_items (COALESCE(processed_at, created_at)) WHERE system_status = 'partial_success_crm_failed' AND COALESCE(feedback_status, '') NOT IN ('feedback_done', 'cleared')",
            "CREATE INDEX IF NOT EXISTS idx_ops_intake_items_submitted_by ON ops_intake_items (submitted_by_user_id, created_at)",
            "CREATE INDEX IF NOT EXISTS idx_ops_intake_items_idempotency ON ops_intake_items (idempotency_key)",
            "CREATE INDEX IF NOT EXISTS idx_ops_intake_items_external_user ON ops_intake_items (source, external_user_id, created_at)",
            "CREATE INDEX IF NOT EXISTS idx_external_phone_backfill_source_account ON external_app_phone_backfill_requests (source, linky_account_id, created_at DESC)",
            "CREATE INDEX IF NOT EXISTS idx_external_phone_backfill_status_created ON external_app_phone_backfill_requests (status, created_at DESC)",
            "CREATE INDEX IF NOT EXISTS idx_ops_timo_intake_status_created ON ops_timo_intake_items (system_status, created_at DESC)",
            "CREATE INDEX IF NOT EXISTS idx_ops_timo_intake_guild_created ON ops_timo_intake_items (guild_name, created_at DESC)",
            "CREATE INDEX IF NOT EXISTS idx_ops_timo_intake_timo_id ON ops_timo_intake_items (timo_id, created_at DESC)",
            "CREATE INDEX IF NOT EXISTS idx_ops_timo_intake_external_user ON ops_timo_intake_items (source, external_user_id, created_at DESC)",
            "CREATE INDEX IF NOT EXISTS idx_ops_intake_history_attempts_created ON ops_intake_binding_history_attempts (created_at DESC, dedupe_key)",
            "CREATE INDEX IF NOT EXISTS idx_ops_intake_history_attempts_dedupe ON ops_intake_binding_history_attempts (dedupe_key, created_at DESC)",
            "CREATE INDEX IF NOT EXISTS idx_ops_intake_history_attempts_guild_created ON ops_intake_binding_history_attempts (guild_name, created_at DESC)",
            "CREATE INDEX IF NOT EXISTS idx_ops_intake_history_attempts_date_created ON ops_intake_binding_history_attempts (created_date_bj, created_at DESC)",
            "CREATE INDEX IF NOT EXISTS idx_ops_intake_history_attempts_submitter ON ops_intake_binding_history_attempts (submitted_by_username, created_at DESC)",
            "CREATE INDEX IF NOT EXISTS idx_ops_intake_history_attempts_status ON ops_intake_binding_history_attempts (current_exception, is_failure, is_duplicate, is_closed, is_success, created_at DESC)",
            "CREATE INDEX IF NOT EXISTS idx_ops_intake_history_attempts_phone_account ON ops_intake_binding_history_attempts (normalized_phone_digits, parsed_account_id, created_at DESC)",
            "CREATE INDEX IF NOT EXISTS idx_ops_sessions_user_id ON ops_sessions (user_id, expires_at)",
            "CREATE INDEX IF NOT EXISTS idx_ops_sessions_expires_at ON ops_sessions (expires_at)",
            "CREATE INDEX IF NOT EXISTS idx_guild_anchor_seen_utc_date ON guild_anchor_seen (guild_executor_key, created_date_utc)",
            "CREATE UNIQUE INDEX IF NOT EXISTS idx_guild_anchor_seen_streamer_sid ON guild_anchor_seen (streamer_sid) WHERE streamer_sid <> ''",
            "CREATE INDEX IF NOT EXISTS idx_timo_anchor_exported_ids_kind ON timo_anchor_exported_ids (export_kind, first_exported_at)",
            "CREATE INDEX IF NOT EXISTS idx_timo_anchor_export_cache_date ON timo_anchor_export_cache (guild_executor_key, export_kind, data_date_bj)",
            "CREATE INDEX IF NOT EXISTS idx_timo_revenue_export_cache_range ON timo_revenue_export_cache (guild_executor_key, export_type, date_from_bj, date_to_bj)",
        ]:
            conn.execute(statement)
        self._backfill_whatsapp_approval_account_created_at(conn)
        self._backfill_whatsapp_approval_baileys_runtime_defaults(conn)
        self._backfill_guild_anchor_seen_created_date_utc(conn)
        conn.commit()

    def _backfill_whatsapp_approval_account_created_at(self, conn: sqlite3.Connection) -> None:
        fallback_now = utc_now()
        conn.execute(
            """
            UPDATE whatsapp_approval_accounts
               SET created_at = COALESCE(NULLIF(updated_at, ''), ?)
             WHERE NULLIF(created_at, '') IS NULL
            """,
            (fallback_now,),
        )

    def _backfill_guild_anchor_seen_created_date_utc(self, conn: sqlite3.Connection) -> None:
        try:
            conn.execute(
                """
                UPDATE guild_anchor_seen
                   SET created_date_utc = date(created_at, 'unixepoch')
                 WHERE COALESCE(created_date_utc, '') = ''
                   AND COALESCE(created_at, 0) > 0
                """
            )
        except sqlite3.OperationalError:
            pass

    def _backfill_whatsapp_approval_baileys_runtime_defaults(self, conn: sqlite3.Connection) -> None:
        try:
            rows = conn.execute(
                """
                SELECT account_key, responsible_type, group_links
                  FROM whatsapp_approval_accounts
                 WHERE responsible_type IN ('registration_group', 'official_group')
                """
            ).fetchall()
        except sqlite3.OperationalError:
            return
        for row in rows:
            account_key = str(row['account_key'] or '').strip()
            responsible_type = str(row['responsible_type'] or '').strip()
            if not account_key or responsible_type not in {'registration_group', 'official_group'}:
                continue
            try:
                parsed = json.loads(str(row['group_links'] or '[]'))
            except Exception:
                parsed = []
            if not isinstance(parsed, list):
                continue
            bindings = [dict(item or {}) for item in parsed if isinstance(item, dict)]
            if not bindings:
                continue
            binding_baileys_account_ids = {first_baileys_account_id(item) for item in bindings}
            binding_baileys_account_ids.discard('')
            if len(binding_baileys_account_ids) > 1:
                continue
            runtime_config = _whatsapp_approval_runtime_config_from_dict(_preferred_group_binding(bindings))
            inherited_baileys_account_id = resolve_baileys_account_id_for_card(
                account_key=account_key,
                explicit_runtime=runtime_config,
                bindings=bindings,
            )
            updated_bindings: list[dict[str, Any]] = []
            changed = False
            for binding in bindings:
                updated_binding = _apply_baileys_runtime_assignment_defaults(
                    binding,
                    responsible_type=responsible_type,
                    baileys_account_id=inherited_baileys_account_id,
                )
                updated_binding['config_fingerprint'] = _whatsapp_approval_binding_config_fingerprint(updated_binding)
                if updated_binding != binding:
                    changed = True
                updated_bindings.append(updated_binding)
            if changed:
                conn.execute(
                    "UPDATE whatsapp_approval_accounts SET group_links = ? WHERE account_key = ?",
                    (json.dumps(updated_bindings, ensure_ascii=False), account_key),
                )


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def parse_iso_datetime(value: str) -> datetime:
    text = str(value or '').strip()
    if text.endswith('Z'):
        text = text[:-1] + '+00:00'
    dt = datetime.fromisoformat(text)
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(timezone.utc)


def create_id(prefix: str) -> str:
    return f"{prefix}_{uuid.uuid4().hex[:12]}"


def _pbkdf2_hash_password(password: str, *, salt: Optional[bytes] = None, iterations: int = 240000) -> str:
    secret = str(password or '')
    if not secret:
        raise ValueError('password_required')
    raw_salt = salt or secrets.token_bytes(16)
    derived = hashlib.pbkdf2_hmac('sha256', secret.encode('utf-8'), raw_salt, int(iterations))
    return 'pbkdf2_sha256${iterations}${salt}${digest}'.format(
        iterations=int(iterations),
        salt=base64.b64encode(raw_salt).decode('ascii'),
        digest=base64.b64encode(derived).decode('ascii'),
    )


def verify_ops_password(password: str, password_hash: str) -> bool:
    try:
        algorithm, iterations_text, salt_b64, digest_b64 = str(password_hash or '').split('$', 3)
        if algorithm != 'pbkdf2_sha256':
            return False
        expected = _pbkdf2_hash_password(
            password,
            salt=base64.b64decode(salt_b64.encode('ascii')),
            iterations=int(iterations_text),
        )
    except Exception:
        return False
    return hmac.compare_digest(expected, str(password_hash or ''))


class OpsAuthManager:
    def __init__(
        self,
        db: Database,
        *,
        session_ttl_hours: int = 12,
        cookie_name: str = OPS_AUTH_SESSION_COOKIE,
        cookie_secure: bool = False,
    ) -> None:
        self.db = db
        self.session_ttl_hours = max(1, int(session_ttl_hours or 12))
        self.cookie_name = str(cookie_name or OPS_AUTH_SESSION_COOKIE).strip() or OPS_AUTH_SESSION_COOKIE
        self.cookie_secure = bool(cookie_secure)

    @staticmethod
    def _serialize_user(row: Optional[sqlite3.Row]) -> Optional[Dict[str, Any]]:
        if row is None:
            return None
        username = str(row['username'] or '').strip()
        display_name = str(row['display_name'] or '').strip()
        serialized = {
            'user_id': row['user_id'],
            'username': username,
            'display_name': display_name or username,
            'role': normalize_ops_role(row['role']),
            'enabled': bool(row['enabled']),
            'last_login_at': row['last_login_at'],
            'created_at': row['created_at'],
            'updated_at': row['updated_at'],
        }
        try:
            if 'session_id' in row.keys():
                serialized['session_id'] = row['session_id']
        except Exception:
            pass
        return serialized

    @staticmethod
    def _hash_session_token(token: str) -> str:
        return hashlib.sha256(str(token or '').encode('utf-8')).hexdigest()

    def has_users(self) -> bool:
        conn = self.db.connect()
        row = conn.execute('SELECT COUNT(*) AS total FROM ops_users').fetchone()
        return bool(row and int(row['total'] or 0) > 0)

    def create_user(
        self,
        *,
        username: str,
        password: str,
        role: str,
        display_name: Optional[str] = None,
        enabled: bool = True,
    ) -> Dict[str, Any]:
        normalized_username = str(username or '').strip().lower()
        if not re.fullmatch(r'[a-z0-9][a-z0-9._-]{2,31}', normalized_username):
            raise ValueError('invalid_username')
        secret = str(password or '')
        if len(secret) < 8:
            raise ValueError('password_too_short')
        now = utc_now()
        conn = self.db.connect()
        with conn:
            existing = conn.execute(
                'SELECT user_id FROM ops_users WHERE username = ?',
                (normalized_username,),
            ).fetchone()
            if existing is not None:
                raise ValueError('username_taken')
            user_id = create_id('ops_user')
            conn.execute(
                '''
                INSERT INTO ops_users (
                    user_id, username, password_hash, role, display_name, enabled, last_login_at, created_at, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                ''',
                (
                    user_id,
                    normalized_username,
                    _pbkdf2_hash_password(secret),
                    normalize_ops_role(role),
                    str(display_name or '').strip() or normalized_username,
                    1 if enabled else 0,
                    None,
                    now,
                    now,
                ),
            )
        return self.get_user_by_id(user_id) or {}

    def bootstrap_admin(self, *, username: str, password: str, display_name: Optional[str] = None) -> Dict[str, Any]:
        if self.has_users():
            raise ValueError('bootstrap_closed')
        return self.create_user(
            username=username,
            password=password,
            role=OPS_AUTH_ROLE_SUPER_ADMIN,
            display_name=display_name,
            enabled=True,
        )

    def get_user_by_id(self, user_id: str) -> Optional[Dict[str, Any]]:
        conn = self.db.connect()
        row = conn.execute(
            'SELECT user_id, username, role, display_name, enabled, last_login_at, created_at, updated_at FROM ops_users WHERE user_id = ?',
            (str(user_id or '').strip(),),
        ).fetchone()
        return self._serialize_user(row)

    def get_user_by_username(self, username: str) -> Optional[Dict[str, Any]]:
        normalized_username = str(username or '').strip().lower()
        if not normalized_username:
            return None
        conn = self.db.connect()
        row = conn.execute(
            'SELECT user_id, username, role, display_name, enabled, last_login_at, created_at, updated_at FROM ops_users WHERE username = ?',
            (normalized_username,),
        ).fetchone()
        return self._serialize_user(row)

    def authenticate(self, username: str, password: str) -> Optional[Dict[str, Any]]:
        normalized_username = str(username or '').strip().lower()
        conn = self.db.connect()
        row = conn.execute('SELECT * FROM ops_users WHERE username = ?', (normalized_username,)).fetchone()
        if row is None or not bool(row['enabled']):
            return None
        if not verify_ops_password(password, str(row['password_hash'] or '')):
            return None
        now = utc_now()
        with conn:
            conn.execute('UPDATE ops_users SET last_login_at = ?, updated_at = ? WHERE user_id = ?', (now, now, row['user_id']))
        return self.get_user_by_id(row['user_id'])

    def list_users(self) -> List[Dict[str, Any]]:
        conn = self.db.connect()
        rows = conn.execute(
            'SELECT user_id, username, role, display_name, enabled, last_login_at, created_at, updated_at FROM ops_users ORDER BY created_at ASC'
        ).fetchall()
        return [user for user in (self._serialize_user(row) for row in rows) if user]

    def update_user(
        self,
        user_id: str,
        *,
        role: Optional[str] = None,
        display_name: Optional[str] = None,
        enabled: Optional[bool] = None,
        password: Optional[str] = None,
    ) -> Dict[str, Any]:
        normalized_user_id = str(user_id or '').strip()
        conn = self.db.connect()
        row = conn.execute('SELECT * FROM ops_users WHERE user_id = ?', (normalized_user_id,)).fetchone()
        if row is None:
            raise ValueError('user_not_found')
        updates = []
        params: List[Any] = []
        if role is not None:
            updates.append('role = ?')
            params.append(normalize_ops_role(role))
        if display_name is not None:
            normalized_display_name = str(display_name or '').strip() or str(row['username'] or '').strip()
            updates.append('display_name = ?')
            params.append(normalized_display_name)
        if enabled is not None:
            updates.append('enabled = ?')
            params.append(1 if enabled else 0)
        if password is not None:
            secret = str(password or '')
            if len(secret) < 8:
                raise ValueError('password_too_short')
            updates.append('password_hash = ?')
            params.append(_pbkdf2_hash_password(secret))
        if not updates:
            return self.get_user_by_id(normalized_user_id) or {}
        params.extend([utc_now(), normalized_user_id])
        with conn:
            conn.execute(f"UPDATE ops_users SET {', '.join(updates)}, updated_at = ? WHERE user_id = ?", tuple(params))
        return self.get_user_by_id(normalized_user_id) or {}

    def delete_user(self, user_id: str) -> bool:
        normalized_user_id = str(user_id or '').strip()
        if not normalized_user_id:
            raise ValueError('user_not_found')
        conn = self.db.connect()
        row = conn.execute('SELECT user_id FROM ops_users WHERE user_id = ?', (normalized_user_id,)).fetchone()
        if row is None:
            raise ValueError('user_not_found')
        with conn:
            conn.execute('DELETE FROM ops_sessions WHERE user_id = ?', (normalized_user_id,))
            conn.execute('DELETE FROM ops_users WHERE user_id = ?', (normalized_user_id,))
        return True

    def change_user_password(self, user_id: str, *, current_password: str, new_password: str) -> Dict[str, Any]:
        normalized_user_id = str(user_id or '').strip()
        conn = self.db.connect()
        row = conn.execute('SELECT * FROM ops_users WHERE user_id = ?', (normalized_user_id,)).fetchone()
        if row is None or not bool(row['enabled']):
            raise ValueError('user_not_found')
        if not verify_ops_password(str(current_password or ''), str(row['password_hash'] or '')):
            raise ValueError('invalid_current_password')
        secret = str(new_password or '')
        if len(secret) < 8:
            raise ValueError('password_too_short')
        with conn:
            conn.execute('UPDATE ops_users SET password_hash = ?, updated_at = ? WHERE user_id = ?', (_pbkdf2_hash_password(secret), utc_now(), normalized_user_id))
        return self.get_user_by_id(normalized_user_id) or {}

    def create_session(self, user: Dict[str, Any], *, ip_address: Optional[str] = None, user_agent: Optional[str] = None) -> str:
        raw_token = secrets.token_urlsafe(32)
        now_dt = datetime.now(timezone.utc)
        now = now_dt.isoformat()
        expires_at = (now_dt + timedelta(hours=self.session_ttl_hours)).isoformat()
        conn = self.db.connect()
        with conn:
            conn.execute('DELETE FROM ops_sessions WHERE expires_at <= ?', (now,))
            conn.execute(
                '''
                INSERT INTO ops_sessions (
                    session_id, user_id, session_token_hash, created_at, expires_at, last_seen_at, ip_address, user_agent
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                ''',
                (
                    create_id('ops_session'),
                    user['user_id'],
                    self._hash_session_token(raw_token),
                    now,
                    expires_at,
                    now,
                    str(ip_address or '').strip() or None,
                    str(user_agent or '').strip() or None,
                ),
            )
        return raw_token

    def revoke_session(self, raw_token: Optional[str]) -> None:
        token = str(raw_token or '').strip()
        if not token:
            return
        conn = self.db.connect()
        with conn:
            conn.execute('DELETE FROM ops_sessions WHERE session_token_hash = ?', (self._hash_session_token(token),))

    def _refresh_session_activity(self, session_id: str, now_dt: datetime) -> str:
        now = now_dt.isoformat()
        expires_at = (now_dt + timedelta(hours=self.session_ttl_hours)).isoformat()
        conn = self.db.connect()
        with conn:
            conn.execute(
                'UPDATE ops_sessions SET last_seen_at = ?, expires_at = ? WHERE session_id = ?',
                (now, expires_at, session_id),
            )
        return expires_at

    def session_user(
        self,
        raw_token: Optional[str],
        *,
        refresh_activity: bool = True,
    ) -> Optional[Dict[str, Any]]:
        token = str(raw_token or '').strip()
        if not token:
            return None
        now_dt = datetime.now(timezone.utc)
        now = now_dt.isoformat()
        query = '''
        SELECT u.user_id, u.username, u.role, u.display_name, u.enabled, u.last_login_at, u.created_at, u.updated_at, s.session_id
        FROM ops_sessions s
        JOIN ops_users u ON u.user_id = s.user_id
        WHERE s.session_token_hash = ? AND s.expires_at > ?
        '''
        params = (self._hash_session_token(token), now)
        if refresh_activity or self.db.db_path == ':memory:':
            row = self.db.connect().execute(query, params).fetchone()
        else:
            readonly_uri = f'file:{quote(str(Path(self.db.db_path).resolve()))}?mode=ro'
            conn = connect_observed_sqlite(
                readonly_uri,
                source='app.main:ops-auth-read',
                timeout=5.0,
                uri=True,
            )
            try:
                conn.row_factory = sqlite3.Row
                conn.execute('PRAGMA busy_timeout=5000')
                conn.execute('PRAGMA query_only=ON')
                row = conn.execute(query, params).fetchone()
            finally:
                conn.close()
        if row is None or not bool(row['enabled']):
            return None
        if refresh_activity:
            try:
                self._refresh_session_activity(row['session_id'], now_dt)
            except sqlite3.OperationalError as exc:
                if 'database is locked' not in str(exc).lower():
                    raise
        return self._serialize_user(row)

    def apply_session_cookie(self, response: Response, raw_token: str) -> None:
        response.set_cookie(
            key=self.cookie_name,
            value=raw_token,
            httponly=True,
            samesite='lax',
            secure=self.cookie_secure,
            path='/',
            max_age=self.session_ttl_hours * 3600,
        )

    def clear_session_cookie(self, response: Response) -> None:
        response.delete_cookie(key=self.cookie_name, path='/')


def _run_registration_group_executor_warmup(executor: Any) -> None:
    try:
        executor.warmup()
    except Exception as exc:
        print(f'Registration group executor warmup degraded at startup: {exc}')


def _schedule_registration_group_executor_warmup(executor: Any) -> str:
    if executor is None or not hasattr(executor, 'warmup') or not callable(getattr(executor, 'warmup')):
        return 'unsupported'
    try:
        asyncio.get_running_loop()
    except RuntimeError:
        _run_registration_group_executor_warmup(executor)
        return 'inline'

    thread = threading.Thread(
        target=_run_registration_group_executor_warmup,
        args=(executor,),
        name='registration-group-executor-warmup',
        daemon=True,
    )
    thread.start()
    return 'threaded_deferred_inside_asyncio_loop'


class LiveLarkReplyAdapter:
    _MESSAGE_MAX_ATTEMPTS = 6

    def __init__(self, *, app_id: str, app_secret: str, domain: str = 'lark') -> None:
        self.app_id = app_id
        self.app_secret = app_secret
        self.base_url = 'https://open.larksuite.com' if domain == 'lark' else 'https://open.feishu.cn'
        self._tenant_access_token: Optional[str] = None

    def _normalize_text_markup(self, text: str) -> str:
        normalized = str(text or '')
        normalized = re.sub(r'\*\*(.+?)\*\*', lambda m: f"<b>{m.group(1)}</b>", normalized)
        return normalized

    def _get_tenant_access_token(self) -> str:
        if self._tenant_access_token:
            return self._tenant_access_token
        response = requests.post(
            f"{self.base_url}/open-apis/auth/v3/tenant_access_token/internal",
            json={"app_id": self.app_id, "app_secret": self.app_secret},
            timeout=15,
        )
        response.raise_for_status()
        body = response.json()
        if body.get('code') != 0:
            raise RuntimeError(f"tenant_access_token failed: {body}")
        self._tenant_access_token = body['tenant_access_token']
        return self._tenant_access_token

    @staticmethod
    def _retry_after_delay_seconds(value: Optional[str], *, attempt: int) -> float:
        raw = str(value or '').strip()
        if raw:
            try:
                parsed = float(raw)
            except (TypeError, ValueError):
                parsed = 0.0
            if parsed > 0:
                return min(max(parsed, 1.0), 30.0)
        return min(2.0 * (2 ** max(attempt - 1, 0)), 30.0)

    def _post_im_message(self, *, url: str, payload: dict) -> dict:
        token = self._get_tenant_access_token()
        last_response = None
        for attempt in range(1, self._MESSAGE_MAX_ATTEMPTS + 1):
            response = requests.post(
                url,
                headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json; charset=utf-8"},
                json=payload,
                timeout=15,
            )
            last_response = response
            if response.status_code == 429 and attempt < self._MESSAGE_MAX_ATTEMPTS:
                time.sleep(self._retry_after_delay_seconds(response.headers.get('Retry-After'), attempt=attempt))
                continue
            response.raise_for_status()
            body = response.json()
            if body.get('code') != 0:
                raise RuntimeError(f"im_message failed: {body}")
            return body
        if last_response is not None:
            last_response.raise_for_status()
        raise RuntimeError('im_message_failed_without_response')

    def reply_text(self, *, message_id: str, text: str) -> dict:
        normalized_text = self._normalize_text_markup(text)
        return self._post_im_message(
            url=f"{self.base_url}/open-apis/im/v1/messages/{message_id}/reply",
            payload={"msg_type": "text", "content": json.dumps({"text": normalized_text}, ensure_ascii=False)},
        )

    def send_text(self, *, chat_id: str, text: str) -> dict:
        if should_suppress_lark_alert(message_text=text):
            return {
                'code': 0,
                'suppressed': True,
                'suppressed_reason': 'invalid_registration_group_invite_404',
            }
        normalized_text = self._normalize_text_markup(text)
        return self._post_im_message(
            url=f"{self.base_url}/open-apis/im/v1/messages?receive_id_type=chat_id",
            payload={"receive_id": chat_id, "msg_type": "text", "content": json.dumps({"text": normalized_text}, ensure_ascii=False)},
        )


class GoogleTranslateCandidateTranslator:
    def __init__(self, *, base_url: str = '', timeout_seconds: float = 20.0) -> None:
        self.base_url = str(base_url or '').strip().rstrip('/') or 'https://translate.googleapis.com'
        self.timeout_seconds = max(3.0, float(timeout_seconds or 20.0))

    def translate(self, text: str, *, role: str = '', language: str = '', region: str = '') -> Dict[str, Any]:
        value = str(text or '').strip()
        if not value:
            raise RuntimeError('translator_empty_text')
        source_language = _normalize_translation_source_language(language, region)
        response = requests.get(
            f'{self.base_url}/translate_a/single',
            params={'client': 'gtx', 'sl': source_language, 'tl': 'zh-CN', 'dt': 't', 'q': value},
            timeout=self.timeout_seconds,
        )
        if response.status_code >= 400:
            raise RuntimeError(f'googletranslate_http_{response.status_code}')
        body = response.json()
        chunks = body[0] if isinstance(body, list) and body else []
        text_zh = ''.join(str(chunk[0] or '') for chunk in chunks if isinstance(chunk, list) and chunk)
        text_zh = re.sub(r'\s+', ' ', text_zh).strip()
        if not text_zh:
            raise RuntimeError('googletranslate_empty_result')
        return {'text_zh': text_zh[:1500], 'status': 'ok', 'source': 'google'}


class LibreTranslateCandidateTranslator:
    def __init__(self, *, base_url: str = '', api_key: str = '', timeout_seconds: float = 20.0) -> None:
        self.base_url = str(base_url or '').strip().rstrip('/') or 'http://127.0.0.1:5000'
        self.api_key = str(api_key or '').strip()
        self.timeout_seconds = max(3.0, float(timeout_seconds or 20.0))

    def translate(self, text: str, *, role: str = '', language: str = '', region: str = '') -> Dict[str, Any]:
        value = str(text or '').strip()
        if not value:
            raise RuntimeError('translator_empty_text')
        source_language = _normalize_translation_source_language(language, region)
        payload = {'q': value, 'source': source_language, 'target': 'zh', 'format': 'text'}
        if self.api_key:
            payload['api_key'] = self.api_key
        response = requests.post(f'{self.base_url}/translate', json=payload, timeout=self.timeout_seconds)
        if response.status_code >= 400:
            raise RuntimeError(f'libretranslate_http_{response.status_code}')
        body = response.json()
        text_zh = str(body.get('translatedText') or '').strip()
        if not text_zh:
            raise RuntimeError('libretranslate_empty_result')
        return {'text_zh': text_zh[:1500], 'status': 'needs_review', 'source': 'libretranslate'}


class GroupAtmosphereAiTranslator:
    def __init__(self, *, api_key: str = '', base_url: str = '', model: str = '', timeout_seconds: float = 20.0) -> None:
        self.api_key = str(api_key or '').strip()
        self.base_url = str(base_url or '').strip().rstrip('/') or 'https://api.openai.com/v1'
        self.model = str(model or '').strip() or 'gpt-4o-mini'
        self.timeout_seconds = max(3.0, float(timeout_seconds or 20.0))

    def translate(self, text: str, *, role: str = '', language: str = '', region: str = '') -> Dict[str, Any]:
        value = str(text or '').strip()
        if not value or not self.api_key:
            raise RuntimeError('translator_not_configured')
        source_language = _normalize_translation_source_language(language, region)
        source_language_name = {'id': '印尼语', 'es': '西班牙语', 'pt': '葡萄牙语'}.get(source_language, source_language or '未知语种')
        prompt = (
            '你是专门服务 WhatsApp 群运营话术的多语种中文翻译，重点处理印尼语、西语、葡语里混杂的口语、缩写、emoji、'
            'Linky/diamond/ID/admin/code 等业务词。请根据地区和语言线索识别源语种，把下面 MCN WhatsApp 群运营话术翻译成运营能看懂的自然中文。\n'
            '硬性要求：\n'
            '1. 输出必须是准确中文翻译，保留原文的段落/换行、称呼和语气；不要总结成“这段话的意思是”，不要加“大意：”。\n'
            '2. 除 Linky、WhatsApp、ID、URL、@用户名外，不要保留源语言单词。diamond=钻石，reward=奖励/收益，screenshot=截图，admin=管理员，grup/grupo=群，kak=朋友/同学或省略。\n'
            '3. 只翻译含义，不改写运营策略，不增删承诺，不加免责声明。emoji 可保留或翻译，但不要丢失业务语气。\n'
            '4. 如果缩写、断句或上下文不确定，仍给出最可能的中文，但 status 设为 needs_review。\n'
            '只输出合法 JSON：{"text_zh":"...","status":"ok或needs_review"}。\n'
            f'地区：{region or "未知"}；推断源语种：{source_language_name}；语言代码：{language or "未知"}；类型：{role or "未知"}\n'
            f'原文：\n{value}'
        )
        max_tokens = min(1200, max(360, int(len(value) * 0.9) + 220))
        response = requests.post(
            f'{self.base_url}/chat/completions',
            headers={'Authorization': f'Bearer {self.api_key}', 'Content-Type': 'application/json'},
            json={
                'model': self.model,
                'messages': [
                    {'role': 'system', 'content': 'Return compact JSON only.'},
                    {'role': 'user', 'content': prompt},
                ],
                'temperature': 0.1,
                'max_tokens': max_tokens,
            },
            timeout=self.timeout_seconds,
        )
        if response.status_code >= 400:
            raise RuntimeError(f'translator_http_{response.status_code}')
        body = response.json()
        content = str((((body.get('choices') or [{}])[0]).get('message') or {}).get('content') or '').strip()
        match = re.search(r'\{.*\}', content, flags=re.S)
        parsed = json.loads(match.group(0) if match else content)
        text_zh = str(parsed.get('text_zh') or '').strip()
        status = str(parsed.get('status') or 'ok').strip() or 'ok'
        if not text_zh:
            raise RuntimeError('translator_empty_result')
        return {'text_zh': text_zh[:1500], 'status': status if status in {'ok', 'needs_review'} else 'ok', 'source': 'ai'}



__all__ = [name for name in globals() if not name.startswith('__')]
